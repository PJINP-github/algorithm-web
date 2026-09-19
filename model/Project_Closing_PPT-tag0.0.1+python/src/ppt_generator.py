"""结项 PPT 核心生成逻辑。"""
from __future__ import annotations

import copy
import os
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .data_loader import (
    filter_by_station,
    load_rows,
    normalize_percent,
    to_int,
)
from .image_generator import image_aspect_ratio, render_table_image

_MODEL_ROOT = Path(__file__).resolve().parent.parent
_CONFIG_PATH = _MODEL_ROOT / "config.yaml"
_TEMPLATE_PATH = _MODEL_ROOT / "model-aiglrithm.pptx"

EMU_PER_PX = 9525  # 96 DPI


# ---------------------------------------------------------------------------
# 配置
# ---------------------------------------------------------------------------

_DEFAULT_CONFIG: dict[str, Any] = {
    "accuracy_threshold": 95.0,
    "low_accuracy_per_page": 6,
    "accuracy_field": "准确率",
    "xlsxway": [],
    "page_titles": {
        "summary": "整体汇总",
        "low_accuracy": "准确率未达标详细",
    },
    "text_style": {
        "font_size": 14,
        "color": "000000",
        "top_px": 42,
        "left_px": 30,
    },
    "summary_image": {
        "width_ratio": 0.60,
        "max_height_ratio": 0.92,
        "min_area_ratio": 0.50,
        "right_padding_px": 20,
        "text_column_ratio": 0.36,
    },
    "low_accuracy_layout": {
        "title_space_after_pt": 10,
        "name_space_after_pt": 3,
        "info_space_after_pt": 3,
        "gap_space_after_pt": 14,
        "image_placeholder_width_px": 520,
        "image_placeholder_height_px": 600,
    },
    "table_style": {
        "header_bg": "#4F81BD",
        "header_fg": "#FFFFFF",
        "header_bold": True,
        "zebra_even_bg": "#F2F2F2",
        "zebra_odd_bg": "#FFFFFF",
        "show_index": True,
        "index_header": "序号",
        "cell_edge_color": "#BFBFBF",
        "cell_edge_width": 0.6,
    },
}


def _deep_merge(base: dict, override: dict) -> dict:
    out = dict(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _deep_merge(out[k], v)
        else:
            out[k] = v
    return out


def load_config() -> dict[str, Any]:
    raw: dict[str, Any] = {}
    if _CONFIG_PATH.exists():
        with _CONFIG_PATH.open("r", encoding="utf-8") as f:
            raw = yaml.safe_load(f) or {}
    return _deep_merge(_DEFAULT_CONFIG, raw)


def _silent_substrings(config: dict[str, Any]) -> list[str]:
    items = config.get("Specify_silent_mode_group") or []
    result: list[str] = []
    for item in items:
        if isinstance(item, dict):
            sub = item.get("substring")
            if sub:
                result.append(str(sub))
        elif isinstance(item, str):
            result.append(item)
    return result


def _px_to_emu(px: float) -> int:
    return int(round(float(px) * EMU_PER_PX))


def _hex_to_rgb(color: str) -> RGBColor:
    s = str(color or "000000").lstrip("#")
    if len(s) != 6:
        s = "000000"
    try:
        return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
    except ValueError:
        return RGBColor(0, 0, 0)


# ---------------------------------------------------------------------------
# 路径解析
# ---------------------------------------------------------------------------

def _normalize_path(raw: Any) -> Path | None:
    if raw is None:
        return None
    s = str(raw).strip().strip('"').strip("'")
    if not s:
        return None
    p = Path(s)
    if p.exists():
        return p
    if os.name == "posix" and "\\" in s:
        alt = Path(s.replace("\\", "/"))
        if alt.exists():
            return alt
    return p


def _resolve_config_xlsx(config: dict[str, Any]) -> Path | None:
    items = config.get("xlsxway") or []
    if isinstance(items, str):
        items = [items]

    candidates: list[tuple[float, Path]] = []
    for raw in items:
        p = _normalize_path(raw)
        if p is None:
            continue
        try:
            if p.exists() and p.is_file():
                candidates.append((p.stat().st_mtime, p))
        except OSError:
            continue

    if not candidates:
        return None
    candidates.sort(key=lambda item: item[0], reverse=True)
    return candidates[0][1]


# ---------------------------------------------------------------------------
# 数据整理
# ---------------------------------------------------------------------------

def _is_silent(name: str, substrings: list[str]) -> bool:
    s = str(name or "")
    return any(sub in s for sub in substrings)


def _classify(
    rows: list[dict[str, Any]], substrings: list[str]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    patrol, silent = [], []
    for row in rows:
        name = str(row.get("算法小类") or "")
        (silent if _is_silent(name, substrings) else patrol).append(row)
    return patrol, silent


def _aggregate(rows: list[dict[str, Any]]) -> dict[str, Any]:
    total_points = 0
    total_correct = 0
    total_blur = 0
    total_blur_after = 0
    for row in rows:
        p = to_int(row.get("点位数量"))
        c = to_int(row.get("准确数量"))
        b = to_int(row.get("拍照模糊数量"))
        ba = to_int(row.get("审核后拍照模糊数量"))
        if p is not None and p > 0 and c is not None:
            total_points += p
            total_correct += c
        if b is not None:
            total_blur += b
        if ba is not None:
            total_blur_after += ba
    if total_points <= 0:
        return {
            "accuracy": None,
            "points": total_points,
            "correct": total_correct,
            "blur": total_blur,
            "blur_after": total_blur_after,
        }
    return {
        "accuracy": round(total_correct / total_points * 100.0, 2),
        "points": total_points,
        "correct": total_correct,
        "blur": total_blur,
        "blur_after": total_blur_after,
    }


def _collect_low_accuracy(
    rows: list[dict[str, Any]],
    accuracy_field: str,
    threshold: float,
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in rows:
        acc = normalize_percent(row.get(accuracy_field))
        if acc is None:
            continue
        if acc < threshold:
            out.append({**row, "_acc": acc})
    out.sort(key=lambda r: r["_acc"])
    return out


# ---------------------------------------------------------------------------
# PPT 基础工具
# ---------------------------------------------------------------------------

def _copy_slide(prs: Presentation, source_index: int):
    source = prs.slides[source_index]
    new_slide = prs.slides.add_slide(source.slide_layout)
    for shp in list(new_slide.shapes):
        shp._element.getparent().remove(shp._element)
    for shp in source.shapes:
        new_slide.shapes._spTree.append(copy.deepcopy(shp._element))
    return new_slide


def _delete_slide(prs: Presentation, index: int) -> None:
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    if 0 <= index < len(slides):
        xml_slides.remove(slides[index])


def _clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def _add_rich_textbox(
    slide,
    left,
    top,
    width,
    height,
    paragraphs: list[dict[str, Any]],
):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    first = True
    for para in paragraphs:
        text = str(para.get("text", ""))
        size = int(para.get("size", 14))
        bold = bool(para.get("bold", False))
        color = para.get("color") or RGBColor(0, 0, 0)
        align = para.get("align", PP_ALIGN.LEFT)
        space_after = para.get("space_after")

        for line in text.split("\n"):
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
            p.alignment = align
            if space_after is not None:
                p.space_after = Pt(float(space_after))
            run = p.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return tb


def _add_placeholder(slide, left, top, width, height, label: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xF7, 0xF7, 0xF7)
    shape.line.color.rgb = RGBColor(0x8C, 0x8C, 0x8C)

    ln = shape.line._get_or_add_ln()
    for dash in ln.findall(qn("a:prstDash")):
        ln.remove(dash)
    ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x7A, 0x7A, 0x7A)
    return shape


def _place_image_right_center(
    slide,
    slide_w: int,
    slide_h: int,
    image_path: Path,
    *,
    text_column_ratio: float,
    right_padding_px: float,
    width_ratio: float,
    max_height_ratio: float,
    min_area_ratio: float,
    top_margin_px: float = 0.0,
) -> None:
    right_pad = _px_to_emu(right_padding_px)
    region_left = int(slide_w * float(text_column_ratio))
    region_w = slide_w - region_left - right_pad
    if region_w <= 0:
        region_w = slide_w // 2

    avail_w = min(region_w, int(slide_w * float(width_ratio)))
    avail_h = int(slide_h * float(max_height_ratio)) - _px_to_emu(top_margin_px)
    if avail_h <= 0:
        avail_h = slide_h // 2

    aspect = image_aspect_ratio(image_path)

    w = avail_w
    h = int(w / aspect) if aspect > 0 else avail_h
    if h > avail_h:
        h = avail_h
        w = int(h * aspect) if aspect > 0 else avail_w

    min_area = float(min_area_ratio) * slide_w * slide_h
    if w * h < min_area:
        h2 = avail_h
        w2 = int(h2 * aspect) if aspect > 0 else avail_w
        if w2 <= avail_w:
            w, h = w2, h2

    left = region_left + (region_w - w) // 2
    top = (slide_h - h) // 2

    slide.shapes.add_picture(
        str(image_path), Emu(left), Emu(top), width=Emu(w), height=Emu(h)
    )


# ---------------------------------------------------------------------------
# 页面渲染
# ---------------------------------------------------------------------------

def _text_metrics(slide_w: int, slide_h: int, config: dict[str, Any]):
    text_style = config["text_style"]
    left_px = float(text_style.get("left_px", 30))
    top_px = float(text_style.get("top_px", 42))

    left = _px_to_emu(left_px)
    top = _px_to_emu(top_px)
    width = slide_w - left - _px_to_emu(left_px)
    if width <= 0:
        width = int(slide_w * 0.5)
    height = slide_h - top - _px_to_emu(top_px)
    if height <= 0:
        height = slide_h - top
    return left, top, width, height, left_px, top_px


def _render_summary_slide(
    slide,
    slide_w: int,
    slide_h: int,
    *,
    station_name: str,
    project_name: str,
    data: dict[str, Any],
    summary_img: Path | None,
    config: dict[str, Any],
) -> None:
    _clear_slide(slide)

    titles = config["page_titles"]
    text_style = config["text_style"]
    img_cfg = config["summary_image"]

    font_size = int(text_style.get("font_size", 14))
    color = _hex_to_rgb(text_style.get("color", "000000"))

    left, top, width, height, left_px, top_px = _text_metrics(
        slide_w, slide_h, config
    )

    title_size = font_size + 8
    patrol = data["patrol"]
    silent = data["silent"]

    # 已移除 "生成时间" 行；"低准确率阈值" 也不再显示字段名
    paragraphs: list[dict[str, Any]] = [
        {
            "text": str(titles.get("summary", "整体汇总")),
            "size": title_size,
            "bold": True,
            "color": color,
            "space_after": 8,
        },
        {"text": f"项目名称：{project_name or '-'}", "size": font_size, "color": color},
        {"text": f"站点名称：{station_name}", "size": font_size, "color": color},
        {"text": "", "size": font_size, "color": color},
        {"text": "【巡视层】", "size": font_size, "bold": True, "color": color},
        {"text": f"整体准确率：{_fmt_acc(patrol['accuracy'])}", "size": font_size, "color": color},
        {"text": f"点位数量：{_fmt_num(patrol['points'])}", "size": font_size, "color": color},
        {"text": f"准确数量：{_fmt_num(patrol['correct'])}", "size": font_size, "color": color},
        {"text": f"拍照模糊数量：{_fmt_num(patrol['blur'])}", "size": font_size, "color": color},
        {
            "text": f"审核后拍照模糊数量：{_fmt_num(patrol['blur_after'])}",
            "size": font_size,
            "color": color,
        },
        {"text": "", "size": font_size, "color": color},
        {"text": "【静默层】", "size": font_size, "bold": True, "color": color},
        {"text": f"整体准确率：{_fmt_acc(silent['accuracy'])}", "size": font_size, "color": color},
        {"text": f"点位数量：{_fmt_num(silent['points'])}", "size": font_size, "color": color},
        {"text": f"准确数量：{_fmt_num(silent['correct'])}", "size": font_size, "color": color},
        {"text": f"拍照模糊数量：{_fmt_num(silent['blur'])}", "size": font_size, "color": color},
        {
            "text": f"审核后拍照模糊数量：{_fmt_num(silent['blur_after'])}",
            "size": font_size,
            "color": color,
        },
        {"text": "", "size": font_size, "color": color},
        {
            "text": f"低准确率阈值：{data['threshold']}%",
            "size": font_size,
            "color": color,
        },
        {
            "text": f"低准确率算法数量：{len(data['low_accuracy'])}",
            "size": font_size,
            "color": color,
        },
    ]

    _add_rich_textbox(slide, left, top, width, height, paragraphs)

    if summary_img and Path(summary_img).exists():
        _place_image_right_center(
            slide,
            slide_w,
            slide_h,
            Path(summary_img),
            text_column_ratio=float(img_cfg.get("text_column_ratio", 0.36)),
            right_padding_px=float(img_cfg.get("right_padding_px", 20)),
            width_ratio=float(img_cfg.get("width_ratio", 0.60)),
            max_height_ratio=float(img_cfg.get("max_height_ratio", 0.92)),
            min_area_ratio=float(img_cfg.get("min_area_ratio", 0.50)),
            top_margin_px=top_px,
        )
    else:
        region_left = int(slide_w * float(img_cfg.get("text_column_ratio", 0.36)))
        _add_placeholder(
            slide,
            region_left,
            top,
            slide_w - region_left - _px_to_emu(20),
            slide_h - top - _px_to_emu(20),
            "（汇总表格图缺失）",
        )


def _render_low_accuracy_slide(
    slide,
    slide_w: int,
    slide_h: int,
    *,
    items: list[dict[str, Any]],
    global_offset: int,
    threshold: float,
    config: dict[str, Any],
) -> None:
    _clear_slide(slide)

    titles = config["page_titles"]
    text_style = config["text_style"]
    low_cfg = config["low_accuracy_layout"]

    font_size = int(text_style.get("font_size", 14))
    color = _hex_to_rgb(text_style.get("color", "000000"))

    left, top, width, height, left_px, top_px = _text_metrics(
        slide_w, slide_h, config
    )
    _ = top_px

    ph_w = _px_to_emu(low_cfg.get("image_placeholder_width_px", 520))
    ph_h = _px_to_emu(low_cfg.get("image_placeholder_height_px", 600))
    ph_w = min(ph_w, int(slide_w * 0.5))
    ph_h = min(ph_h, slide_h - _px_to_emu(80))
    ph_left = slide_w - ph_w - _px_to_emu(left_px)
    ph_top = (slide_h - ph_h) // 2
    _add_placeholder(slide, ph_left, ph_top, ph_w, ph_h, "（在此处插入证明图片）")

    text_w = ph_left - left - _px_to_emu(left_px)
    if text_w <= 0:
        text_w = width
    text_h = height

    title_size = font_size + 6
    name_size = max(10, font_size - 2)
    info_size = max(9, font_size - 4)

    title_gap = float(low_cfg.get("title_space_after_pt", 10))
    name_gap = float(low_cfg.get("name_space_after_pt", 3))
    info_gap = float(low_cfg.get("info_space_after_pt", 3))
    block_gap = float(low_cfg.get("gap_space_after_pt", 14))

    paragraphs: list[dict[str, Any]] = [
        {
            "text": str(titles.get("low_accuracy", "准确率未达标详细")),
            "size": title_size,
            "bold": True,
            "color": color,
            "space_after": title_gap,
        },
    ]

    for idx, item in enumerate(items):
        name = str(item.get("算法小类") or "-")
        acc = item.get("_acc")
        acc_text = f"{acc:.2f}%" if isinstance(acc, (int, float)) else "-"
        p = _fmt_num(item.get("点位数量"))
        c = _fmt_num(item.get("准确数量"))
        b = _fmt_num(item.get("拍照模糊数量"))
        ba = _fmt_num(item.get("审核后拍照模糊数量"))

        paragraphs.append(
            {
                "text": f"{global_offset + idx + 1}. {name}    准确率 {acc_text}",
                "size": name_size,
                "bold": True,
                "color": color,
                "space_after": name_gap,
            }
        )
        paragraphs.append(
            {
                "text": f"     点位 {p} / 准确 {c} / 拍照模糊 {b} / 审核后拍照模糊 {ba}",
                "size": info_size,
                "color": color,
                "space_after": info_gap,
            }
        )
        paragraphs.append(
            {
                "text": "     原因：______________________________________________________",
                "size": info_size,
                "color": color,
                "space_after": block_gap,
            }
        )

    _add_rich_textbox(slide, left, top, text_w, text_h, paragraphs)


# ---------------------------------------------------------------------------
# 辅助
# ---------------------------------------------------------------------------

def _fmt_acc(value: Any) -> str:
    if value is None:
        return "-"
    try:
        return f"{float(value):.2f}%"
    except (TypeError, ValueError):
        return "-"


def _fmt_num(value: Any) -> str:
    if value is None:
        return "-"
    n = to_int(value)
    if n is None:
        return str(value)
    return f"{n:,}"


def _safe_filename(name: str) -> str:
    bad = '<>:"/\\|?*\n\r\t'
    out = "".join("_" if c in bad else c for c in str(name))
    return out.strip().strip(".") or "output"


def _resolve_accuracy_field(config: dict[str, Any], override: str | None) -> str:
    if override:
        return override
    return str(config.get("accuracy_field") or "准确率")


# ---------------------------------------------------------------------------
# 核心生成
# ---------------------------------------------------------------------------

def _generate_core(
    xlsx_path: Path,
    station_name: str,
    output_dir: Path,
    config: dict[str, Any],
    accuracy_field_override: str | None = None,
    source_label: str = "",
) -> dict[str, Any]:
    if not _TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"模板文件缺失：{_TEMPLATE_PATH}（请把 model-aiglrithm.pptx 放到模型目录根）"
        )

    rows = load_rows(str(xlsx_path))
    if not rows:
        raise ValueError(f"xlsx 中没有数据：{xlsx_path}")

    matched = filter_by_station(rows, station_name)
    if not matched:
        raise ValueError(
            f"站点匹配失败：{station_name!r} 在 {xlsx_path.name} 中没有匹配行"
        )

    project_name = ""
    for row in matched:
        pn = row.get("项目名称")
        if pn:
            project_name = str(pn).strip()
            break

    substrings = _silent_substrings(config)
    patrol_rows, silent_rows = _classify(matched, substrings)

    threshold = float(config.get("accuracy_threshold", 95.0))
    field = _resolve_accuracy_field(config, accuracy_field_override)

    patrol_stats = _aggregate(patrol_rows)
    silent_stats = _aggregate(silent_rows)
    low_accuracy = _collect_low_accuracy(matched, field, threshold)

    data = {
        "patrol": patrol_stats,
        "silent": silent_stats,
        "low_accuracy": low_accuracy,
        "threshold": threshold,
        "field": field,
        "source_label": source_label or str(xlsx_path),
    }

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    image_dir = output_dir / "images"
    image_dir.mkdir(parents=True, exist_ok=True)

    footer_text = f"*{station_name}-算法交付数据"
    summary_img = render_table_image(
        matched,
        image_dir / "summary.png",
        footer=footer_text,
        style=config.get("table_style") or {},
    )

    prs = Presentation(str(_TEMPLATE_PATH))
    if not prs.slides:
        raise ValueError("模板 pptx 中没有幻灯片，无法复制为母版页")
    template_index = 0

    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)

    summary_slide = _copy_slide(prs, template_index)
    _render_summary_slide(
        summary_slide,
        slide_w,
        slide_h,
        station_name=station_name,
        project_name=project_name,
        data=data,
        summary_img=summary_img,
        config=config,
    )

    low_per_page = int(config.get("low_accuracy_per_page", 6)) or 6
    for start in range(0, len(low_accuracy), low_per_page):
        page_items = low_accuracy[start : start + low_per_page]
        slide = _copy_slide(prs, template_index)
        _render_low_accuracy_slide(
            slide,
            slide_w,
            slide_h,
            items=page_items,
            global_offset=start,
            threshold=threshold,
            config=config,
        )

    _delete_slide(prs, template_index)

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    safe_station = _safe_filename(station_name)
    output_name = f"{safe_station}_算法结项_{ts}.pptx"
    output_path = output_dir / output_name
    prs.save(str(output_path))

    return {
        "status": "completed",
        "summary": (
            f"已生成 {output_name}；巡视层 {_fmt_acc(patrol_stats['accuracy'])}，"
            f"静默层 {_fmt_acc(silent_stats['accuracy'])}，"
            f"低准确率算法 {len(low_accuracy)} 个"
        ),
        "files": [output_name, "images/summary.png"],
        "folders": ["images"],
        "station": station_name,
        "project_name": project_name,
        "threshold": threshold,
        "accuracy_field": field,
        "source_xlsx": str(xlsx_path),
    }


# ---------------------------------------------------------------------------
# 对外入口
# ---------------------------------------------------------------------------

def generate_from_request(request: dict[str, Any]) -> dict[str, Any]:
    config = load_config()

    parameters = request.get("parameters") or {}
    station_name = (
        parameters.get("station_name") or request.get("station_name") or ""
    ).strip()
    if not station_name:
        raise ValueError("缺少 station_name 参数")

    accuracy_field_override = (
        parameters.get("accuracy_field")
        or request.get("accuracy_field")
        or None
    )

    uploads = request.get("uploads") or {}
    input_path = uploads.get("input_file") or request.get("input_path")

    source_label = ""
    if input_path:
        xlsx_path = Path(input_path)
        if xlsx_path.is_dir():
            candidates = sorted(xlsx_path.glob("*.xlsx"))
            if not candidates:
                raise FileNotFoundError(f"目录内没有 xlsx：{xlsx_path}")
            xlsx_path = candidates[0]
        if not xlsx_path.exists():
            raise FileNotFoundError(f"输入 xlsx 不存在：{xlsx_path}")
        source_label = f"导入文件：{xlsx_path}"
    else:
        cfg_xlsx = _resolve_config_xlsx(config)
        if cfg_xlsx is None:
            raise FileNotFoundError(
                "未上传 xlsx，且 config.yaml 的 xlsxway 中没有可用的 xlsx 文件"
            )
        xlsx_path = cfg_xlsx
        source_label = f"config.xlsxway：{xlsx_path}"

    output_dir = Path(request["output_dir"])

    return _generate_core(
        xlsx_path=xlsx_path,
        station_name=station_name,
        output_dir=output_dir,
        config=config,
        accuracy_field_override=accuracy_field_override,
        source_label=source_label,
    )


def generate_local() -> int:
    config = load_config()

    print("=== 结项 PPT 生成（本地模式）===")

    cfg_xlsx = _resolve_config_xlsx(config)
    if cfg_xlsx is not None:
        print(f"[config.xlsxway] 使用最新修改的文件：{cfg_xlsx}")
    else:
        print("[config.xlsxway] 未找到可用 xlsx；请手动输入路径。")

    raw = input(
        "请输入 xlsx 文件路径（可拖拽；留空则使用 config.yaml 中 xlsxway 的最新文件）："
    ).strip().strip('"').strip("'")

    if raw:
        xlsx_path = Path(raw).expanduser().resolve()
        if not xlsx_path.exists():
            print(f"文件不存在：{xlsx_path}")
            return 1
        source_label = f"手动导入：{xlsx_path}"
    else:
        if cfg_xlsx is None:
            print("未提供路径，且 config.xlsxway 中没有可用文件，退出。")
            return 1
        xlsx_path = cfg_xlsx
        source_label = f"config.xlsxway：{xlsx_path}"

    station_name = input("请输入站点名称（可只写前缀，如 马王风电场）：").strip()
    if not station_name:
        print("未提供站点名称，退出。")
        return 1

    default_field = str(config.get("accuracy_field") or "准确率")
    print(f"可用准确率字段：准确率 / 审核后准确率（直接回车使用默认：{default_field}）")
    field_choice = input("请选择准确率字段：").strip() or default_field
    if field_choice not in {"准确率", "审核后准确率"}:
        print(f"未识别的字段 {field_choice!r}，回退到默认：{default_field}")
        field_choice = default_field

    output_dir = xlsx_path.parent

    try:
        result = _generate_core(
            xlsx_path=xlsx_path,
            station_name=station_name,
            output_dir=output_dir,
            config=config,
            accuracy_field_override=field_choice,
            source_label=source_label,
        )
    except Exception as exc:
        import traceback

        traceback.print_exc()
        print(f"[FAIL] {exc}")
        return 1

    print(f"[OK] {result['summary']}")
    print(f"输出目录：{output_dir}")
    return 0