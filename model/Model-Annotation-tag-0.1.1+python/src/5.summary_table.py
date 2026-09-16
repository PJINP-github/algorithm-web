#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
脚本 5.summary_table.py
读取 Work-txt 下的 txt 结果，生成结构化 Excel 汇总表。
不生成“文件清单”和“原文附录”。
"""

from __future__ import annotations

import re
import sys
import math
from collections import Counter
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit(
        "错误: 缺少 openpyxl，请使用 ./start.sh 5 或 "
        "venv/bin/python3 src/5.summary_table.py 执行"
    ) from exc

BASE_DIR = Path(__file__).resolve().parent

if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from ask_data_utils import (  # noqa: E402
    ASK_DATA_FILE,
    ASK_DATA_T2_FILE,
    CLASS_OUT_FILE,
    DATA_CLASS_FILE,
    MODEL_FILE,
    OUTLOOK_FILE,
    OUTLOOK_SORTED_FILE,
    find_text_mappings,
    excluded_task_ids,
    load_category_mapping_exclude_substrings,
    load_category_mapping_rule_items,
    load_category_mapping_rules,
    load_config,
    parse_ask_lines,
    record_exclude_reason,
    remove_excluded_task_refs,
    resolve_mapping_file,
    resolve_work_dir,
    sort_unique_records,
    text_exclude_reason,
)

WORK_DIR = resolve_work_dir(BASE_DIR)
OUTPUT_FILE = WORK_DIR / "数字-9-模型标注训练清单.xlsx"
CONFIG = load_config()
MAPPING_EXCLUDE_SUBSTRINGS = load_category_mapping_exclude_substrings(
    resolve_mapping_file(WORK_DIR)
)


TASK_REF_RE = re.compile(r"\b(?:autoedge|cvat)@\d+\b", re.IGNORECASE)
OUTLOOK_TASK_RE = re.compile(r"^---Task\s+(\d+)\s+(未训练|已训练)\s+(.+)$")
EXCEL_CELL_LIMIT = 32767

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
FILL_TRAINED = PatternFill(fill_type="solid", fgColor="E2F0D9")
FILL_UNTRAINED = PatternFill(fill_type="solid", fgColor="FCE4D6")
FILL_WARN = PatternFill(fill_type="solid", fgColor="FFF2CC")
FILL_INFO = PatternFill(fill_type="solid", fgColor="DDEBF7")
FILL_FILL_SOURCE = PatternFill(fill_type="solid", fgColor="E4DFEC")
FILL_WHITE = PatternFill(fill_type="solid", fgColor="FFFFFF")
WHITE_FONT = Font(color="FFFFFF", bold=True)
NORMAL_FONT = Font(color="000000")
TASK_DETAIL_FONT = InlineFont(color="000000")
TASK_CATEGORY_FONT = InlineFont(color="548235")
TASK_PROJECT_FONT = InlineFont(color="2F75B5")
TASK_NOTE_FONT = InlineFont(color="BF9000")
TASK_MODEL_HIT_FONT = InlineFont(color="C00000", b=True)
THIN_BLACK = Side(style="thin", color="000000")
NO_BORDER = Border()
CELL_BORDER = Border(
    left=THIN_BLACK,
    right=THIN_BLACK,
    top=THIN_BLACK,
    bottom=THIN_BLACK,
)
CELL_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


def read_text(path: Path) -> str:
    if not path.exists():
        return ""
    return path.read_text(encoding="utf-8", errors="replace")


def read_nonempty_lines(path: Path) -> list[str]:
    return [line.strip() for line in read_text(path).splitlines() if line.strip()]


def save_workbook_if_available(wb: Workbook, output_file: Path) -> Path | None:
    try:
        wb.save(output_file)
        return output_file
    except PermissionError:
        print(f"Warning: {output_file} is in use; not overwritten.")
        return None


def excel_value(value):
    if value is None:
        return ""
    if isinstance(value, CellRichText):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if len(text) > EXCEL_CELL_LIMIT:
        return text[: EXCEL_CELL_LIMIT - 20] + "...[单元格截断]"
    return text


def extract_task_refs(text: str) -> list[str]:
    refs = []
    seen = set()
    for match in TASK_REF_RE.finditer(text or ""):
        ref = match.group(0).lower()
        if ref in seen:
            continue
        seen.add(ref)
        refs.append(ref)
    return refs


def task_ref_sort_key(ref: str) -> tuple[int, str, int]:
    normalized = ref.strip().lower()
    match = TASK_REF_RE.fullmatch(normalized)
    if not match:
        return 99, normalized, 0
    prefix, task_id = normalized.split("@", 1)
    prefix_order = {"autoedge": 0, "cvat": 1}.get(prefix, 10)
    return prefix_order, prefix, -int(task_id)


def sort_task_refs(refs: list[str]) -> list[str]:
    unique = {}
    for ref in refs:
        normalized = ref.strip().lower()
        if normalized and normalized not in unique:
            unique[normalized] = normalized
    return sorted(unique.values(), key=task_ref_sort_key)


def format_refs(refs: list[str]) -> str:
    return ",".join(sort_task_refs(refs))


def task_ids_from_refs(refs: list[str]) -> set[str]:
    ids = set()
    for ref in refs:
        match = TASK_REF_RE.fullmatch(ref.strip().lower())
        if match:
            ids.add(ref.split("@", 1)[1])
    return ids


def count_refs_by_prefix(refs: list[str]) -> tuple[int, int]:
    sorted_refs = sort_task_refs(refs)
    autoedge_count = sum(1 for ref in sorted_refs if ref.startswith("autoedge@"))
    cvat_count = sum(1 for ref in sorted_refs if ref.startswith("cvat@"))
    return autoedge_count, cvat_count


def split_model_line(line: str) -> tuple[str, str]:
    if "  完整Task:" in line:
        key, content = line.split("  完整Task:", 1)
    elif "完整Task:" in line:
        key, content = line.split("完整Task:", 1)
    else:
        key, content = line, ""
    return key.strip(), content.strip()


def split_fill_source_key(key: str) -> tuple[str, str]:
    marker = "  补齐Task来源:"
    if marker not in key:
        return key.strip(), ""
    clean_key, source_text = key.split(marker, 1)
    return clean_key.strip(), source_text.strip()


def extract_latest_task_ids(key: str) -> list[str]:
    match = re.search(r"最新Task\s+(.*?)(?:\s+备注\s+|$)", key)
    if not match:
        return []
    return re.findall(r"\d+", match.group(1))


def parse_model_key(key: str) -> dict:
    clean_key, source_text = split_fill_source_key(key)
    latest_ids = extract_latest_task_ids(clean_key)
    note_part = clean_key
    start_time = ""

    for marker in ("  训练开始时间 ", "  开始时间 "):
        if marker in note_part:
            note_part, start_time = note_part.rsplit(marker, 1)
            start_time = start_time.strip()
            break

    if "  备注 " in note_part:
        _prefix, note = note_part.split("  备注 ", 1)
    else:
        note = note_part.replace("最新Task", "").strip()

    return {
        "latest_ids": latest_ids,
        "note": note.strip(),
        "start_time": start_time,
        "fill_source_text": source_text,
    }


def parse_fill_sources(source_text: str) -> list[dict]:
    sources = []
    for segment in (source_text or "").split(";"):
        segment = segment.strip()
        if not segment or "=" not in segment:
            continue
        label, refs_text = segment.split("=", 1)
        refs = sort_task_refs(extract_task_refs(refs_text))
        if refs:
            sources.append(
                {
                    "label": label.strip(),
                    "refs": refs,
                    "count": len(refs),
                }
            )
    return sources


def format_fill_sources(source_text: str) -> str:
    parts = []
    for source in parse_fill_sources(source_text):
        parts.append(
            f"{source['label']}({source['count']}):"
            f"{','.join(source['refs'])}"
        )
    return "; ".join(parts)


def load_model_entries(excluded_ids: set[str] | None = None) -> list[dict]:
    rows = []
    excluded_ids = {str(task_id) for task_id in (excluded_ids or set())}
    for line in read_nonempty_lines(WORK_DIR / MODEL_FILE):
        if not line.startswith("最新Task"):
            continue
        if text_exclude_reason(line, MAPPING_EXCLUDE_SUBSTRINGS):
            continue
        refs = extract_task_refs(line)
        if refs and all(
            ref.split("@", 1)[1] in excluded_ids
            for ref in refs
        ):
            continue
        line = remove_excluded_task_refs(line, excluded_ids)
        key, content = split_model_line(line)
        meta = parse_model_key(key)
        meta["latest_ids"] = [
            task_id
            for task_id in meta["latest_ids"]
            if task_id not in excluded_ids
        ]
        refs = sort_task_refs(extract_task_refs(content))
        autoedge_count, cvat_count = count_refs_by_prefix(refs)
        fill_sources = parse_fill_sources(meta["fill_source_text"])
        fill_ref_count = sum(source["count"] for source in fill_sources)
        rows.append(
            {
                "latest_ids": meta["latest_ids"],
                "note": meta["note"],
                "start_time": meta["start_time"],
                "fill_source_text": meta["fill_source_text"],
                "fill_source_summary": format_fill_sources(meta["fill_source_text"]),
                "fill_source_count": len(fill_sources),
                "fill_ref_count": fill_ref_count,
                "refs": refs,
                "autoedge_count": autoedge_count,
                "cvat_count": cvat_count,
            }
        )
    return rows


def attach_model_categories(model_entries: list[dict], outlook_entries: list[dict]):
    categories = [
        entry.get("model_category", "")
        for entry in outlook_entries
        if entry.get("match_status") == "正常"
    ]
    for index, entry in enumerate(model_entries):
        entry["model_category"] = categories[index] if index < len(categories) else ""


def parse_outlook_task_line(line: str) -> dict | None:
    match = OUTLOOK_TASK_RE.match(line.strip())
    if not match:
        return None
    task_id, status, rest = match.groups()
    rest = rest.strip()
    category = rest
    project = ""
    if " " in rest:
        category, project = rest.rsplit(None, 1)
    return {
        "task_id": task_id,
        "status": status,
        "category": category.strip(),
        "project": project.strip().rstrip("&"),
        "ref": f"autoedge@{task_id}",
    }


def summarize_counter(counter: Counter, limit: int = 12) -> str:
    if not counter:
        return ""
    parts = []
    for key, count in counter.most_common(limit):
        parts.append(f"{key}({count})")
    if len(counter) > limit:
        parts.append(f"...共{len(counter)}项")
    return "、".join(parts)


def split_blocks(text: str) -> list[str]:
    return [block for block in re.split(r"\n\s*\n", text or "") if block.strip()]


def extract_braced_block_info(lines: list[str]) -> str:
    for line in reversed(lines):
        stripped = line.strip()
        if stripped.startswith("{") and stripped.endswith("}"):
            return stripped[1:-1].strip()
    return ""


def site_generalize_terms() -> list[str]:
    values = ((CONFIG or {}).get("summary") or {}).get(
        "site_generalize_substrings",
        [],
    )
    if isinstance(values, str):
        values = [values]
    return [
        str(value).strip()
        for value in values
        if str(value).strip()
    ]


def normalize_block_site(site: str, generalize_terms: list[str]) -> str:
    for term in generalize_terms:
        if term in site:
            return term
    return site


def clean_block_info_sites(
    value: str,
    exclude_substrings: list[str] | None = None,
) -> str:
    generalize_terms = site_generalize_terms()
    sites = []
    seen = set()
    for site in str(value or "").split("&"):
        site = site.strip()
        if not site:
            continue
        if text_exclude_reason(site, exclude_substrings or MAPPING_EXCLUDE_SUBSTRINGS):
            continue
        site = normalize_block_site(site, generalize_terms)
        if site in seen:
            continue
        seen.add(site)
        sites.append(site)
    return "&".join(sites)


def load_outlook_entries(
    excluded_ids: set[str] | None = None,
) -> tuple[list[dict], list[str]]:
    excluded_ids = {str(task_id) for task_id in (excluded_ids or set())}
    text = read_text(WORK_DIR / OUTLOOK_SORTED_FILE) or read_text(WORK_DIR / OUTLOOK_FILE)
    entries = []
    unmatched_blocks = []
    model_categories = read_nonempty_lines(WORK_DIR / CLASS_OUT_FILE)

    for block_index, block in enumerate(split_blocks(text)):
        lines = [line.rstrip() for line in block.splitlines() if line.strip()]
        if not lines:
            continue
        header = lines[0]
        model_category = (
            model_categories[block_index]
            if block_index < len(model_categories)
            else ""
        )
        if not header.startswith("最新Task"):
            unmatched_blocks.append(block.strip())
            keyword_match = re.search(r"#\s*关键词\s*\[(.*?)\]", header)
            entries.append(
                {
                    "model_category": keyword_match.group(1).strip() if keyword_match else model_category,
                    "latest_ids": [],
                    "note": "无匹配的 model 或 ask 行",
                    "start_time": "",
                    "trained_count": 0,
                    "untrained_count": 0,
                    "total_task_rows": 0,
                    "coverage": None,
                    "trained_tasks": [],
                    "untrained_tasks": [],
                    "trained_refs": [],
                    "untrained_refs": [],
                    "final_refs": [],
                    "block_info": "",
                    "category_summary": "",
                    "project_summary": "",
                    "fill_source_summary": "",
                    "match_status": "无匹配",
                }
            )
            continue

        key, header_content = split_model_line(header)
        meta = parse_model_key(key)
        block_info = extract_braced_block_info(lines)
        tasks = []
        for line in lines[1:]:
            parsed = parse_outlook_task_line(line)
            if parsed:
                if text_exclude_reason(line, MAPPING_EXCLUDE_SUBSTRINGS):
                    continue
                if parsed["task_id"] in excluded_ids:
                    continue
                tasks.append(parsed)

        final_refs = []
        for line in reversed(lines):
            stripped = line.strip()
            if stripped.startswith("#") and stripped.endswith("#"):
                refs = extract_task_refs(stripped)
                if refs:
                    final_refs = sort_task_refs(refs)
                    break
        if not final_refs:
            final_refs = sort_task_refs(extract_task_refs(header_content))
        final_refs = [
            ref
            for ref in final_refs
            if ref.split("@", 1)[1] not in excluded_ids
        ]

        trained_refs = [task["ref"] for task in tasks if task["status"] == "已训练"]
        untrained_refs = [task["ref"] for task in tasks if task["status"] == "未训练"]
        trained_tasks = [task for task in tasks if task["status"] == "已训练"]
        untrained_tasks = [task for task in tasks if task["status"] == "未训练"]
        category_counts = Counter(task["category"] for task in tasks if task["category"])
        project_counts = Counter(task["project"] for task in tasks if task["project"])
        total_task_rows = len(tasks)
        trained_count = len(trained_refs)
        untrained_count = len(untrained_refs)
        coverage = trained_count / total_task_rows if total_task_rows else 0

        entries.append(
            {
                "model_category": model_category,
                "latest_ids": meta["latest_ids"],
                "note": meta["note"],
                "start_time": meta["start_time"],
                "trained_count": trained_count,
                "untrained_count": untrained_count,
                "total_task_rows": total_task_rows,
                "coverage": coverage,
                "trained_tasks": trained_tasks,
                "untrained_tasks": untrained_tasks,
                "trained_refs": sort_task_refs(trained_refs),
                "untrained_refs": sort_task_refs(untrained_refs),
                "final_refs": final_refs,
                "block_info": block_info,
                "category_summary": summarize_counter(category_counts),
                "project_summary": summarize_counter(project_counts),
                "fill_source_summary": format_fill_sources(meta["fill_source_text"]),
                "match_status": "正常",
            }
        )

    return entries, unmatched_blocks


def load_mapping_rules() -> list[tuple[str, str]]:
    return [
        (source, target)
        for _lineno, source, target in load_category_mapping_rules(
            resolve_mapping_file(WORK_DIR)
        )
    ]


def apply_mapping(value: str, rules: list[tuple[str, str]]) -> str:
    matches = find_text_mappings(value, rules, {"mode": "longest_substring"})
    return "、".join(dict.fromkeys(target for _pattern, target in matches))


def load_highlight_terms_by_target() -> dict[str, list[str]]:
    terms_by_target: dict[str, list[str]] = {}
    mapping_path = resolve_mapping_file(WORK_DIR)
    for item in load_category_mapping_rule_items(mapping_path):
        target = str(item.get("model_target") or item.get("target") or "").strip()
        if not target:
            continue
        terms = terms_by_target.setdefault(target, [])
        for value in (
            target,
            item.get("target", ""),
            item.get("model_target", ""),
            item.get("task_source", ""),
            item.get("source", ""),
        ):
            value = str(value or "").strip()
            if value and value not in terms:
                terms.append(value)
    return terms_by_target


def task_detail_parts(task: dict, ask_lookup: dict[str, object]) -> list[str]:
    record = ask_lookup.get(task.get("task_id"))
    category = getattr(record, "category", "") if record else task.get("category", "")
    project = getattr(record, "project", "") if record else task.get("project", "")
    note = getattr(record, "note", "") if record else ""
    parts = [f"Task {task.get('task_id', '')}", category, project, note]
    return [
        str(part).strip()
        for part in parts
        if str(part).strip() and str(part).strip() != "-"
    ]


def task_detail_text(task: dict, ask_lookup: dict[str, object]) -> str:
    return "{" + " | ".join(task_detail_parts(task, ask_lookup)) + "}"


def append_highlighted_text(
    rich_text: CellRichText,
    text: str,
    base_font: InlineFont,
    highlight_terms: list[str],
):
    terms = [
        term
        for term in sorted(set(highlight_terms), key=len, reverse=True)
        if term
    ]
    if not text or not terms:
        rich_text.append(TextBlock(base_font, text))
        return

    index = 0
    while index < len(text):
        hit = None
        for term in terms:
            if text.startswith(term, index):
                hit = term
                break
        if hit:
            rich_text.append(TextBlock(TASK_MODEL_HIT_FONT, hit))
            index += len(hit)
            continue
        next_index = index + 1
        while next_index < len(text):
            if any(text.startswith(term, next_index) for term in terms):
                break
            next_index += 1
        rich_text.append(TextBlock(base_font, text[index:next_index]))
        index = next_index


def task_detail_rich_text(
    tasks: list[dict],
    ask_lookup: dict[str, object],
    model_category: str = "",
    highlight_terms: list[str] | None = None,
    separator: str = "\n",
):
    if not tasks:
        return "无"
    rich_text = CellRichText()
    highlight_terms = list(highlight_terms or [])
    model_category = str(model_category or "").strip()
    if model_category and model_category not in highlight_terms:
        highlight_terms.append(model_category)
    for task_index, task in enumerate(tasks):
        if task_index:
            rich_text.append(TextBlock(TASK_DETAIL_FONT, separator))
        parts = task_detail_parts(task, ask_lookup)
        rich_text.append(TextBlock(TASK_DETAIL_FONT, "{"))
        for part_index, part in enumerate(parts):
            if part_index:
                rich_text.append(TextBlock(TASK_DETAIL_FONT, " | "))
            if part_index == 0:
                font = TASK_DETAIL_FONT
            elif part_index == 1:
                font = TASK_CATEGORY_FONT
            elif part_index == 2 and len(parts) >= 4:
                font = TASK_PROJECT_FONT
            else:
                font = TASK_NOTE_FONT
            if part_index == 0:
                rich_text.append(TextBlock(font, part))
            else:
                append_highlighted_text(rich_text, part, font, highlight_terms)
        rich_text.append(TextBlock(TASK_DETAIL_FONT, "}"))
    return rich_text


def ask_record_detail_text(record) -> str:
    parts = [
        f"Task {record.task_id}",
        record.task_status,
        record.pre_label_status,
        record.category,
        record.project,
        record.note,
    ]
    return " | ".join(str(part) for part in parts if part)


def trained_image_count(entry: dict, ask_lookup: dict[str, object]) -> int:
    total = 0
    seen_task_ids = set()
    for task in entry.get("trained_tasks", []):
        task_id = str(task.get("task_id", ""))
        if not task_id or task_id in seen_task_ids:
            continue
        seen_task_ids.add(task_id)
        record = ask_lookup.get(task_id)
        try:
            total += max(int(getattr(record, "image_count", 0)), 0)
        except (TypeError, ValueError):
            continue
    return total


def build_overview_model_rows(
    outlook_entries: list[dict],
    ask_records,
    highlight_terms_by_target: dict[str, list[str]] | None = None,
) -> list[list]:
    ask_lookup = {record.task_id: record for record in ask_records}
    highlight_terms_by_target = highlight_terms_by_target or {}
    rows = []
    for index, entry in enumerate(outlook_entries, start=1):
        total = entry["total_task_rows"]
        untrained_count = entry["untrained_count"]
        if total:
            completion_rate = f"{entry['coverage']:.1%}"
            all_trained = "是" if untrained_count == 0 else "否"
        else:
            completion_rate = "无可判断"
            all_trained = "无匹配数据"

        untrained_notes = task_detail_rich_text(
            entry.get("untrained_tasks", []),
            ask_lookup,
            model_category=entry.get("model_category", ""),
            highlight_terms=highlight_terms_by_target.get(
                entry.get("model_category", ""),
                [],
            ),
            separator="---",
        )
        trained_images = trained_image_count(entry, ask_lookup)

        rows.append(
            [
                index,
                entry.get("model_category", ""),
                " ".join(entry.get("latest_ids", [])) or "无",
                entry.get("note", ""),
                entry.get("start_time", "") or "无",
                entry["trained_count"],
                trained_images,
                untrained_count,
                untrained_notes,
                completion_rate,
                all_trained,
            ]
        )
    return rows


def build_not_in_training_rows(ask_records, outlook_entries: list[dict]) -> list[list]:
    final_task_ids = set()
    for entry in outlook_entries:
        final_task_ids.update(task_ids_from_refs(entry.get("final_refs", [])))

    rows = []
    for record in ask_records:
        if record.task_id in final_task_ids:
            continue
        rows.append(
            [
                int(record.task_id) if record.task_id.isdigit() else record.task_id,
                record.task_status,
                record.pre_label_status,
                record.category,
                record.project,
                record.note,
                "未进入任何最终Task群",
            ]
        )

    rows.sort(key=lambda row: int(row[0]) if str(row[0]).isdigit() else -1, reverse=True)
    if not rows:
        return [["无", "", "", "", "", "", "全部标注任务均进入最终Task群"]]
    return rows


def build_annotation_rows(ask_records, ask_t2_records) -> list[list]:
    rows = []
    for source_name, records in ((ASK_DATA_FILE, ask_records), (ASK_DATA_T2_FILE, ask_t2_records)):
        for record in records:
            rows.append(
                [
                    source_name,
                    int(record.task_id) if record.task_id.isdigit() else record.task_id,
                    record.task_status,
                    record.pre_label_status,
                    record.category,
                    record.project,
                    record.note,
                    record.task_name,
                ]
            )
    rows.sort(key=lambda row: int(row[1]) if str(row[1]).isdigit() else -1, reverse=True)
    return rows


def build_model_rows(model_entries: list[dict]) -> list[list]:
    rows = []
    for index, entry in enumerate(model_entries, start=1):
        complete_refs = format_refs(entry["refs"])
        rows.append(
            [
                index,
                entry.get("model_category", ""),
                " ".join(entry["latest_ids"]),
                entry["note"],
                entry["start_time"],
                entry["fill_source_summary"],
                entry["fill_source_count"],
                entry["fill_ref_count"],
                entry["autoedge_count"],
                entry["cvat_count"],
                len(entry["refs"]),
                complete_refs,
            ]
        )
    return rows


def build_outlook_rows(
    outlook_entries: list[dict],
    ask_records,
    highlight_terms_by_target: dict[str, list[str]] | None = None,
) -> list[list]:
    rows = []
    for index, entry in enumerate(outlook_entries, start=1):
        untrained_notes = clean_block_info_sites(
            entry.get("block_info", ""),
            MAPPING_EXCLUDE_SUBSTRINGS,
        )
        rows.append(
            [
                index,
                entry.get("model_category", ""),
                " ".join(entry["latest_ids"]),
                entry["note"],
                entry["start_time"],
                entry["trained_count"],
                entry["untrained_count"],
                entry["total_task_rows"],
                f"{entry['coverage']:.1%}" if entry["total_task_rows"] else "无可判断",
                format_refs(entry["trained_refs"]),
                format_refs(entry["untrained_refs"]),
                untrained_notes,
                format_refs(entry["final_refs"]),
                entry["category_summary"],
                entry["project_summary"],
                entry["fill_source_summary"],
            ]
        )
    return rows


def build_mapping_rows(rules: list[tuple[str, str]]) -> list[list]:
    rows = []
    for source, target in rules:
        rows.append(["映射规则", source, target, "原始->标准"])

    for category in read_nonempty_lines(WORK_DIR / DATA_CLASS_FILE):
        mapped = apply_mapping(category, rules)
        rows.append(
            [
                "原始类别清单",
                category,
                mapped or "",
                "已匹配" if mapped else "未匹配",
            ]
        )

    for category in read_nonempty_lines(WORK_DIR / CLASS_OUT_FILE):
        rows.append(["标准类别清单", "", category, "标准类别"])

    return rows


def mapping_wide_headers(group_size: int = 8) -> tuple[list[str], set[int]]:
    headers = []
    buffer_columns = set()
    for group_index in range(group_size):
        if group_index:
            headers.append("")
            buffer_columns.add(len(headers))
        headers.extend(["信息来源", "原始类别", "标准类别", "说明"])
    return headers, buffer_columns


def build_mapping_wide_rows(records: list[list], group_size: int = 8) -> list[list]:
    rows = []
    width = group_size * 4 + (group_size - 1)
    for index in range(0, len(records), group_size):
        row = []
        for group_index in range(group_size):
            if group_index:
                row.append("")
            record_index = index + group_index
            if record_index < len(records):
                row.extend(records[record_index])
            else:
                row.extend([""] * 4)
        row.extend([""] * (width - len(row)))
        rows.append(row)
    return rows


def annotation_fill(headers: list[str], row: list) -> PatternFill | None:
    pre_label = row[headers.index("预标注状态")]
    task_status = row[headers.index("任务状态")]
    if task_status != "已完成":
        return FILL_WARN
    if pre_label == "完成":
        return FILL_TRAINED
    if pre_label == "未执行":
        return FILL_UNTRAINED
    return FILL_WARN


def model_fill(headers: list[str], row: list) -> PatternFill | None:
    fill_source = row[headers.index("补齐来源")]
    return FILL_FILL_SOURCE if fill_source else FILL_TRAINED


def outlook_fill(headers: list[str], row: list) -> PatternFill | None:
    total_count = row[headers.index("总Task行数")]
    try:
        if int(total_count) == 0:
            return FILL_WARN
    except (TypeError, ValueError):
        return FILL_WARN
    untrained_count = row[headers.index("未训练数")]
    try:
        return FILL_UNTRAINED if int(untrained_count) > 0 else FILL_TRAINED
    except (TypeError, ValueError):
        return FILL_WARN


def mapping_fill(headers: list[str], row: list) -> PatternFill | None:
    status = row[headers.index("说明")]
    source = row[headers.index("来源")]
    if "未匹配" in str(status):
        return FILL_UNTRAINED
    if source == "映射规则":
        return FILL_INFO
    if source == "标准类别清单":
        return FILL_TRAINED
    return None


def overview_model_fill(headers: list[str], row: list) -> PatternFill | None:
    status = row[headers.index("是否全部训练完成")]
    if status == "是":
        return FILL_TRAINED
    if status == "否":
        return FILL_UNTRAINED
    return FILL_WARN


def not_in_training_fill(headers: list[str], row: list) -> PatternFill | None:
    return FILL_UNTRAINED if row[0] != "无" else FILL_TRAINED


def apply_table_style(ws, start_row: int, start_col: int, headers: list[str], rows: list[list], fill_func=None):
    for col_offset, header in enumerate(headers):
        cell = ws.cell(start_row, start_col + col_offset, excel_value(header))
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.border = CELL_BORDER
        cell.alignment = CELL_ALIGNMENT

    for row_offset, row in enumerate(rows, start=1):
        values = [excel_value(value) for value in row]
        row_fill = fill_func(headers, values) if fill_func else None
        for col_offset, value in enumerate(values):
            cell = ws.cell(start_row + row_offset, start_col + col_offset, value)
            cell.font = NORMAL_FONT
            cell.border = CELL_BORDER
            cell.alignment = CELL_ALIGNMENT
            if row_fill:
                cell.fill = row_fill


def fit_columns(ws):
    long_headers = {
        "完整Task群",
        "最终Task群",
        "未训练项目信息",
        "已训练Task群",
        "未训练Task群",
        "未训练Task备注信息",
        "标注说明",
        "补齐来源",
        "说明",
        "备注",
        "训练备注",
        "项目分布",
        "类别分布",
        "判断原因",
    }
    for column_cells in ws.columns:
        header = str(column_cells[0].value or "")
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        cap = 100 if header in long_headers else 42
        width = min(max(max_len + 2, 10), cap)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def estimate_wrapped_lines(value, column_width) -> int:
    """Estimate displayed lines using the final Excel column width."""
    text = "" if value is None else str(value)
    if not text:
        return 1
    usable_width = max(float(column_width or 10) - 2, 1)
    line_count = 0
    for line in text.splitlines() or [""]:
        visual_length = sum(2 if ord(char) > 127 else 1 for char in line)
        line_count += max(1, math.ceil(visual_length / usable_width))
    return max(line_count, 1)


def fit_overview_row_heights(ws):
    """Set each overview row to the largest estimated wrapped cell height."""
    max_row_height = 409.5
    for row_index in range(2, ws.max_row + 1):
        line_count = 1
        for column_index in range(1, ws.max_column + 1):
            column_letter = get_column_letter(column_index)
            width = ws.column_dimensions[column_letter].width or 10
            line_count = max(
                line_count,
                estimate_wrapped_lines(
                    ws.cell(row_index, column_index).value,
                    width,
                ),
            )
        height = min(max(30, line_count * 18 + 6), max_row_height)
        ws.row_dimensions[row_index].height = height


def write_overview_sheet(
    wb,
    outlook_entries: list[dict],
    ask_records,
    highlight_terms_by_target: dict[str, list[str]] | None = None,
):
    ws = wb.create_sheet("总览")
    main_headers = [
        "序号",
        "模型类别",
        "最新Task",
        "训练备注",
        "训练时间",
        "已训练Task数量",
        "已训练图片数量",
        "未训练Task数量",
        "未训练Task备注信息",
        "训练完成率",
        "是否全部训练完成",
    ]
    side_headers = [
        "不纳入训练Task",
        "任务状态",
        "预标注状态",
        "类别",
        "项目",
        "标注说明",
        "判断原因",
    ]
    model_rows = build_overview_model_rows(
        outlook_entries,
        ask_records,
        highlight_terms_by_target=highlight_terms_by_target,
    )
    not_in_training_rows = build_not_in_training_rows(ask_records, outlook_entries)

    side_start_col = len(main_headers) + 2
    apply_table_style(
        ws,
        1,
        1,
        main_headers,
        model_rows,
        fill_func=overview_model_fill,
    )
    apply_table_style(
        ws,
        1,
        side_start_col,
        side_headers,
        not_in_training_rows,
        fill_func=not_in_training_fill,
    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(main_headers))}{len(model_rows) + 1}"
    ws.row_dimensions[1].height = 30
    fit_columns(ws)
    fit_overview_row_heights(ws)


def style_sheet(ws, headers: list[str], fill_func=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row_index, row in enumerate(ws.iter_rows(), start=1):
        values = [cell.value for cell in row]
        row_fill = None if row_index == 1 or not fill_func else fill_func(headers, values)
        for cell in row:
            cell.border = CELL_BORDER
            cell.alignment = CELL_ALIGNMENT
            if row_index == 1:
                cell.fill = HEADER_FILL
                cell.font = WHITE_FONT
            else:
                cell.font = NORMAL_FONT
                if row_fill:
                    cell.fill = row_fill

    ws.row_dimensions[1].height = 28
    rich_note_col = (
        headers.index("未训练Task备注信息") + 1
        if "未训练Task备注信息" in headers
        else None
    )
    for row_index in range(2, ws.max_row + 1):
        height = 36
        if rich_note_col:
            note_value = ws.cell(row_index, rich_note_col).value
            text_len = len(str(note_value or ""))
            height = min(max(44, (text_len // 80 + 1) * 18), 220)
        ws.row_dimensions[row_index].height = height

    fit_columns(ws)


def write_sheet(wb, title: str, headers: list[str], rows: list[list], fill_func=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append([excel_value(value) for value in row])
    style_sheet(ws, headers, fill_func=fill_func)
    if title == "训练状态归纳":
        # Freeze the serial number and model category columns while scrolling.
        ws.freeze_panes = "C2"


def mapping_record_fill(record: list) -> PatternFill | None:
    if len(record) < 4:
        return None
    source, _raw, _target, status = record
    if "未匹配" in str(status):
        return FILL_UNTRAINED
    if source == "映射规则":
        return FILL_INFO
    if source == "标准类别清单":
        return FILL_TRAINED
    return None


def write_mapping_sheet(wb, records: list[list], group_size: int = 8):
    ws = wb.create_sheet("类别映射")
    headers, buffer_columns = mapping_wide_headers(group_size=group_size)

    ws.append(headers)
    wide_rows = build_mapping_wide_rows(records, group_size=group_size)
    for row in wide_rows:
        ws.append([excel_value(value) for value in row])

    for cell in ws[1]:
        if cell.column in buffer_columns:
            cell.fill = FILL_WHITE
            cell.font = NORMAL_FONT
            cell.border = NO_BORDER
        else:
            cell.fill = HEADER_FILL
            cell.font = WHITE_FONT
            cell.border = CELL_BORDER
        cell.alignment = CELL_ALIGNMENT

    for row_index, row in enumerate(wide_rows, start=2):
        for group_index in range(group_size):
            buffer_col = group_index * 5 if group_index else None
            if buffer_col:
                cell = ws.cell(row_index, buffer_col)
                cell.fill = FILL_WHITE
                cell.border = NO_BORDER
                cell.alignment = CELL_ALIGNMENT

            start = group_index * 5
            record = row[start:start + 4]
            row_fill = mapping_record_fill(record) if any(record) else None
            for col_offset in range(4):
                cell = ws.cell(row_index, start + col_offset + 1)
                cell.font = NORMAL_FONT
                cell.border = CELL_BORDER
                cell.alignment = CELL_ALIGNMENT
                if row_fill:
                    cell.fill = row_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 28
    for row_index in range(2, ws.max_row + 1):
        ws.row_dimensions[row_index].height = 36
    fit_columns(ws)
    for column_index in buffer_columns:
        column_letter = get_column_letter(column_index)
        ws.column_dimensions[column_letter].width = 3
        for row_index in range(1, ws.max_row + 1):
            cell = ws.cell(row_index, column_index)
            cell.fill = FILL_WHITE
            cell.border = NO_BORDER
            cell.alignment = CELL_ALIGNMENT


def main():
    if not WORK_DIR.exists():
        raise SystemExit(f"错误: 工作目录不存在 {WORK_DIR}")

    ask_records = sort_unique_records(
        record
        for record in parse_ask_lines(read_text(WORK_DIR / ASK_DATA_FILE).splitlines())
        if not record_exclude_reason(record, MAPPING_EXCLUDE_SUBSTRINGS)
    )
    ask_t2_records = sort_unique_records(
        record
        for record in parse_ask_lines(read_text(WORK_DIR / ASK_DATA_T2_FILE).splitlines())
        if not record_exclude_reason(record, MAPPING_EXCLUDE_SUBSTRINGS)
    )
    raw_ask_records = parse_ask_lines(
        read_text(WORK_DIR / ASK_DATA_FILE).splitlines()
    )
    excluded_ids = excluded_task_ids(
        raw_ask_records,
        MAPPING_EXCLUDE_SUBSTRINGS,
    )
    model_entries = load_model_entries(excluded_ids)
    outlook_entries, _unmatched_blocks = load_outlook_entries(excluded_ids)
    attach_model_categories(model_entries, outlook_entries)
    mapping_rules = load_mapping_rules()
    highlight_terms_by_target = load_highlight_terms_by_target()

    wb = Workbook()
    wb.remove(wb.active)

    write_overview_sheet(
        wb,
        outlook_entries,
        ask_records,
        highlight_terms_by_target=highlight_terms_by_target,
    )
    write_sheet(
        wb,
        "项目训练明细",
        [
            "序号",
            "模型",
            "最新Task",
            "备注",
            "训练开始时间",
            "补齐来源",
            "补齐来源数",
            "补齐Task数",
            "autoedge数量",
            "cvat数量",
            "完整Task数量",
            "完整Task群",
        ],
        build_model_rows(model_entries),
        fill_func=model_fill,
    )
    write_sheet(
        wb,
        "训练状态归纳",
        [
            "序号",
            "模型类别",
            "最新Task",
            "备注",
            "训练开始时间",
            "已训练数",
            "未训练数",
            "总Task行数",
            "覆盖率",
            "已训练Task群",
            "未训练Task群",
            "未训练Task备注信息",
            "最终Task群",
            "类别分布",
            "项目分布",
            "补齐来源",
        ],
        build_outlook_rows(
            outlook_entries,
            ask_records,
            highlight_terms_by_target=highlight_terms_by_target,
        ),
        fill_func=outlook_fill,
    )
    write_sheet(
        wb,
        "标注任务明细",
        ["来源", "Task ID", "任务状态", "预标注状态", "类别", "项目", "标注说明", "任务名称"],
        build_annotation_rows(ask_records, ask_t2_records),
        fill_func=annotation_fill,
    )
    write_mapping_sheet(wb, build_mapping_rows(mapping_rules))

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    saved_file = save_workbook_if_available(wb, OUTPUT_FILE)
    if saved_file:
        print(f"Done, output file: {saved_file}")
    else:
        print(f"Done, file is in use and was not overwritten: {OUTPUT_FILE}")
    print("已生成 sheet: " + ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
