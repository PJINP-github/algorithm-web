# -*- coding: utf-8 -*-
"""
周报生成脚本
用法: python weekly_report.py <站点关键字>
示例: python weekly_report.py 马王

通过对比配置目录中的旧数据 Excel 与新数据 Excel，
生成指定站点的 Markdown 周报。所有展示逻辑由 config.yaml 灵活控制。
"""
import sys
import os
import re
import json
import socket
import subprocess
import shutil
import time
from datetime import date, datetime, timedelta
import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ========== 配置 ==========
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.environ.get("RUST_PORTAL_WEEKLY_OUTPUT_DIR", "").strip() or os.path.join(BASE_DIR, "output")
OUTPUT_RETENTION_DAYS = 7
GENERATED_FILE_PREFIX = "周报_"
GENERATED_FILE_EXTENSIONS = (".md", ".png", ".xlsx")
DEFAULT_DATA_DIR = os.path.join(BASE_DIR, "date_review", "0Work")
DEFAULT_OLD_FILENAME = "old.xlsx"
DEFAULT_NEW_FILENAME = "result-2.xlsx"
CONFIG_FILE = os.path.join(BASE_DIR, "config.yaml")
LAN_UPLOAD_RESULT_FILENAME = "lan_temp_result-2.xlsx"
LAN_UPLOAD_RESULT_META_FILENAME = "lan_temp_result-2.json"
LAN_UPLOAD_OLD_FILENAME = "lan_temp_old.xlsx"
LAN_UPLOAD_OLD_META_FILENAME = "lan_temp_old.json"
LAN_UPLOAD_FILENAME = LAN_UPLOAD_RESULT_FILENAME
LAN_UPLOAD_META_FILENAME = LAN_UPLOAD_RESULT_META_FILENAME
LAN_SESSIONS_DIRNAME = "lan_sessions"
LAN_UPLOAD_KINDS = {
    "old": {
        "filename": LAN_UPLOAD_OLD_FILENAME,
        "meta_filename": LAN_UPLOAD_OLD_META_FILENAME,
        "default_name": DEFAULT_OLD_FILENAME,
        "temp_prefix": "lan_temp_old_",
    },
    "result": {
        "filename": LAN_UPLOAD_RESULT_FILENAME,
        "meta_filename": LAN_UPLOAD_RESULT_META_FILENAME,
        "default_name": DEFAULT_NEW_FILENAME,
        "temp_prefix": "lan_temp_result_",
    },
}

# 中文章节编号
CN_NUMS = ["一", "二", "三", "四", "五", "六", "七", "八", "九"]

HEADERS = [
    "业务线", "项目名称", "工程组", "站点", "服务器状态", "算法状态", "达标率", "更新计划",
    "算法小类", "任务日期", "点位数量", "准确数量", "审核后准确数量", "拍照模糊数量",
    "审核后拍照模糊数量", "准确率", "审核后准确率", "误报率", "审核后误报率", "漏报率",
    "审核后漏报率", "当前版本", "计划数据上传日期", "计划训练日期", "下一版本当前阶段",
    "下一版本计划更新日期"
]

PERCENT_FIELDS = {
    "准确率",
    "审核后准确率",
    "误报率",
    "审核后误报率",
}


def load_config():
    """加载 yaml 配置"""
    with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def resolve_config_path(value, base_dir):
    """解析 YAML 路径，支持环境变量、相对路径和绝对路径。"""
    path = os.path.expandvars(os.path.expanduser(str(value or "").strip()))
    if os.path.isabs(path):
        return os.path.normpath(path)
    return os.path.normpath(os.path.join(base_dir, path))


def resolve_data_files(cfg):
    """根据 YAML 配置解析旧数据和新数据文件路径。"""
    source_cfg = cfg.get("data_source", {}) or {}
    portal_data_dir = os.environ.get("RUST_PORTAL_DATA_DIR", "").strip()
    if portal_data_dir:
        source_cfg = dict(source_cfg)
        source_cfg["directory"] = portal_data_dir
    data_dir = resolve_config_path(
        source_cfg.get("directory", DEFAULT_DATA_DIR),
        BASE_DIR,
    )
    old_file = resolve_config_path(
        source_cfg.get("old_file", DEFAULT_OLD_FILENAME),
        data_dir,
    )
    new_file = resolve_config_path(
        source_cfg.get("new_file", DEFAULT_NEW_FILENAME),
        data_dir,
    )
    return old_file, new_file


def normalize_lan_upload_kind(kind):
    normalized = str(kind or "result").strip().lower()
    if normalized in ("result", "new"):
        return "result"
    if normalized in ("old", "previous"):
        return "old"
    raise ValueError("unsupported LAN upload kind: %s" % kind)


def get_lan_upload_paths(output_dir, kind="result"):
    """返回局域网临时上传表和元数据路径。"""
    upload_kind = normalize_lan_upload_kind(kind)
    upload_cfg = LAN_UPLOAD_KINDS[upload_kind]
    return (
        os.path.join(output_dir, upload_cfg["filename"]),
        os.path.join(output_dir, upload_cfg["meta_filename"]),
    )


def validate_xlsx_file(path):
    """确认上传文件是可读取的 Excel 工作簿。"""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            _ = wb.sheetnames
        finally:
            wb.close()
        return True, None
    except Exception as exc:
        return False, str(exc)


def write_lan_upload_metadata(meta_path, original_name, saved_path, kind="result"):
    upload_kind = normalize_lan_upload_kind(kind)
    upload_cfg = LAN_UPLOAD_KINDS[upload_kind]
    metadata = {
        "kind": upload_kind,
        "logicalName": upload_cfg["default_name"],
        "originalName": original_name or upload_cfg["default_name"],
        "savedName": os.path.basename(saved_path),
        "size": os.path.getsize(saved_path),
        "mtime": datetime.now().isoformat(timespec="seconds"),
    }
    with open(meta_path, "w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False, indent=2)


def read_lan_upload_metadata(output_dir, kind="result"):
    upload_kind = normalize_lan_upload_kind(kind)
    upload_cfg = LAN_UPLOAD_KINDS[upload_kind]
    upload_path, meta_path = get_lan_upload_paths(output_dir, upload_kind)
    if not os.path.isfile(upload_path):
        return None
    metadata = {}
    if os.path.isfile(meta_path):
        try:
            with open(meta_path, "r", encoding="utf-8") as f:
                metadata = json.load(f) or {}
        except Exception:
            metadata = {}
    stat = os.stat(upload_path)
    metadata["kind"] = upload_kind
    metadata.setdefault("logicalName", upload_cfg["default_name"])
    metadata.setdefault("originalName", upload_cfg["default_name"])
    metadata["savedName"] = os.path.basename(upload_path)
    metadata["size"] = stat.st_size
    metadata.setdefault(
        "mtime",
        datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
    )
    return metadata


def read_lan_uploads_metadata(output_dir):
    return {
        kind: read_lan_upload_metadata(output_dir, kind)
        for kind in LAN_UPLOAD_KINDS
    }


def resolve_effective_data_files(old_file, new_file, output_dir):
    """根据局域网临时上传状态决定本次使用的新旧数据。"""
    old_upload_path, _ = get_lan_upload_paths(output_dir, "old")
    result_upload_path, _ = get_lan_upload_paths(output_dir, "result")
    upload_state = {
        "old": os.path.isfile(old_upload_path),
        "result": os.path.isfile(result_upload_path),
    }
    effective_old_file = old_file
    effective_new_file = new_file
    if upload_state["old"]:
        effective_old_file = old_upload_path
    elif upload_state["result"]:
        effective_old_file = new_file
    if upload_state["result"]:
        effective_new_file = result_upload_path
    return effective_old_file, effective_new_file, upload_state


def ensure_output_dir():
    """确保周报产物输出目录存在。"""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    return OUTPUT_DIR


def ensure_lan_session_dir(output_dir, session_dir=None, session_id=None):
    """为局域网 Web 模式准备本次启动专属临时目录。"""
    if session_dir:
        path = os.path.abspath(session_dir)
    else:
        session_id = session_id or datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        safe_id = re.sub(r"[^0-9A-Za-z_.-]+", "_", str(session_id)).strip("._")
        if not safe_id:
            safe_id = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        path = os.path.join(output_dir, LAN_SESSIONS_DIRNAME, safe_id)
    os.makedirs(path, exist_ok=True)
    return path


def get_session_id_from_dir(session_dir):
    return os.path.basename(os.path.abspath(session_dir))


def parse_report_date_from_filename(filename):
    """从周报产物文件名中解析 YYYYMMDD 日期。"""
    match = re.match(
        r"^%s.+_(\d{8})(?:_|\.)(?:.*)$" % re.escape(GENERATED_FILE_PREFIX),
        filename,
    )
    if not match:
        return None
    try:
        return datetime.strptime(match.group(1), "%Y%m%d").date()
    except ValueError:
        return None


def cleanup_old_outputs(output_dir, retention_days=OUTPUT_RETENTION_DAYS):
    """删除 output 中超过保留天数的周报产物。"""
    deleted = []
    failures = []
    if not os.path.isdir(output_dir):
        return deleted, failures

    cutoff_date = date.today() - timedelta(days=retention_days)
    for entry in os.scandir(output_dir):
        if not entry.is_file():
            continue
        filename = entry.name
        if not filename.startswith(GENERATED_FILE_PREFIX):
            continue
        if not filename.lower().endswith(GENERATED_FILE_EXTENSIONS):
            continue

        report_date = parse_report_date_from_filename(filename)
        try:
            file_date = (
                report_date
                if report_date is not None
                else datetime.fromtimestamp(entry.stat().st_mtime).date()
            )
            if file_date < cutoff_date:
                os.remove(entry.path)
                deleted.append(entry.path)
        except OSError as exc:
            failures.append((entry.path, exc))

    return deleted, failures


def load_site_rows(fname, keyword):
    """加载指定站点(子串匹配)的所有数据行"""
    wb = openpyxl.load_workbook(fname, data_only=True)
    try:
        ws = wb.active
        actual_headers = []
        for cell in ws[1]:
            header = str(cell.value).strip() if cell.value is not None else ""
            actual_headers.append(header)
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            data = dict(zip(HEADERS, row))
            for header, value in zip(actual_headers, row):
                if header:
                    data[header] = value
            site = str(data.get("站点") or "")
            if keyword and keyword in site:
                # 源表中的 '-' 表示该百分比为 0，读取时标准化，避免影响平均值和阈值判断。
                for field in PERCENT_FIELDS:
                    if str(data.get(field) or "").strip() == "-":
                        data[field] = "0%"
                rows.append(data)
        return rows
    finally:
        wb.close()


def parse_percent(v):
    """解析百分比为 float，无效返回 None
    支持格式: '96%', '96', 0.96(小数百分比=96%), - 等
    """
    if v is None:
        return None
    if isinstance(v, (int, float)):
        # 数值类型: <=1 视为小数百分比(0.96=96%), >1 直接视为百分比(96=96%)
        val = float(v)
        if val <= 1:
            return val * 100
        return val
    s = str(v).strip()
    if s in ("", "-", "--"):
        return None
    s = s.replace("%", "").strip()
    try:
        val = float(s)
        # Excel 文本单元格也可能存成 '0.86'，与数值型 0.86 的含义保持一致。
        return val * 100 if val <= 1 else val
    except ValueError:
        return None


def parse_int(v):
    """解析整数，支持 Excel 中的千分位分隔符，无效返回 0。"""
    if v is None:
        return 0
    s = str(v).strip()
    if s in ("", "-", "--"):
        return 0
    s = s.replace(",", "").replace("，", "")
    try:
        return int(float(s))
    except ValueError:
        return 0


def is_blank_or_dash(v):
    """判断原始单元格是否为空、短横线或双短横线。"""
    return v is None or str(v).strip() in ("", "-", "--")


def has_effective_points(v):
    """判断点位数量是否是可进入统计和明细表的有效正数。"""
    return not is_blank_or_dash(v) and parse_int(v) > 0


def has_accuracy_count_value(v):
    """判断准确数量字段是否有值；'-'和'无数据'由parse_int按0计入。"""
    return v is not None and str(v).strip() != ""


def has_zero_points(v):
    """判断点位数量是否明确为 0。"""
    return not is_blank_or_dash(v) and parse_int(v) == 0


def parse_task_date(v):
    """解析 'M-D' 格式任务日期为 date 对象，无效返回 None"""
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "-", "--"):
        return None
    try:
        month, day = s.split("-")
        month, day = int(month), int(day)
        today = date.today()
        year = today.year
        candidate = date(year, month, day)
        if candidate > today:
            candidate = date(year - 1, month, day)
        return candidate
    except (ValueError, TypeError):
        return None


def fmt_percent(v):
    """格式化百分比显示"""
    if v is None:
        return "-"
    return "%.0f%%" % v


def fmt_date(d):
    """格式化日期显示"""
    if d is None:
        return "-"
    return d.strftime("%Y-%m-%d")


def fmt_table_date(v):
    """格式化表格中的日期单元格，兼容 Excel 日期和文本日期。"""
    if v is None or is_blank_or_dash(v):
        return "-"
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def get_algorithm_build_date(row):
    """获取算法无数据风险表中的建立日期，旧表头缺失时回退到任务日期。"""
    for field in ("建立日期", "任务日期"):
        value = row.get(field)
        if value is not None and not is_blank_or_dash(value):
            return value
    return None


def clean_site_name(name):
    """去掉站点名中的日期部分，如 '马王风电场 2024-09-30' → '马王风电场'"""
    if not name:
        return name
    return re.sub(r"\s+\d{4}-\d{2}-\d{2}$", "", str(name)).strip()


def get_silent_substrings(cfg):
    """读取静默算法匹配子字符串，未配置时兼容旧逻辑。"""
    configured = cfg.get("Specify_silent_mode_group", None)
    if configured is None:
        return ["静默"]
    if isinstance(configured, dict):
        configured = configured.get("substrings", configured.get("substring", []))
    if isinstance(configured, str):
        configured = [configured]
    if not isinstance(configured, (list, tuple, set)):
        return ["静默"]
    return [
        str(item.get("substring", "") if isinstance(item, dict) else item).strip()
        for item in configured
        if str(item.get("substring", "") if isinstance(item, dict) else item).strip()
    ]


def is_silent_algo(name, silent_substrings=None):
    """算法小类名称包含任一配置子字符串时视为静默算法。"""
    if silent_substrings is None:
        silent_substrings = ["静默"]
    algorithm_name = str(name or "")
    return any(substring in algorithm_name for substring in silent_substrings)


def get_position_accuracy_selection(cfg):
    """返回点位准确数量字段，1=准确数量，2=审核后准确数量。"""
    try:
        selection = int(cfg.get("AlgorithmPositionAccuracySelection", 1))
    except (TypeError, ValueError):
        selection = 1
    return "审核后准确数量" if selection == 2 else "准确数量"


def calculate_position_accuracy(rows, correct_field):
    """按有效点位数量汇总点位准确率，缺失准确数量按0计入。"""
    valid_rows = [
        row for row in rows
        if has_effective_points(row.get("点位数量"))
        and has_accuracy_count_value(row.get(correct_field))
    ]
    total_points = sum(parse_int(row.get("点位数量")) for row in valid_rows)
    total_correct = sum(parse_int(row.get(correct_field)) for row in valid_rows)
    if total_points <= 0:
        return None
    return total_correct / total_points * 100


def average_percent(rows, field):
    """计算指定百分比字段的算法平均值，结果单位为百分数。"""
    values = [parse_percent(row.get(field)) for row in rows]
    values = [value for value in values if value is not None]
    return sum(values) / len(values) if values else None


def get_zero_accuracy_algorithms(rows, accuracy_field, silent_substrings=None):
    """获取有效点位且当前选定准确率为 0% 的算法小类。"""
    zero_accuracy_algos = []
    for row in rows:
        if not has_effective_points(row.get("点位数量")):
            continue
        if parse_percent(row.get(accuracy_field)) == 0:
            zero_accuracy_algos.append(row["算法小类"])
    return zero_accuracy_algos


def format_accuracy_detail(row, display_mode, silent_substrings=None):
    """格式化算法明细中的准确率和数量信息。"""
    display_mode = str(display_mode or "audited").lower()
    raw_acc = parse_percent(row.get("准确率"))
    audited_acc = parse_percent(row.get("审核后准确率"))
    accuracy_parts = []

    if display_mode in ("raw", "both") and raw_acc is not None:
        accuracy_parts.append("准确率 %s" % fmt_percent(raw_acc))
    if display_mode in ("audited", "both") and audited_acc is not None:
        accuracy_parts.append("审核后 %s" % fmt_percent(audited_acc))
    if not accuracy_parts:
        accuracy_parts.append("审核后 -")

    points = parse_int(row.get("点位数量"))
    accurate = parse_int(row.get("准确数量"))
    errors = max(0, points - accurate)
    accuracy_parts.append("点位：%d" % points)
    accuracy_parts.append("准确：%d" % accurate)

    if is_silent_algo(row.get("算法小类"), silent_substrings):
        accuracy_parts.append("误报：%d" % errors)
    elif errors > 0:
        accuracy_parts.append("错误：%d" % errors)

    return "，".join(accuracy_parts)


def get_accuracy_sort_value(row, display_mode, accuracy_source):
    """获取准确率表的排序值。"""
    display_mode = str(display_mode or "audited").lower()
    accuracy_source = str(accuracy_source or "audited").lower()
    raw_value = parse_percent(row.get("准确率"))
    audited_value = parse_percent(row.get("审核后准确率"))

    if display_mode == "raw":
        return raw_value if raw_value is not None else audited_value
    if display_mode == "audited":
        return audited_value if audited_value is not None else raw_value

    if accuracy_source == "raw":
        return raw_value if raw_value is not None else audited_value
    return audited_value if audited_value is not None else raw_value


def get_accuracy_cell_style(value, style_cfg):
    """根据准确率值返回准确率单元格文字样式。"""
    if value is None:
        return None
    style_cfg = style_cfg or {}
    warning_threshold = parse_percent(style_cfg.get("warning_threshold", 95))
    danger_threshold = parse_percent(style_cfg.get("danger_threshold", 50))
    warning_text = normalize_hex_color(style_cfg.get("warning_text"), "C88700")
    danger_text = normalize_hex_color(style_cfg.get("danger_text"), "E57373")
    if danger_threshold is not None and value < danger_threshold:
        return {"text_color": danger_text}
    if warning_threshold is not None and value < warning_threshold:
        return {"text_color": warning_text}
    return None


def build_accuracy_table_data(
    rows, display_mode, overview_cfg, accuracy_source, silent_substrings=None
):
    """生成整体准确率表格数据。"""
    display_mode = str(display_mode or "audited").lower()
    overview_cfg = overview_cfg or {}
    columns_cfg = overview_cfg.get("accuracy_detail_columns", {}) or {}
    sort_mode = str(overview_cfg.get("accuracy_detail_sort", "desc")).strip().lower()
    highlight_cfg = overview_cfg.get("accuracy_detail_colors", {}) or {}

    show_points = bool(columns_cfg.get("show_points", True))
    show_correct = bool(columns_cfg.get("show_correct", True))
    show_errors = bool(columns_cfg.get("show_errors", True))
    show_accuracy = bool(columns_cfg.get("show_accuracy", True))

    headers = ["算法小类"]
    if show_points:
        headers.append("点位数量")
    if show_correct:
        headers.append("准确数量")
    if show_errors:
        headers.append("误报/错误数量")
    if show_accuracy:
        if display_mode in ("raw", "both"):
            headers.append("准确率")
        if display_mode in ("audited", "both"):
            headers.append("审核后准确率")

    records = []
    for row in rows:
        points = parse_int(row.get("点位数量"))
        accurate = parse_int(row.get("准确数量"))
        errors = max(0, points - accurate)
        accuracy_value = get_accuracy_sort_value(row, display_mode, accuracy_source)
        values = [str(row.get("算法小类") or "")]
        cell_styles = [None]
        if show_points:
            values.append(points)
            cell_styles.append(None)
        if show_correct:
            values.append(accurate)
            cell_styles.append(None)
        if show_errors:
            values.append(errors)
            cell_styles.append(None)
        if show_accuracy:
            if display_mode in ("raw", "both"):
                raw_accuracy = parse_percent(row.get("准确率"))
                values.append(fmt_percent(raw_accuracy))
                cell_styles.append(get_accuracy_cell_style(raw_accuracy, highlight_cfg))
            if display_mode in ("audited", "both"):
                audited_accuracy = parse_percent(row.get("审核后准确率"))
                values.append(fmt_percent(audited_accuracy))
                cell_styles.append(get_accuracy_cell_style(audited_accuracy, highlight_cfg))
        records.append({
            "values": values,
            "name": values[0],
            "sort_value": accuracy_value,
            "cell_styles": cell_styles,
        })

    if sort_mode in ("asc", "ascending", "low_first"):
        records.sort(key=lambda item: (
            item["sort_value"] is None,
            item["sort_value"] if item["sort_value"] is not None else 0,
            item["name"],
        ))
    elif sort_mode not in ("none", "off", "false", "no"):
        records.sort(key=lambda item: (
            item["sort_value"] is None,
            -(item["sort_value"] if item["sort_value"] is not None else 0),
            item["name"],
        ))

    table_rows = [item["values"] for item in records]
    cell_styles = [item["cell_styles"] for item in records]
    return headers, table_rows, cell_styles


def build_accuracy_markdown_table(headers, table_rows):
    """生成 Markdown 备用表格。"""
    def escape(value):
        return str(value).replace("|", "\\|").replace("\n", " ")

    lines = [
        "| %s |" % " | ".join(headers),
        "| %s |" % " | ".join("---:" if index else "---" for index in range(len(headers))),
    ]
    for row in table_rows:
        lines.append("| %s |" % " | ".join(escape(value) for value in row))
    return lines


def create_accuracy_xlsx(path, headers, table_rows, sheet_title="整体准确率明细", row_styles=None, cell_styles=None):
    """生成带边框和列宽的表格 Excel。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for row in table_rows:
        ws.append(row)

    thin_black = Side(style="thin", color="000000")
    border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    white_fill = PatternFill("solid", fgColor="FFFFFF")

    row_styles = normalize_style_rows(row_styles, len(table_rows))
    cell_styles = normalize_style_matrix(cell_styles, len(table_rows), len(headers))

    def excel_color(value, default):
        color = normalize_hex_color(value, default)
        if color is None:
            return default
        return color

    for cell in ws[1]:
        cell.font = Font(name="Microsoft YaHei", bold=False, color="000000", size=10)
        cell.fill = white_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[1].height = 22

    for row in ws.iter_rows(min_row=2):
        row_index = row[0].row - 2
        row_style = row_styles[row_index] if row_index < len(row_styles) else None
        for index, cell in enumerate(row):
            cell_style = None
            if row_index < len(cell_styles) and index < len(cell_styles[row_index]):
                cell_style = cell_styles[row_index][index]
            fill_color = resolve_style_color(cell_style, "fill", None)
            if fill_color is None:
                fill_color = resolve_style_color(row_style, "fill", "FFFFFF")
            text_color = resolve_style_color(cell_style, "text_color", None)
            if text_color is None:
                text_color = resolve_style_color(row_style, "text_color", "000000")
            fill_color = excel_color(fill_color, "FFFFFF")
            text_color = excel_color(text_color, "000000")
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)
            else:
                cell.fill = white_fill
            cell.font = Font(name="Microsoft YaHei", color=text_color or "000000", size=10)
            cell.alignment = Alignment(
                horizontal="left" if index == 0 else "center",
                vertical="center",
            )
            cell.border = border
        ws.row_dimensions[row[0].row].height = 20

    widths = {}
    for index, header in enumerate(headers):
        values = [header] + [row[index] for row in table_rows]
        max_len = max(len(str(value)) for value in values)
        header_text = str(header)
        if index == 0:
            width = min(38, max(16, max_len + 3))
        elif "本周准确率" in header_text or "上周准确率" in header_text:
            width = min(34, max(24, max_len + 8))
        elif "准确率" in header_text or "变化" in header_text:
            width = min(22, max(12, max_len + 3))
        elif "数量" in str(header):
            width = min(18, max(12, max_len + 2))
        else:
            width = min(18, max(12, max_len + 2))
        widths[openpyxl.utils.get_column_letter(index + 1)] = width
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    wb.save(path)


def find_table_font(bold=False):
    """查找 Windows 中文字体，供 PNG 表格渲染使用。"""
    fonts_dir = os.path.join(os.environ.get("WINDIR") or BASE_DIR, "Fonts")
    names = (
        ["msyhbd.ttc", "simhei.ttf", "simsunb.ttf", "Dengb.ttf"]
        if bold else
        ["msyh.ttc", "simsun.ttc", "Deng.ttf", "arial.ttf"]
    )
    for name in names:
        path = os.path.join(fonts_dir, name)
        if os.path.exists(path):
            return path
    return None


def create_table_png(
    path,
    headers,
    table_rows,
    split_long_tables=False,
    split_ratio=0.5,
    column_gap=0,
    row_styles=None,
    cell_styles=None,
):
    """将表格渲染为 PNG，长表可拆成左右紧贴的两块。"""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return False

    if not headers or not table_rows:
        return False

    row_styles = normalize_style_rows(row_styles, len(table_rows))
    cell_styles = normalize_style_matrix(cell_styles, len(table_rows), len(headers))

    regular_path = find_table_font()
    bold_path = find_table_font(bold=True) or regular_path
    try:
        font = ImageFont.truetype(regular_path, 18) if regular_path else ImageFont.load_default()
        header_font = ImageFont.truetype(regular_path or bold_path, 18) if (regular_path or bold_path) else font
    except OSError:
        font = ImageFont.load_default()
        header_font = font

    canvas = Image.new("RGB", (10, 10), "white")
    draw = ImageDraw.Draw(canvas)

    def text_width(text, current_font):
        box = draw.textbbox((0, 0), str(text), font=current_font)
        return box[2] - box[0]

    def wrap_text(text, current_font, max_width):
        text = str(text)
        if not text:
            return [""]
        lines = []
        current = ""
        for char in text:
            candidate = current + char
            if current and text_width(candidate, current_font) > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines or [""]

    padding_x = 12
    padding_y = 7
    line_height = 24
    natural_widths = []
    for index, header in enumerate(headers):
        header_text = str(header)
        values = [header] + [row[index] for row in table_rows]
        natural = max(text_width(value, header_font if value == header else font) for value in values)
        if index == 0:
            natural_widths.append(min(360, max(190, natural + padding_x * 2)))
        elif "本周准确率" in header_text or "上周准确率" in header_text:
            natural_widths.append(min(320, max(200, natural + padding_x * 2)))
        elif index == 3 and len(headers) == 4 and "/" in header_text:
            natural_widths.append(min(220, max(130, natural + padding_x * 2)))
        else:
            natural_widths.append(min(180, max(120, natural + padding_x * 2)))

    header_height = line_height + padding_y * 2
    def prepare_rows(rows):
        prepared = []
        for row in rows:
            prepared.append([
                wrap_text(row[index], font, natural_widths[index] - padding_x * 2)
                for index in range(len(headers))
            ])
        return prepared

    def table_size(wrapped_rows):
        row_heights = [
            max(
                line_height + padding_y * 2,
                max(len(lines) for lines in row) * line_height + padding_y * 2,
            )
            for row in wrapped_rows
        ]
        return (
            sum(natural_widths) + 2,
            header_height + sum(row_heights) + 2,
            row_heights,
        )

    wrapped_rows = prepare_rows(table_rows)
    single_width, single_height, single_row_heights = table_size(wrapped_rows)
    try:
        split_ratio = max(0.0, float(split_ratio))
    except (TypeError, ValueError):
        split_ratio = 0.5
    try:
        column_gap = max(0, int(column_gap))
    except (TypeError, ValueError):
        column_gap = 0

    should_split = (
        bool(split_long_tables)
        and len(table_rows) > 1
        and single_height > single_width * split_ratio
    )
    if should_split:
        left_count = (len(table_rows) + 1) // 2
        left_rows = prepare_rows(table_rows[:left_count])
        right_rows = prepare_rows(table_rows[left_count:])
        left_width, left_height, left_row_heights = table_size(left_rows)
        right_width, right_height, right_row_heights = table_size(right_rows)
        left_row_styles = row_styles[:left_count]
        right_row_styles = row_styles[left_count:]
        left_cell_styles = cell_styles[:left_count]
        right_cell_styles = cell_styles[left_count:]
        canvas_width = left_width + column_gap + right_width
        canvas_height = max(left_height, right_height)
    else:
        left_rows = wrapped_rows
        right_rows = []
        left_width = single_width
        left_height = single_height
        left_row_heights = single_row_heights
        right_width = 0
        right_height = 0
        right_row_heights = []
        left_row_styles = row_styles
        right_row_styles = []
        left_cell_styles = cell_styles
        right_cell_styles = []
        canvas_width = single_width
        canvas_height = single_height

    image = Image.new("RGB", (canvas_width, canvas_height), "white")
    draw = ImageDraw.Draw(image)

    def draw_cell(
        left, top, width, height, value_lines, current_font, fill, centered,
        text_color="#1F2933"
    ):
        draw.rectangle(
            [left, top, left + width - 1, top + height - 1],
            fill=fill,
            outline="#000000",
            width=1,
        )
        text_height = len(value_lines) * line_height
        text_top = top + (height - text_height) // 2 + 2
        for line_index, line in enumerate(value_lines):
            line_width = text_width(line, current_font)
            x = left + (width - line_width) // 2 if centered else left + padding_x
            y = text_top + line_index * line_height
            draw.text((x, y), str(line), font=current_font, fill=text_color)

    def draw_table(offset_x, wrapped, row_heights, row_style_items, cell_style_items):
        x = offset_x + 1
        for index, header in enumerate(headers):
            draw_cell(
                x, 1, natural_widths[index], header_height, [header],
                header_font, "#FFFFFF", index != 0, "#000000"
            )
            x += natural_widths[index]

        y = 1 + header_height
        for row_index, row in enumerate(wrapped):
            x = offset_x + 1
            height = row_heights[row_index]
            row_style = row_style_items[row_index] if row_index < len(row_style_items) else None
            for index, value_lines in enumerate(row):
                cell_style = None
                if row_index < len(cell_style_items) and index < len(cell_style_items[row_index]):
                    cell_style = cell_style_items[row_index][index]
                fill = resolve_style_color(cell_style, "fill", None)
                if fill is None:
                    fill = resolve_style_color(row_style, "fill", "#FFFFFF")
                fill = color_for_pillow(fill, "#FFFFFF")
                text_color = resolve_style_color(cell_style, "text_color", None)
                if text_color is None:
                    text_color = resolve_style_color(row_style, "text_color", "#1F2933")
                text_color = color_for_pillow(text_color, "#1F2933")
                draw_cell(
                    x, y, natural_widths[index], height, value_lines, font,
                    fill, index != 0, text_color
                )
                x += natural_widths[index]
            y += height

    draw_table(0, left_rows, left_row_heights, left_row_styles, left_cell_styles)
    if right_rows:
        draw_table(
            left_width + column_gap,
            right_rows,
            right_row_heights,
            right_row_styles,
            right_cell_styles,
        )

    image.save(path, "PNG")
    return True


def create_accuracy_png(path, headers, table_rows, row_styles=None, cell_styles=None):
    """将整体准确率表格渲染为带边框的 PNG。"""
    return create_table_png(
        path,
        headers,
        table_rows,
        row_styles=row_styles,
        cell_styles=cell_styles,
    )


def create_sublist_table_asset(
    base_path,
    suffix,
    headers,
    table_rows,
    table_cfg,
    compact_columns=True,
):
    """根据配置生成嵌套列表对应的 PNG 表格。"""
    if not table_cfg.get("enabled", True) or not table_rows:
        return None

    if compact_columns:
        headers, table_rows = compact_table_columns(headers, table_rows)
    path = "%s_%s.png" % (base_path, suffix)
    has_png = create_table_png(
        path,
        headers,
        table_rows,
        split_long_tables=table_cfg.get("split_long_tables", True),
        split_ratio=table_cfg.get("split_ratio", 0.5),
        column_gap=table_cfg.get("column_gap", 0),
    )
    return path if has_png else None


def append_table_reference(lines, path, alt):
    """向 Markdown 报告插入表格图片引用。"""
    lines.append("")
    lines.append("![%s](%s)" % (alt, os.path.basename(path)))
    lines.append("")


def create_accuracy_table_assets(
    base_path,
    rows,
    display_mode,
    overview_cfg,
    accuracy_source,
    table_cfg=None,
    silent_substrings=None,
):
    """生成整体准确率的 xlsx、PNG 表格资产。"""
    table_cfg = table_cfg or {}
    headers, table_rows, cell_styles = build_accuracy_table_data(
        rows,
        display_mode,
        overview_cfg,
        accuracy_source,
        silent_substrings,
    )
    headers, table_rows = compact_table_columns(headers, table_rows)
    if not table_rows:
        return None
    xlsx_path = base_path + "_准确率明细.xlsx"
    png_path = base_path + "_准确率明细.png"
    create_accuracy_xlsx(
        xlsx_path,
        headers,
        table_rows,
        sheet_title="整体准确率明细",
        cell_styles=cell_styles,
    )
    has_png = create_table_png(
        png_path,
        headers,
        table_rows,
        split_long_tables=table_cfg.get("split_long_tables", True),
        split_ratio=table_cfg.get("split_ratio", 0.5),
        column_gap=table_cfg.get("column_gap", 0),
        cell_styles=cell_styles,
    )
    return {
        "xlsx_path": xlsx_path,
        "png_path": png_path if has_png else None,
        "markdown_lines": build_accuracy_markdown_table(headers, table_rows),
    }


def build_accuracy_trend_table_data(old_map, new_map, accuracy_field, today, changes_cfg):
    """生成本周变化中的准确率趋势表数据。"""
    changes_cfg = changes_cfg or {}
    week_labels = get_week_range_labels(today)
    show_point_change = bool(changes_cfg.get("show_point_change", False))
    sort_mode = str(changes_cfg.get("metric_sort", "desc")).strip().lower()
    color_cfg = changes_cfg.get("accuracy_trend_colors", {}) or {}

    positive_text = normalize_hex_color(color_cfg.get("positive_text"), "2E7D32")
    negative_text = normalize_hex_color(color_cfg.get("negative_text"), "C62828")
    neutral_text = normalize_hex_color(color_cfg.get("neutral_text"), "666666")

    headers = [
        "算法名称",
        "本周准确率(%s)" % week_labels["current_label"],
        "上周准确率(%s)" % week_labels["previous_label"],
        "准确率变化",
    ]
    if show_point_change:
        headers.extend(["本周点位数量", "上周点位数量", "点位变化"])

    records = []
    for name, new_row in new_map.items():
        old_row = old_map.get(name)
        if not old_row:
            continue

        current_accuracy = parse_percent(new_row.get(accuracy_field))
        previous_accuracy = parse_percent(old_row.get(accuracy_field))
        current_points = parse_int(new_row.get("点位数量"))
        previous_points = parse_int(old_row.get("点位数量"))
        delta_accuracy = (
            None if current_accuracy is None or previous_accuracy is None
            else current_accuracy - previous_accuracy
        )
        delta_points = current_points - previous_points

        visible_change = current_accuracy != previous_accuracy
        if show_point_change and current_points != previous_points:
            visible_change = True
        if not visible_change:
            continue

        values = [
            name,
            fmt_percent(current_accuracy),
            fmt_percent(previous_accuracy),
            format_signed_percent(delta_accuracy),
        ]
        cell_styles = [None] * len(values)

        if delta_accuracy is not None:
            if delta_accuracy > 0:
                cell_styles[3] = {"text_color": positive_text}
            elif delta_accuracy < 0:
                cell_styles[3] = {"text_color": negative_text}
            else:
                cell_styles[3] = {"text_color": neutral_text}

        if show_point_change:
            values.extend([
                current_points,
                previous_points,
                format_signed_int(delta_points),
            ])
            cell_styles.extend([None, None, None])

        sort_value = current_accuracy
        records.append({
            "values": values,
            "cell_styles": cell_styles,
            "sort_value": sort_value,
            "name": name,
        })

    if sort_mode in ("asc", "ascending", "low_first"):
        records.sort(key=lambda item: (
            item["sort_value"] is None,
            item["sort_value"] if item["sort_value"] is not None else 0,
            item["name"],
        ))
    elif sort_mode not in ("none", "off", "false", "no"):
        records.sort(key=lambda item: (
            item["sort_value"] is None,
            -(item["sort_value"] if item["sort_value"] is not None else 0),
            item["name"],
        ))

    table_rows = [item["values"] for item in records]
    cell_styles = [item["cell_styles"] for item in records]
    return headers, table_rows, cell_styles


def create_accuracy_trend_table_assets(base_path, old_map, new_map, accuracy_field, today, changes_cfg, table_cfg=None):
    """生成本周变化里的准确率趋势表资产。"""
    table_cfg = table_cfg or {}
    headers, table_rows, cell_styles = build_accuracy_trend_table_data(
        old_map,
        new_map,
        accuracy_field,
        today,
        changes_cfg,
    )
    if not table_rows:
        return None

    xlsx_path = base_path + "_算法准确率变化.xlsx"
    png_path = base_path + "_算法准确率变化.png"
    create_accuracy_xlsx(
        xlsx_path,
        headers,
        table_rows,
        sheet_title="算法准确率变化",
        cell_styles=cell_styles,
    )
    has_png = create_table_png(
        png_path,
        headers,
        table_rows,
        split_long_tables=table_cfg.get("split_long_tables", True),
        split_ratio=table_cfg.get("split_ratio", 0.5),
        column_gap=table_cfg.get("column_gap", 0),
        cell_styles=cell_styles,
    )
    return {
        "xlsx_path": xlsx_path,
        "png_path": png_path if has_png else None,
        "markdown_lines": build_accuracy_markdown_table(headers, table_rows),
    }


def format_compare_value(field, value):
    """格式化指标变化中的单元格值，百分比统一输出为百分数。"""
    if field in PERCENT_FIELDS:
        if str(value or "").strip() == "-":
            value = "0%"
        return fmt_percent(parse_percent(value))
    if field == "点位数量":
        if value is None or str(value).strip() in ("", "-", "--"):
            return "-"
        return str(parse_int(value))
    if value is None or str(value).strip() in ("", "-", "--"):
        return "-"
    return str(value).strip()


def format_unchanged_compare_value(field, row):
    """格式化指标变化表中未变化字段的当前值。"""
    return format_compare_value(field, (row or {}).get(field))


def format_signed_percent(value):
    """格式化带正负号的百分比差值。"""
    if value is None:
        return "-"
    sign = "+" if value > 0 else ""
    return "%s%.0f%%" % (sign, value)


def format_signed_int(value):
    """格式化带正负号的整数差值。"""
    if value is None:
        return "-"
    sign = "+" if value > 0 else ""
    return "%s%d" % (sign, int(value))


def get_week_range_labels(today):
    """根据报告日期推算本周和上周的显示区间。"""
    current_start = today - timedelta(days=today.weekday())
    previous_start = current_start - timedelta(days=7)
    previous_end = current_start - timedelta(days=1)
    return {
        "current_start": current_start,
        "current_end": today,
        "previous_start": previous_start,
        "previous_end": previous_end,
        "current_label": "%s-%s" % (
            current_start.strftime("%m.%d"),
            today.strftime("%m.%d"),
        ),
        "previous_label": "%s~%s" % (
            previous_start.strftime("%m.%d"),
            previous_end.strftime("%m.%d"),
        ),
    }


def normalize_hex_color(value, default=None):
    """把 '#RRGGBB' 之类的颜色值规范成 RRGGBB。"""
    if value is None:
        return default
    s = str(value).strip()
    if not s:
        return default
    s = s.lstrip("#")
    if len(s) == 6 and re.fullmatch(r"[0-9a-fA-F]{6}", s):
        return s.upper()
    if len(s) == 8 and re.fullmatch(r"[0-9a-fA-F]{8}", s):
        return s.upper()
    return default


def color_for_pillow(value, default):
    """给 Pillow 使用的颜色值补 '#' 前缀。"""
    color = normalize_hex_color(value, default)
    if color is None:
        return default
    return color if str(color).startswith("#") else "#" + str(color)


def normalize_style_rows(styles, row_count):
    """把行样式补齐到指定长度。"""
    normalized = list(styles or [])[:row_count]
    if len(normalized) < row_count:
        normalized.extend([None] * (row_count - len(normalized)))
    return normalized


def normalize_style_matrix(styles, row_count, col_count):
    """把单元格样式矩阵补齐到指定尺寸。"""
    normalized = []
    styles = list(styles or [])
    for row_index in range(row_count):
        row = styles[row_index] if row_index < len(styles) else None
        row = list(row or [])[:col_count]
        if len(row) < col_count:
            row.extend([None] * (col_count - len(row)))
        normalized.append(row)
    return normalized


def resolve_style_color(style, key, default=None):
    """从样式字典中取颜色。"""
    if not style:
        return default
    if isinstance(style, str):
        return normalize_hex_color(style, default) if key == "fill" else default
    if not isinstance(style, dict):
        return default
    return normalize_hex_color(style.get(key), default)


def coerce_text_list(value):
    """把 YAML 中的单条或多条文案统一成列表。"""
    if value is None:
        return []
    if isinstance(value, str):
        return [value]
    if isinstance(value, (list, tuple)):
        return [str(item) for item in value if str(item).strip()]
    return []


def render_template_text(text, context):
    """渲染 YAML 文案模板，不存在的变量保留原样。"""
    class SafeDict(dict):
        def __missing__(self, key):
            return "{" + key + "}"

    try:
        return str(text).format_map(SafeDict(context or {}))
    except (ValueError, KeyError):
        return str(text)


def condition_matches(condition, context):
    """判断 YAML 文案条件是否满足。"""
    condition = str(condition or "always").strip()
    if not condition or condition == "always":
        return True
    aliases = {
        "has_blur_risk": "has_blur",
        "has_stale_risk": "has_stale",
        "has_idle_risk": "has_idle",
        "has_high_fp": "has_high_fp",
        "has_low_accuracy": "has_low_accuracy",
        "has_todo_log": "has_todo_log",
        "has_todo_upload": "has_todo_upload",
    }
    if condition in context:
        return bool(context.get(condition))
    if condition in aliases:
        return bool(context.get(aliases[condition]))
    return False


def append_configured_texts(lines, section_cfg, context, item_no):
    """按 YAML 条件追加配置化文案。"""
    templates = section_cfg.get("templates", []) if section_cfg else []
    if not isinstance(templates, list):
        return item_no
    for template in templates:
        if isinstance(template, str):
            condition = "always"
            texts = [template]
        elif isinstance(template, dict):
            condition = template.get("when", "always")
            texts = coerce_text_list(template.get("text", template.get("texts")))
        else:
            continue
        if not condition_matches(condition, context):
            continue
        for text in texts:
            rendered = render_template_text(text, context).strip()
            if rendered:
                lines.append("%d、%s" % (item_no, rendered))
                item_no += 1
    return item_no


def table_field_label(field):
    """生成表格列头，审核后字段使用简洁的基础字段名。"""
    field = str(field or "").strip()
    return field[3:] if field.startswith("审核后") else field


def is_empty_table_value(value):
    """判断表格单元格是否为空占位值。"""
    return is_blank_or_dash(value)


def compact_table_columns(headers, table_rows):
    """删除整列为空或仅为 '-' 的表格列，保留首列算法小类。"""
    if not headers or not table_rows:
        return headers, table_rows

    keep_indexes = [0]
    for index in range(1, len(headers)):
        if any(
            index < len(row) and not is_empty_table_value(row[index])
            for row in table_rows
        ):
            keep_indexes.append(index)

    return (
        [headers[index] for index in keep_indexes],
        [[row[index] for index in keep_indexes] for row in table_rows],
    )


def build_cf_html(fragment):
    """构造 Windows 剪贴板 HTML Format 数据。"""
    html_doc = (
        "<!DOCTYPE html><html><head><meta charset=\"utf-8\"></head>"
        "<body style=\"font-family:'Microsoft YaHei',Arial,sans-serif;"
        "font-size:14px;line-height:1.55;color:#000000;\">"
        "<!--StartFragment-->%s<!--EndFragment--></body></html>"
    ) % fragment
    header_tpl = (
        "Version:0.9\r\n"
        "SourceURL:about:blank\r\n"
        "StartHTML:%010d\r\n"
        "EndHTML:%010d\r\n"
        "StartFragment:%010d\r\n"
        "EndFragment:%010d\r\n"
    )
    dummy_header = header_tpl % (0, 0, 0, 0)
    start_html = len(dummy_header)
    start_marker = "<!--StartFragment-->"
    end_marker = "<!--EndFragment-->"
    html_bytes = html_doc.encode("utf-8")
    start_fragment = start_html + len(html_doc[:html_doc.index(start_marker) + len(start_marker)].encode("utf-8"))
    end_fragment = start_html + len(html_doc[:html_doc.index(end_marker)].encode("utf-8"))
    end_html = start_html + len(html_bytes)
    header = header_tpl % (start_html, end_html, start_fragment, end_fragment)
    return header.encode("ascii") + html_bytes


def report_to_clipboard_html(report, base_dir):
    """把 Markdown 周报转换成剪贴板 HTML，图片以内嵌 base64 输出。"""
    import base64
    import html

    html_lines = []
    image_pattern = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.*?)\)$")
    for line in report.splitlines():
        image_match = image_pattern.match(line.strip())
        if image_match:
            image_path = image_match.group("path").strip()
            if not os.path.isabs(image_path):
                image_path = os.path.join(base_dir, image_path)
            if os.path.exists(image_path):
                with open(image_path, "rb") as f:
                    encoded = base64.b64encode(f.read()).decode("ascii")
                alt = html.escape(image_match.group("alt") or "整体准确率明细")
                html_lines.append(
                    "<div style=\"margin:8px 0;\">"
                    "<img alt=\"%s\" src=\"data:image/png;base64,%s\" "
                    "style=\"max-width:100%%;height:auto;display:block;\" />"
                    "</div>" % (alt, encoded)
                )
            continue
        if line == "":
            continue
        else:
            html_lines.append(
                "<div style=\"white-space:pre-wrap;\">%s</div>" %
                html.escape(line)
            )
    return build_cf_html("\n".join(html_lines))


def report_to_share_html(report, base_dir, site_name, report_date, upload_metas=None):
    """生成局域网复制页面 HTML。"""
    import base64
    import html

    body_parts = []
    image_pattern = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.*?)\)$")
    plain_text = report_to_plain_clipboard_text(report)
    if upload_metas and "old" not in upload_metas and "result" not in upload_metas:
        upload_metas = {"old": None, "result": upload_metas}
    upload_metas = upload_metas or {}
    upload_payload = {}
    for kind in ("old", "result"):
        meta = upload_metas.get(kind)
        upload_payload[kind] = dict(meta, exists=True) if meta else None
    for line in report.splitlines():
        image_match = image_pattern.match(line.strip())
        if image_match:
            image_path = image_match.group("path").strip()
            if not os.path.isabs(image_path):
                image_path = os.path.join(base_dir, image_path)
            if os.path.exists(image_path):
                with open(image_path, "rb") as f:
                    encoded = base64.b64encode(f.read()).decode("ascii")
                alt = html.escape(image_match.group("alt") or "表格图片")
                body_parts.append(
                    '<img alt="%s" src="data:image/png;base64,%s">' %
                    (alt, encoded)
                )
            continue
        if line:
            body_parts.append("<p>%s</p>" % html.escape(line))
        else:
            body_parts.append("<br>")

    title = "%s 周报 %s" % (site_name, report_date)
    return """<!doctype html>
<html>
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{title}</title>
  <style>
    body {{ font-family: "Microsoft YaHei", Arial, sans-serif; margin: 0; color: #111; background: #f6f7f9; }}
    .bar {{ position: sticky; top: 0; z-index: 2; display: flex; gap: 8px; align-items: center; flex-wrap: wrap; padding: 12px; background: #fff; border-bottom: 1px solid #ddd; }}
    button {{ height: 36px; padding: 0 14px; border: 1px solid #222; background: #111; color: #fff; cursor: pointer; }}
    .status {{ color: #555; font-size: 14px; }}
    .hint {{ width: 100%; color: #777; font-size: 13px; }}
    .upload {{ display: flex; gap: 8px; align-items: center; flex-wrap: wrap; width: 100%; padding-top: 4px; }}
    .upload-label {{ min-width: 90px; font-weight: 600; }}
    .upload input {{ max-width: 260px; }}
    .upload button[disabled] {{ opacity: .45; cursor: not-allowed; }}
    .upload-state {{ color: #333; font-size: 14px; min-width: 210px; }}
    .page {{ max-width: 960px; margin: 18px auto; padding: 24px; background: #fff; }}
    p {{ margin: 0 0 8px; white-space: pre-wrap; line-height: 1.55; }}
    img {{ max-width: 100%; height: auto; display: block; margin: 10px 0 16px; }}
    textarea {{ display: block; width: calc(100% - 24px); max-width: 960px; min-height: 180px; margin: 0 auto 18px; padding: 12px; border: 1px solid #ddd; font: 13px/1.5 Consolas, "Microsoft YaHei", monospace; }}
    textarea.hidden {{ position: fixed; left: -9999px; top: -9999px; width: 1px; height: 1px; }}
  </style>
</head>
<body>
  <div class="bar">
    <button onclick="copyText()">复制文字</button>
    <button onclick="copyHtml()">复制图文</button>
    <button onclick="showText()">显示文字</button>
    <span id="status" class="status">正在准备剪贴板文字...</span>
    <div class="hint">同一局域网电脑打开本页后，点击复制文字即可粘贴；若浏览器拦截剪贴板，可点显示文字后手动复制。</div>
    <div class="upload" data-kind="old">
      <span class="upload-label">old.xlsx</span>
      <span id="oldUploadState" class="upload-state"></span>
      <input id="oldXlsxFile" type="file" accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet">
      <button id="oldUploadButton" onclick="uploadXlsx('old')">暂存 old</button>
      <button id="oldDeleteButton" onclick="deleteXlsx('old')">删除 old</button>
    </div>
    <div class="upload" data-kind="result">
      <span class="upload-label">result-2.xlsx</span>
      <span id="resultUploadState" class="upload-state"></span>
      <input id="resultXlsxFile" type="file" accept=".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet">
      <button id="resultUploadButton" onclick="uploadXlsx('result')">暂存最新table</button>
      <button id="resultDeleteButton" onclick="deleteXlsx('result')">删除最新table</button>
      <div class="hint">网页最多保留一份 old.xlsx 和一份 result-2.xlsx；old 未删除时优先作为老数据，result-2 未删除时作为新数据。若只上传 result-2，本机 result-2.xlsx 本次作为老数据。</div>
    </div>
  </div>
  <div id="report" class="page">
{body}
  </div>
  <textarea id="plain" class="hidden" readonly></textarea>
    <script>
    const plainText = {plain_json};
    let uploadMetas = {upload_json};
    const uploadDefs = {{
      old: {{ label: "old.xlsx", endpoint: "/api/temp-old", emptyText: "未暂存 old。", uploadedText: "已暂存 old。", fileId: "oldXlsxFile", stateId: "oldUploadState", uploadId: "oldUploadButton", deleteId: "oldDeleteButton" }},
      result: {{ label: "result-2.xlsx", endpoint: "/api/temp-result", emptyText: "未暂存最新table。", uploadedText: "已暂存最新table。", fileId: "resultXlsxFile", stateId: "resultUploadState", uploadId: "resultUploadButton", deleteId: "resultDeleteButton" }}
    }};
    const statusEl = document.getElementById("status");
    const plainEl = document.getElementById("plain");
    document.getElementById("plain").value = plainText;
    function setStatus(text) {{
      statusEl.textContent = text;
    }}
    function selectPlainText() {{
      plainEl.classList.remove("hidden");
      plainEl.focus();
      plainEl.select();
    }}
    async function copyText() {{
      try {{
        await navigator.clipboard.writeText(plainText);
        setStatus("已复制文字。");
        plainEl.classList.add("hidden");
      }} catch (e) {{
        selectPlainText();
        const ok = document.execCommand("copy");
        setStatus(ok ? "已复制文字。" : "浏览器未允许自动复制，请手动复制下方文字。");
      }}
    }}
    async function copyHtml() {{
      const report = document.getElementById("report");
      try {{
        const item = new ClipboardItem({{
          "text/html": new Blob([report.innerHTML], {{ type: "text/html" }}),
          "text/plain": new Blob([plainText], {{ type: "text/plain" }})
        }});
        await navigator.clipboard.write([item]);
        setStatus("已复制图文。");
      }} catch (e) {{
        const range = document.createRange();
        range.selectNodeContents(report);
        const selection = window.getSelection();
        selection.removeAllRanges();
        selection.addRange(range);
        const ok = document.execCommand("copy");
        selection.removeAllRanges();
        if (ok) {{
          setStatus("已尝试复制图文。");
        }} else {{
          await copyText();
        }}
      }}
    }}
    function showText() {{
      selectPlainText();
      setStatus("文字已显示，可手动复制。");
    }}
    function formatSize(size) {{
      if (!size) return "";
      if (size < 1024) return size + " B";
      if (size < 1024 * 1024) return (size / 1024).toFixed(1) + " KB";
      return (size / 1024 / 1024).toFixed(2) + " MB";
    }}
    function renderUploadState(kind) {{
      const def = uploadDefs[kind];
      const meta = uploadMetas[kind];
      const uploadStateEl = document.getElementById(def.stateId);
      const fileEl = document.getElementById(def.fileId);
      const uploadButton = document.getElementById(def.uploadId);
      const deleteButton = document.getElementById(def.deleteId);
      if (meta) {{
        const details = [meta.originalName || def.label, formatSize(meta.size), meta.mtime]
          .filter(Boolean)
          .join(" / ");
        uploadStateEl.textContent = "已暂存：" + details;
        fileEl.disabled = true;
        uploadButton.disabled = true;
        deleteButton.disabled = false;
      }} else {{
        uploadStateEl.textContent = def.emptyText;
        fileEl.disabled = false;
        uploadButton.disabled = false;
        deleteButton.disabled = true;
      }}
    }}
    async function refreshUploadState(kind) {{
      const def = uploadDefs[kind];
      try {{
        const response = await fetch(def.endpoint);
        const meta = await response.json();
        uploadMetas[kind] = meta && meta.exists ? meta : null;
        renderUploadState(kind);
      }} catch (e) {{
        renderUploadState(kind);
      }}
    }}
    async function uploadXlsx(kind) {{
      const def = uploadDefs[kind];
      const fileEl = document.getElementById(def.fileId);
      if (uploadMetas[kind]) {{
        setStatus("已有暂存 " + def.label + "，请先删除后再添加。");
        return;
      }}
      const file = fileEl.files && fileEl.files[0];
      if (!file) {{
        setStatus("请选择 " + def.label + "。");
        return;
      }}
      const form = new FormData();
      form.append("file", file);
      const response = await fetch(def.endpoint, {{ method: "POST", body: form }});
      const data = await response.json();
      if (!response.ok || !data.ok) {{
        setStatus(data.error || "上传失败。");
        return;
      }}
      uploadMetas[kind] = data.file;
      fileEl.value = "";
      renderUploadState(kind);
      setStatus(def.uploadedText);
    }}
    async function deleteXlsx(kind) {{
      const def = uploadDefs[kind];
      const response = await fetch(def.endpoint, {{ method: "DELETE" }});
      const data = await response.json();
      if (!response.ok || !data.ok) {{
        setStatus(data.error || "删除失败。");
        return;
      }}
      uploadMetas[kind] = null;
      renderUploadState(kind);
      setStatus("已删除暂存 " + def.label + "。");
    }}
    window.addEventListener("load", () => {{
      Object.keys(uploadDefs).forEach((kind) => {{
        renderUploadState(kind);
        refreshUploadState(kind);
      }});
      copyText().catch(() => setStatus("点击复制文字即可粘贴。"));
    }});
  </script>
</body>
</html>
""".format(
        title=html.escape(title),
        body="\n".join("    " + part for part in body_parts),
        plain_json=json.dumps(plain_text, ensure_ascii=False),
        upload_json=json.dumps(upload_payload, ensure_ascii=False),
    )


def get_lan_ip():
    """获取优先用于局域网访问的本机 IP。"""
    candidates = []
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            sock.connect(("8.8.8.8", 80))
            candidates.append(sock.getsockname()[0])
        finally:
            sock.close()
    except OSError:
        pass

    try:
        hostname = socket.gethostname()
        candidates.extend(socket.gethostbyname_ex(hostname)[2])
    except OSError:
        pass

    for ip in candidates:
        if ip and not ip.startswith("127.") and not ip.startswith("169.254."):
            return ip
    return "127.0.0.1"


def is_tcp_port_open(host, port):
    try:
        with socket.create_connection((host, port), timeout=0.3):
            return True
    except OSError:
        return False


def is_clipboard_share_server(port, session_id=None):
    """判断端口上是否已经有可复用的局域网复制页服务。"""
    try:
        from urllib.request import urlopen

        with urlopen("http://127.0.0.1:%d/api/status" % port, timeout=0.8) as response:
            content = response.read(4096).decode("utf-8", errors="ignore")
        data = json.loads(content)
        if not bool(data.get("ok")) or data.get("service") != "weekly_report_share":
            return False
        if session_id is None:
            return True
        return data.get("sessionId") == session_id
    except Exception:
        return False


def start_lan_share_service(host, port, output_dir, upload_dir):
    creationflags = 0
    startupinfo = None
    if os.name == "nt":
        creationflags = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0)
        startupinfo = subprocess.STARTUPINFO()
        startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startupinfo.wShowWindow = 0
    try:
        process = subprocess.Popen(
            [
                sys.executable,
                os.path.abspath(__file__),
                "--lan-share-server",
                host,
                str(port),
                output_dir,
                upload_dir,
            ],
            cwd=BASE_DIR,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            stdin=subprocess.DEVNULL,
            creationflags=creationflags,
            startupinfo=startupinfo,
        )
    except OSError:
        return False
    session_id = get_session_id_from_dir(upload_dir)
    for _ in range(25):
        if process.poll() is not None:
            return False
        if is_clipboard_share_server(port, session_id=session_id):
            return True
        time.sleep(0.2)
    return False


def start_lan_clipboard_server(output_dir, clipboard_cfg, session_dir=None):
    """按配置启动局域网访问用的静态文件服务。"""
    if not bool(clipboard_cfg.get("lan_share", False)):
        return None

    host = str(clipboard_cfg.get("lan_host", "0.0.0.0") or "0.0.0.0")
    start_port = int(clipboard_cfg.get("lan_port", 8765) or 8765)
    upload_dir = ensure_lan_session_dir(output_dir, session_dir=session_dir)
    session_id = get_session_id_from_dir(upload_dir)
    port = None
    should_start = False
    for candidate_port in range(start_port, start_port + 10):
        if is_clipboard_share_server(candidate_port, session_id=session_id):
            port = candidate_port
            should_start = False
            break
        if not is_tcp_port_open("127.0.0.1", candidate_port):
            port = candidate_port
            should_start = True
            break
    if port is None:
        return None

    if should_start:
        if not start_lan_share_service(host, port, output_dir, upload_dir):
            return None

    return "http://%s:%d/clipboard_share.html" % (get_lan_ip(), port)


def run_lan_share_server(host, port, output_dir, upload_dir=None):
    """运行局域网复制页服务，并支持上传/删除临时 old/result-2 表。"""
    import http.server
    import tempfile
    from urllib.parse import unquote

    output_dir = os.path.abspath(output_dir)
    upload_dir = ensure_lan_session_dir(output_dir, session_dir=upload_dir)
    session_id = get_session_id_from_dir(upload_dir)
    os.makedirs(output_dir, exist_ok=True)

    def parse_multipart_file(content_type, body):
        boundary_match = re.search(
            r'boundary=(?:"([^"]+)"|([^;]+))',
            content_type,
            flags=re.I,
        )
        if not boundary_match:
            return None, None, "缺少上传边界。"
        boundary = (boundary_match.group(1) or boundary_match.group(2)).strip()
        marker = b"--" + boundary.encode("utf-8")
        for raw_part in body.split(marker):
            part = raw_part
            if not part or part in (b"--", b"--\r\n"):
                continue
            if part.startswith(b"\r\n"):
                part = part[2:]
            if part.endswith(b"\r\n"):
                part = part[:-2]
            if part.endswith(b"--"):
                part = part[:-2]
                if part.endswith(b"\r\n"):
                    part = part[:-2]
            header_bytes, sep, payload = part.partition(b"\r\n\r\n")
            if not sep:
                continue
            headers = header_bytes.decode("utf-8", errors="ignore")
            disposition = ""
            for line in headers.splitlines():
                if line.lower().startswith("content-disposition:"):
                    disposition = line
                    break
            if 'name="file"' not in disposition:
                continue
            filename_match = re.search(r'filename="([^"]*)"', disposition)
            filename = filename_match.group(1) if filename_match else ""
            if not filename:
                return None, None, "没有收到文件。"
            return filename, payload, None
        return None, None, "没有收到文件。"

    def upload_kind_from_path(path):
        request_path = path.split("?", 1)[0]
        if request_path == "/api/temp-old":
            return "old"
        if request_path == "/api/temp-result":
            return "result"
        return None

    class LanShareHandler(http.server.SimpleHTTPRequestHandler):
        server_version = "WeeklyReportShare/1.0"

        def __init__(self, *args, **kwargs):
            super().__init__(*args, directory=output_dir, **kwargs)

        def handle(self):
            try:
                super().handle()
            except (ConnectionAbortedError, ConnectionResetError, BrokenPipeError):
                return

        def log_message(self, format, *args):
            return

        def send_json(self, status, payload):
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            self.send_response(status)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            if self.path.split("?", 1)[0] == "/api/status":
                self.send_json(200, {
                    "ok": True,
                    "service": "weekly_report_share",
                    "sessionId": session_id,
                })
                return
            upload_kind = upload_kind_from_path(self.path)
            if upload_kind:
                meta = read_lan_upload_metadata(upload_dir, upload_kind)
                if meta:
                    meta["exists"] = True
                    self.send_json(200, meta)
                else:
                    self.send_json(200, {"exists": False})
                return
            return super().do_GET()

        def do_DELETE(self):
            upload_kind = upload_kind_from_path(self.path)
            if not upload_kind:
                self.send_error(404)
                return
            upload_path, meta_path = get_lan_upload_paths(upload_dir, upload_kind)
            removed = False
            for path in (upload_path, meta_path):
                try:
                    if os.path.exists(path):
                        os.remove(path)
                        removed = True
                except OSError as exc:
                    self.send_json(500, {
                        "ok": False,
                        "error": "删除失败: %s" % exc,
                    })
                    return
            self.send_json(200, {"ok": True, "removed": removed})

        def do_POST(self):
            upload_kind = upload_kind_from_path(self.path)
            if not upload_kind:
                self.send_error(404)
                return

            upload_cfg = LAN_UPLOAD_KINDS[upload_kind]
            upload_path, meta_path = get_lan_upload_paths(upload_dir, upload_kind)
            if os.path.exists(upload_path):
                self.send_json(409, {
                    "ok": False,
                    "error": "已有暂存 %s，请先删除后再添加。" % upload_cfg["default_name"],
                })
                return

            content_type = self.headers.get("Content-Type", "")
            if "multipart/form-data" not in content_type:
                self.send_json(400, {
                    "ok": False,
                    "error": "请选择 .xlsx 文件上传。",
                })
                return

            try:
                content_length = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                content_length = 0
            if content_length <= 0:
                self.send_json(400, {
                    "ok": False,
                    "error": "没有收到文件。",
                })
                return

            body = self.rfile.read(content_length)
            filename, payload, parse_error = parse_multipart_file(content_type, body)
            if parse_error:
                self.send_json(400, {
                    "ok": False,
                    "error": parse_error,
                })
                return

            original_name = os.path.basename(unquote(filename))
            if not original_name.lower().endswith(".xlsx"):
                self.send_json(400, {
                    "ok": False,
                    "error": "只支持上传 Microsoft Excel 工作表 (.xlsx)。",
                })
                return

            if not payload:
                self.send_json(400, {
                    "ok": False,
                    "error": "上传文件为空。",
                })
                return

            fd, tmp_path = tempfile.mkstemp(
                prefix=upload_cfg["temp_prefix"],
                suffix=".xlsx",
                dir=upload_dir,
            )
            try:
                with os.fdopen(fd, "wb") as tmp_file:
                    tmp_file.write(payload)
                ok, error = validate_xlsx_file(tmp_path)
                if not ok:
                    self.send_json(400, {
                        "ok": False,
                        "error": "上传文件不是可读取的 .xlsx: %s" % error,
                    })
                    return
                os.replace(tmp_path, upload_path)
                write_lan_upload_metadata(
                    meta_path,
                    original_name,
                    upload_path,
                    upload_kind,
                )
                meta = read_lan_upload_metadata(upload_dir, upload_kind) or {}
                meta["exists"] = True
                self.send_json(200, {"ok": True, "file": meta})
            finally:
                try:
                    if os.path.exists(tmp_path):
                        os.remove(tmp_path)
                except OSError:
                    pass

    server = http.server.ThreadingHTTPServer((host, int(port)), LanShareHandler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


def report_to_clipboard_rtf(report, base_dir):
    """把 Markdown 周报转换成带内嵌 PNG 的 RTF 剪贴板数据。"""
    import binascii

    image_pattern = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.*?)\)$")

    def escape_text(text):
        # 使用 RTF Unicode 控制字，避免系统代码页导致中文丢失。
        units = []
        encoded = str(text).encode("utf-16-le")
        for index in range(0, len(encoded), 2):
            value = encoded[index] | (encoded[index + 1] << 8)
            if value >= 0x8000:
                value -= 0x10000
            units.append("\\u%d?" % value)
        return "".join(units)

    parts = [
        r"{\rtf1\ansi\deff0",
        r"{\fonttbl{\f0 Microsoft YaHei;}}",
        r"\viewkind4\uc1\f0\fs28 ",
    ]
    for line in report.splitlines():
        image_match = image_pattern.match(line.strip())
        if image_match:
            image_path = image_match.group("path").strip()
            if not os.path.isabs(image_path):
                image_path = os.path.join(base_dir, image_path)
            if os.path.exists(image_path):
                with open(image_path, "rb") as f:
                    encoded = binascii.hexlify(f.read()).decode("ascii")
                parts.append(r"\pard\sa120\qc ")
                parts.append(r"{\pict\pngblip ")
                parts.append(encoded)
                parts.append(r"}\par ")
                continue

        if line:
            parts.append(r"\pard\sa120\ql ")
            parts.append(escape_text(line))
        parts.append(r"\par ")

    parts.append("}")
    return "".join(parts).encode("ascii")


def create_report_preview_image(report, base_dir):
    """将报告文本和其中的 PNG 合成为原生图片剪贴板兜底图。"""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return None

    regular_path = find_table_font()
    try:
        font = ImageFont.truetype(regular_path, 20) if regular_path else ImageFont.load_default()
    except OSError:
        font = ImageFont.load_default()

    image_pattern = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.*?)\)$")
    canvas_width = 1000
    padding_x = 24
    line_height = 30
    text_width_limit = canvas_width - padding_x * 2
    measure = Image.new("RGB", (10, 10), "white")
    measure_draw = ImageDraw.Draw(measure)

    def text_width(text):
        box = measure_draw.textbbox((0, 0), str(text), font=font)
        return box[2] - box[0]

    def wrap_text(text):
        text = str(text)
        if not text:
            return [""]
        lines = []
        current = ""
        for char in text:
            candidate = current + char
            if current and text_width(candidate) > text_width_limit:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines or [""]

    blocks = []
    for line in report.splitlines():
        image_match = image_pattern.match(line.strip())
        if image_match:
            image_path = image_match.group("path").strip()
            if not os.path.isabs(image_path):
                image_path = os.path.join(base_dir, image_path)
            if os.path.exists(image_path):
                try:
                    with Image.open(image_path) as source:
                        embedded = source.convert("RGB").copy()
                    max_image_width = canvas_width - padding_x * 2
                    if embedded.width > max_image_width:
                        ratio = max_image_width / float(embedded.width)
                        embedded = embedded.resize(
                            (max_image_width, max(1, int(embedded.height * ratio))),
                            Image.Resampling.LANCZOS,
                        )
                    blocks.append(("image", embedded))
                    continue
                except (OSError, ValueError):
                    pass
            blocks.append(("text", wrap_text(image_match.group("alt") or "表格图片")))
        elif line:
            blocks.append(("text", wrap_text(line)))
        else:
            blocks.append(("blank", None))

    margin_y = 24
    block_heights = []
    for block_type, value in blocks:
        if block_type == "image":
            block_heights.append(value.height + 16)
        elif block_type == "blank":
            block_heights.append(14)
        else:
            block_heights.append(max(line_height, len(value) * line_height))
    canvas_height = margin_y * 2 + sum(block_heights)
    image = Image.new("RGB", (canvas_width, canvas_height), "white")
    draw = ImageDraw.Draw(image)
    y = margin_y
    for (block_type, value), block_height in zip(blocks, block_heights):
        if block_type == "image":
            x = (canvas_width - value.width) // 2
            image.paste(value, (x, y))
            y += block_height
        elif block_type == "blank":
            y += block_height
        else:
            for line in value:
                draw.text((padding_x, y), line, font=font, fill="#000000")
                y += line_height
            y += max(0, block_height - len(value) * line_height)
    return image


def image_to_png_bytes(image):
    """将 Pillow 图片编码为 PNG 字节。"""
    from io import BytesIO

    buffer = BytesIO()
    image.save(buffer, "PNG")
    return buffer.getvalue()


def image_to_dib_bytes(image):
    """将 Pillow 图片编码为 Windows CF_DIB 数据。"""
    import struct
    from PIL import Image

    rgb = image.convert("RGB")
    width, height = rgb.size
    row_size = width * 3
    padded_row_size = (row_size + 3) & ~3
    top_down_bgr = rgb.tobytes("raw", "BGR")
    rows = []
    for row_index in range(height - 1, -1, -1):
        row = top_down_bgr[row_index * row_size:(row_index + 1) * row_size]
        rows.append(row + b"\x00" * (padded_row_size - row_size))
    pixels = b"".join(rows)
    header = struct.pack(
        "<IiiHHIIiiII",
        40,
        width,
        height,
        1,
        24,
        0,
        len(pixels),
        0,
        0,
        0,
        0,
    )
    return header + pixels


def report_to_plain_clipboard_text(report):
    """生成普通文本剪贴板兜底内容，避免粘贴出 Markdown 图片路径。"""
    lines = []
    image_pattern = re.compile(r"^!\[(?P<alt>.*?)\]\((?P<path>.*?)\)$")
    for line in report.splitlines():
        image_match = image_pattern.match(line.strip())
        if image_match:
            lines.append("[整体准确率明细图片]")
        else:
            lines.append(line)
    return "\n".join(lines)


def copy_rich_to_clipboard(report, base_dir, pure_image=False):
    """复制报告到剪贴板，可选整份报告纯图片输出。"""
    import ctypes
    import locale

    user32 = ctypes.windll.user32
    kernel32 = ctypes.windll.kernel32
    CF_TEXT = 1
    CF_UNICODETEXT = 13
    CF_DIB = 8
    user32.RegisterClipboardFormatW.argtypes = [ctypes.c_wchar_p]
    user32.RegisterClipboardFormatW.restype = ctypes.c_uint
    user32.OpenClipboard.argtypes = [ctypes.c_void_p]
    user32.OpenClipboard.restype = ctypes.c_bool
    user32.EmptyClipboard.restype = ctypes.c_bool
    user32.CloseClipboard.restype = ctypes.c_bool
    user32.SetClipboardData.argtypes = [ctypes.c_uint, ctypes.c_void_p]
    user32.SetClipboardData.restype = ctypes.c_void_p
    kernel32.GetACP.restype = ctypes.c_uint
    kernel32.GlobalAlloc.argtypes = [ctypes.c_uint, ctypes.c_size_t]
    kernel32.GlobalAlloc.restype = ctypes.c_void_p
    kernel32.GlobalLock.argtypes = [ctypes.c_void_p]
    kernel32.GlobalLock.restype = ctypes.c_void_p
    kernel32.GlobalUnlock.argtypes = [ctypes.c_void_p]
    kernel32.GlobalUnlock.restype = ctypes.c_bool
    kernel32.GlobalFree.argtypes = [ctypes.c_void_p]
    kernel32.GlobalFree.restype = ctypes.c_void_p
    html_format = user32.RegisterClipboardFormatW("HTML Format")
    rtf_format = user32.RegisterClipboardFormatW("Rich Text Format")
    png_format = user32.RegisterClipboardFormatW("PNG")
    if not html_format or not rtf_format:
        return False

    def alloc_global(data_bytes):
        GMEM_MOVEABLE = 0x0002
        hglobal = kernel32.GlobalAlloc(GMEM_MOVEABLE, len(data_bytes))
        if not hglobal:
            raise ctypes.WinError(ctypes.get_last_error())
        ptr = kernel32.GlobalLock(hglobal)
        if not ptr:
            kernel32.GlobalFree(hglobal)
            raise ctypes.WinError(ctypes.get_last_error())
        ctypes.memmove(ptr, data_bytes, len(data_bytes))
        kernel32.GlobalUnlock(hglobal)
        return hglobal

    def set_clipboard_formats(plain_text):
        html_bytes = report_to_clipboard_html(report, base_dir)
        rtf_bytes = report_to_clipboard_rtf(report, base_dir)
        native_formats = []
        if pure_image:
            try:
                preview = create_report_preview_image(report, base_dir)
                if preview is not None:
                    native_formats.append((CF_DIB, image_to_dib_bytes(preview)))
                    if png_format:
                        native_formats.append((png_format, image_to_png_bytes(preview)))
            except Exception:
                native_formats = []
        try:
            ansi_name = "cp%d" % kernel32.GetACP()
            ansi_bytes = plain_text.encode(ansi_name, errors="replace") + b"\x00"
        except Exception:
            ansi_bytes = plain_text.encode(locale.getpreferredencoding(False), errors="replace") + b"\x00"
        unicode_bytes = plain_text.encode("utf-16le") + b"\x00\x00"
        formats = [
            (CF_UNICODETEXT, unicode_bytes),
            (CF_TEXT, ansi_bytes),
            (html_format, html_bytes),
            (rtf_format, rtf_bytes),
        ]
        if pure_image and native_formats:
            formats = native_formats
        if not user32.OpenClipboard(None):
            raise ctypes.WinError(ctypes.get_last_error())
        handles = []
        try:
            if not user32.EmptyClipboard():
                raise ctypes.WinError(ctypes.get_last_error())
            for fmt, data_bytes in formats:
                hglobal = alloc_global(data_bytes)
                handles.append(hglobal)
                if not user32.SetClipboardData(fmt, hglobal):
                    raise ctypes.WinError(ctypes.get_last_error())
            return True
        finally:
            user32.CloseClipboard()

    try:
        if set_clipboard_formats(report):
            return "rich"
    except Exception:
        pass

    try:
        plain_text = report_to_plain_clipboard_text(report)
        ansi_name = "cp%d" % kernel32.GetACP()
        ansi_bytes = plain_text.encode(ansi_name, errors="replace")
        unicode_bytes = plain_text.encode("utf-16le") + b"\x00\x00"
        if not user32.OpenClipboard(None):
            return False
        try:
            if not user32.EmptyClipboard():
                return False
            for fmt, data_bytes in ((CF_UNICODETEXT, unicode_bytes), (CF_TEXT, ansi_bytes + b"\x00")):
                hglobal = alloc_global(data_bytes)
                if not user32.SetClipboardData(fmt, hglobal):
                    raise ctypes.WinError(ctypes.get_last_error())
            return "text"
        finally:
            user32.CloseClipboard()
    except Exception:
        return False


def get_field_cfg(cfg, field):
    """获取字段配置，不存在时返回空 dict"""
    fields = cfg.get("fields", {})
    return fields.get(field, {}) or {}


def get_priority_value(row, field, cfg):
    """
    获取优先展示值（用于准确率/误报率等单数值展示）
    优先审核后值，审核后无效时用原始值
    """
    field_cfg = get_field_cfg(cfg, field)
    priority = field_cfg.get("priority", "audited")

    is_audited = "审核后" in field
    raw_field = field.replace("审核后", "") if is_audited else field
    aud_field = "审核后" + field if not is_audited else field

    raw_val = parse_percent(row.get(raw_field))
    aud_val = parse_percent(row.get(aud_field))

    if priority == "audited":
        return aud_val if aud_val is not None else raw_val
    else:
        return raw_val if raw_val is not None else aud_val


def is_below_threshold(row, field, cfg):
    """判断字段值是否低于阈值（准确率类）"""
    field_cfg = get_field_cfg(cfg, field)
    threshold = field_cfg.get("threshold")
    if threshold is None:
        return False
    val = get_priority_value(row, field, cfg)
    return val is not None and val < threshold


def is_above_threshold(row, field, cfg):
    """判断字段值是否高于阈值（误报率/漏报率类）"""
    field_cfg = get_field_cfg(cfg, field)
    threshold = field_cfg.get("threshold")
    if threshold is None:
        return False
    val = get_priority_value(row, field, cfg)
    return val is not None and val > threshold


def get_algorithm_control(cfg):
    """读取算法小类显示和计算控制配置。"""
    return cfg.get("algorithm_control", {}) or {}


def algorithm_control_names(control):
    """获取受控算法小类名称列表。"""
    names = control.get(
        "algorithm_classes",
        control.get("算法小类", []),
    )
    if isinstance(names, str):
        names = [names]
    return {
        str(name).strip()
        for name in (names or [])
        if str(name).strip()
    }


def algorithm_control_enabled(control):
    return bool(control.get("enabled", True)) and bool(algorithm_control_names(control))


def should_keep_algorithm(row, control, table_key=None, for_calculation=False):
    """判断算法小类是否保留在计算或指定表格中。"""
    if not algorithm_control_enabled(control):
        return True

    name = str(row.get("算法小类") or "").strip()
    if name not in algorithm_control_names(control):
        return True

    if for_calculation:
        return bool(control.get(
            "include_in_calculation",
            control.get("参与计算", True),
        ))

    tables = control.get("tables", control.get("表格", {})) or {}
    if table_key and not bool(tables.get(table_key, True)):
        return True

    return bool(control.get("display", control.get("显示", True)))


def filter_algorithm_rows(rows, control, table_key=None, for_calculation=False):
    return [
        row for row in rows
        if should_keep_algorithm(row, control, table_key, for_calculation)
    ]


def filter_algorithm_items(items, control, table_key=None):
    """过滤以算法小类名称开头的字符串或元组列表。"""
    if not algorithm_control_enabled(control):
        return list(items)

    filtered = []
    names = algorithm_control_names(control)
    tables = control.get("tables", control.get("表格", {})) or {}
    display = bool(control.get("display", control.get("显示", True)))
    table_enabled = bool(tables.get(table_key, True)) if table_key else True
    for item in items:
        name = item if isinstance(item, str) else (item[0] if item else "")
        if (
            str(name).strip() in names
            and table_enabled
            and not display
        ):
            continue
        filtered.append(item)
    return filtered


def parse_clipboard_mode_arg(value):
    """解析命令行剪贴板模式，返回 None 表示使用 YAML。"""
    mode = str(value or "").strip().lower()
    if not mode:
        return None
    text_modes = {"text", "rich", "mixed", "normal", "y", "yes", "false", "0"}
    image_modes = {"image", "pure_image", "pure-image", "n", "no", "true", "1"}
    if mode in text_modes:
        return False
    if mode in image_modes:
        return True
    raise ValueError("未知剪贴板模式: %s" % value)


def parse_accuracy_source_arg(value):
    """解析准确率数据源，返回 raw 或 audited，None 表示使用 YAML。"""
    source = str(value or "").strip().lower()
    if not source:
        return None
    raw_modes = {"raw", "accuracy", "准确率", "y", "yes", "1"}
    audited_modes = {
        "audited", "audit", "review", "审核后", "审核后准确率", "n", "no", "0"
    }
    if source in raw_modes:
        return "raw"
    if source in audited_modes:
        return "audited"
    raise ValueError("未知准确率数据源: %s" % value)


def parse_main_options(argv):
    """解析周报主流程的可选参数。前三个位置参数保持向后兼容。"""
    options = {
        "local_clipboard": False,
        "lan_share": False,
        "lan_session_dir": None,
        "lan_upload_dir": None,
        "no_clipboard": False,
    }
    positional = []
    i = 1
    while i < len(argv):
        arg = argv[i]
        if arg == "--local-clipboard":
            options["local_clipboard"] = True
            i += 1
        elif arg == "--lan-share":
            options["lan_share"] = True
            i += 1
        elif arg == "--lan-session-dir":
            if i + 1 >= len(argv):
                raise ValueError("--lan-session-dir 缺少目录参数")
            options["lan_session_dir"] = argv[i + 1]
            i += 2
        elif arg == "--lan-upload-dir":
            if i + 1 >= len(argv):
                raise ValueError("--lan-upload-dir 缺少目录参数")
            options["lan_upload_dir"] = argv[i + 1]
            i += 2
        elif arg == "--no-clipboard":
            options["no_clipboard"] = True
            i += 1
        else:
            positional.append(arg)
            i += 1
    return positional, options


def configured_accuracy_source(cfg):
    """读取 YAML 中准确率计算默认使用的数据源。"""
    field_cfg = get_field_cfg(cfg, "准确率")
    priority = str(field_cfg.get("priority", "")).strip().lower()
    if priority in ("raw", "audited"):
        return priority

    display = str(
        (cfg.get("overview", {}) or {}).get("accuracy_detail_display", "")
    ).strip().lower()
    if display in ("raw", "audited"):
        return display
    return "audited"


def apply_accuracy_source_override(cfg, source):
    """将本次运行选择的数据源同步到准确率展示和优先取值配置。"""
    overview_cfg = cfg.setdefault("overview", {})
    fields_cfg = cfg.setdefault("fields", {})
    accuracy_cfg = fields_cfg.setdefault("准确率", {})
    overview_cfg["accuracy_detail_display"] = source
    accuracy_cfg["display"] = source
    accuracy_cfg["priority"] = source


def main():
    try:
        positional_args, run_options = parse_main_options(sys.argv)
    except ValueError as exc:
        print("错误: %s" % exc)
        return 1

    if len(positional_args) < 1:
        # 无参数时交互式输入
        print("=" * 44)
        print("  周报生成工具")
        print("=" * 44)
        print("")
        keyword = input("请输入站点关键字（如：马王）: ").strip()
        print("")
    else:
        keyword = positional_args[0].strip()

    if not keyword:
        print("错误: 站点关键字不能为空")
        return 1

    cfg = load_config()
    old_file, new_file = resolve_data_files(cfg)
    output_dir = ensure_output_dir()
    clipboard_cfg = cfg.get("clipboard", {}) or {}
    local_clipboard_mode = bool(run_options.get("local_clipboard"))
    lan_share_mode = bool(run_options.get("lan_share"))
    if local_clipboard_mode:
        clipboard_cfg = dict(clipboard_cfg)
        clipboard_cfg["lan_share"] = False
    elif lan_share_mode:
        clipboard_cfg = dict(clipboard_cfg)
        clipboard_cfg["lan_share"] = True
    lan_session_dir = None
    if run_options.get("lan_upload_dir"):
        lan_session_dir = ensure_lan_session_dir(
            output_dir,
            session_dir=run_options.get("lan_upload_dir"),
        )
        print("[局域网Web] 本次临时目录: %s" % lan_session_dir)
    elif bool(clipboard_cfg.get("lan_share", False)):
        lan_session_dir = ensure_lan_session_dir(
            output_dir,
            session_dir=run_options.get("lan_session_dir"),
        )
        print("[局域网Web] 本次临时目录: %s" % lan_session_dir)
    data_temp_dir = lan_session_dir or output_dir
    old_file, new_file, using_lan_upload = resolve_effective_data_files(
        old_file,
        new_file,
        data_temp_dir,
    )
    thresholds = cfg.get("thresholds", {})
    sections = cfg.get("sections", {})
    overview_cfg = cfg.get("overview", {})
    changes_cfg = cfg.get("changes", {})
    handled_cfg = cfg.get("handled", {})
    next_steps_cfg = cfg.get("next_steps", {})
    risks_cfg = cfg.get("risks", {})
    table_layout_cfg = cfg.get("table_layout", {}) or {}
    algorithm_control_cfg = get_algorithm_control(cfg)
    deleted_outputs = []
    cleanup_failures = []
    accuracy_source_override = None
    if len(positional_args) >= 3:
        try:
            accuracy_source_override = parse_accuracy_source_arg(positional_args[2])
        except ValueError as exc:
            print("错误: %s，可用值: raw 或 audited" % exc)
            return 1
        if accuracy_source_override is not None:
            apply_accuracy_source_override(cfg, accuracy_source_override)
    accuracy_source = accuracy_source_override or configured_accuracy_source(cfg)
    accuracy_field = "准确率" if accuracy_source == "raw" else "审核后准确率"
    accuracy_field_label = accuracy_field
    pure_image_clipboard = bool(clipboard_cfg.get("pure_image", False))
    if len(positional_args) >= 2:
        try:
            pure_image_override = parse_clipboard_mode_arg(positional_args[1])
        except ValueError as exc:
            print("错误: %s，可用值: text 或 image" % exc)
            return 1
        if pure_image_override is not None:
            pure_image_clipboard = pure_image_override
    generated_table_paths = []

    for label, path in (("旧数据文件", old_file), ("新数据文件", new_file)):
        if not os.path.isfile(path):
            print("错误: %s不存在，请检查config.yaml中的data_source配置: %s" % (label, path))
            return 1
    if using_lan_upload.get("old"):
        print("[局域网临时表] 网页 old.xlsx 本次作为旧数据: %s" % old_file)
    elif using_lan_upload.get("result"):
        print("[局域网临时表] 网页未暂存 old.xlsx，本机 result-2.xlsx 本次作为旧数据: %s" % old_file)
    if using_lan_upload.get("result"):
        print("[局域网临时表] 网页 result-2.xlsx 本次作为新数据: %s" % new_file)

    old_rows = load_site_rows(old_file, keyword)
    new_rows = load_site_rows(new_file, keyword)

    if not new_rows:
        print("错误: 在 %s 中未找到包含 '%s' 的站点" % (new_file, keyword))
        return 1

    deleted_outputs, cleanup_failures = cleanup_old_outputs(output_dir)

    all_new_rows = new_rows
    all_old_rows = old_rows
    controlled_old_rows = filter_algorithm_rows(
        all_old_rows,
        algorithm_control_cfg,
        for_calculation=True,
    )
    controlled_new_rows = filter_algorithm_rows(
        all_new_rows,
        algorithm_control_cfg,
        for_calculation=True,
    )
    no_data_rows = [
        r for r in controlled_new_rows
        if has_zero_points(r.get("点位数量"))
    ]
    no_data_names = {
        str(r.get("算法小类") or "").strip()
        for r in no_data_rows
        if str(r.get("算法小类") or "").strip()
    }
    old_rows = [
        r for r in controlled_old_rows
        if str(r.get("算法小类") or "").strip() not in no_data_names
        and not has_zero_points(r.get("点位数量"))
    ]
    new_rows = [
        r for r in controlled_new_rows
        if str(r.get("算法小类") or "").strip() not in no_data_names
        and not has_zero_points(r.get("点位数量"))
    ]

    # 站点和项目名称
    site_name = clean_site_name(all_new_rows[0]["站点"])
    project_names = []
    for row in all_new_rows:
        project_name = str(row.get("项目名称") or "").strip()
        if project_name and project_name not in project_names:
            project_names.append(project_name)
    project_name_text = "、".join(project_names) if project_names else "-"
    today = date.today()
    report_stem = os.path.join(
        output_dir, "周报_%s_%s" % (site_name, today.strftime("%Y%m%d"))
    )

    # 按算法小类分组
    old_map = {}
    for r in old_rows:
        old_map[r["算法小类"]] = r
    new_map = {}
    for r in new_rows:
        new_map[r["算法小类"]] = r

    # ========== 数据统计 ==========
    total_points = sum(parse_int(r["点位数量"]) for r in new_rows)
    algo_count = len(new_map)

    silent_substrings = get_silent_substrings(cfg)
    silent_algos = [
        r["算法小类"]
        for r in new_rows
        if is_silent_algo(r["算法小类"], silent_substrings)
    ]
    silent_text = "有（" + "、".join(silent_algos) + "）" if silent_algos else "无"

    # 整体准确率和明细表只纳入点位数量为有效正数的算法。
    accuracy_calc_rows = [
        r for r in new_rows
        if has_effective_points(r.get("点位数量"))
    ]
    accuracy_display_rows = filter_algorithm_rows(
        accuracy_calc_rows,
        algorithm_control_cfg,
        table_key="accuracy_detail",
    )

    # 准确率按巡视（非静默）和静默分别计算，避免两类算法混合后失真。
    patrol_rows = [
        r for r in accuracy_calc_rows
        if not is_silent_algo(r["算法小类"], silent_substrings)
    ]
    silent_rows = [
        r for r in accuracy_calc_rows
        if is_silent_algo(r["算法小类"], silent_substrings)
    ]
    position_accuracy_field = get_position_accuracy_selection(cfg)
    position_accuracy = calculate_position_accuracy(
        accuracy_calc_rows,
        position_accuracy_field,
    )
    accuracy_groups = [
        (
            "巡视",
            patrol_rows,
            average_percent(patrol_rows, accuracy_field),
        ),
        (
            "静默",
            silent_rows,
            average_percent(silent_rows, accuracy_field),
        ),
    ]

    # 准确率低于阈值
    low_acc = []
    for r in new_rows:
        if has_effective_points(r.get("点位数量")) and is_below_threshold(
            r, accuracy_field, cfg
        ):
            val = get_priority_value(r, accuracy_field, cfg)
            low_acc.append((r["算法小类"], val))

    # 误报率高于阈值
    high_fp = []
    for r in new_rows:
        if is_above_threshold(r, "误报率", cfg) or is_above_threshold(r, "审核后误报率", cfg):
            val = get_priority_value(r, "误报率", cfg)
            high_fp.append((r["算法小类"], val))

    # 漏报率高于阈值
    high_leak = []
    for r in new_rows:
        if is_above_threshold(r, "漏报率", cfg) or is_above_threshold(r, "审核后漏报率", cfg):
            val = get_priority_value(r, "漏报率", cfg)
            high_leak.append((r["算法小类"], val))

    # 拍照模糊（具体到算法）
    blur_items = []
    for r in new_rows:
        blur_cnt = parse_int(r["拍照模糊数量"])
        blur_aud_cnt = parse_int(r["审核后拍照模糊数量"])
        if blur_cnt > 0 or blur_aud_cnt > 0:
            blur_items.append((r["算法小类"], blur_cnt, blur_aud_cnt))
    blur_total = sum(b[1] for b in blur_items)

    # 下一步计划只关注有有效点位且当前选定准确率明确为 0% 的算法。
    idle_algos = get_zero_accuracy_algorithms(
        new_rows,
        accuracy_field,
        silent_substrings,
    )

    # 更新日期久远
    stale_days = thresholds.get("stale_days", 15)
    stale_algos = []
    for r in new_rows:
        d = parse_task_date(r["任务日期"])
        if d is not None and (today - d).days > stale_days:
            stale_algos.append((r["算法小类"], d, parse_int(r["点位数量"])))

    # 点位数量为空或为 "-" 的算法单独列为数据风险，不与有效点位统计混淆。
    empty_point_algos = []
    for r in new_rows:
        if is_blank_or_dash(r.get("点位数量")):
            empty_point_algos.append((r["算法小类"], r.get("点位数量")))

    # 点位数量明确为 0 的算法只进入算法无数据风险，不参与其他统计和表格。
    no_data_algos = [
        (
            r["算法小类"],
            get_algorithm_build_date(r),
            parse_int(r.get("点位数量")),
        )
        for r in no_data_rows
    ]

    low_acc_display = filter_algorithm_items(
        low_acc, algorithm_control_cfg, "low_accuracy"
    )
    high_fp_display = filter_algorithm_items(
        high_fp, algorithm_control_cfg, "high_false_positive"
    )
    blur_items_display = filter_algorithm_items(
        blur_items, algorithm_control_cfg, "blur"
    )
    blur_total_display = sum(item[1] for item in blur_items_display)
    blur_risk_items_display = filter_algorithm_items(
        blur_items, algorithm_control_cfg, "risks"
    )
    stale_algos_display = filter_algorithm_items(
        stale_algos, algorithm_control_cfg, "risks"
    )
    idle_algos_display = filter_algorithm_items(
        idle_algos, algorithm_control_cfg, "risks"
    )
    empty_point_algos_display = filter_algorithm_items(
        empty_point_algos, algorithm_control_cfg, "risks"
    )
    no_data_algos_display = filter_algorithm_items(
        no_data_algos, algorithm_control_cfg, "risks"
    )
    site_attainment = parse_percent(all_new_rows[0].get("达标率"))

    # ========== 三、对比新老数据 ==========
    added = [k for k in new_map if k not in old_map]
    removed = [k for k in old_map if k not in new_map]

    compare_fields = list(changes_cfg.get("compare_fields", []))
    if accuracy_source_override is not None:
        compare_fields = [
            accuracy_field if field in ("准确率", "审核后准确率") else field
            for field in compare_fields
        ]
    metric_mode = str(changes_cfg.get("metric_mode", "accuracy_trend")).strip().lower()
    changed = []
    for k in new_map:
        if k in old_map:
            o, n = old_map[k], new_map[k]
            diffs = []
            for field in compare_fields:
                ov, nv = o.get(field), n.get(field)
                old_display = format_compare_value(field, ov)
                new_display = format_compare_value(field, nv)
                if old_display != new_display:
                    diffs.append((field, old_display, new_display))
            if diffs:
                changed.append((k, diffs))
    changed_display = filter_algorithm_items(
        changed, algorithm_control_cfg, "changes"
    )
    accuracy_trend_assets = None
    if changes_cfg.get("show_metric_changes", True) and metric_mode == "accuracy_trend":
        accuracy_trend_assets = create_accuracy_trend_table_assets(
            report_stem,
            old_map,
            new_map,
            accuracy_field,
            today,
            changes_cfg,
            table_layout_cfg,
        )

    # ========== 四、已处理（对比上周已解决的事项） ==========
    # 根据 yaml 关键字配置识别待办类型
    keywords_cfg = get_field_cfg(cfg, "下一版本当前阶段").get("keywords", {})
    todo_log = []      # 本周仍"填写问题日志"未完成（供下一步计划使用）
    todo_upload = []   # 本周仍"准确率上传"未完成（供下一步计划使用）
    for a, r in new_map.items():
        stage = str(r.get("下一版本当前阶段") or "")
        for kw, todo_type in keywords_cfg.items():
            if kw in stage:
                if todo_type == "problem_log":
                    todo_log.append(a)
                elif todo_type == "upload_pending":
                    todo_upload.append(a)

    # 1、准确率已达标：上周<95%，本周>=95%
    handled_acc = []
    for a in new_map:
        if a in old_map:
            o, n = old_map[a], new_map[a]
            old_val = get_priority_value(o, accuracy_field, cfg)
            new_val = get_priority_value(n, accuracy_field, cfg)
            if old_val is not None and old_val < thresholds.get("accuracy", 95) and \
               new_val is not None and new_val >= thresholds.get("accuracy", 95):
                handled_acc.append((a, old_val, new_val))

    # 2、误报率已降低：上周>5%，本周<=5%
    handled_fp = []
    for a in new_map:
        if a in old_map:
            o, n = old_map[a], new_map[a]
            old_val = get_priority_value(o, "误报率", cfg)
            new_val = get_priority_value(n, "误报率", cfg)
            if old_val is not None and old_val > thresholds.get("false_positive", 5) and \
               new_val is not None and new_val <= thresholds.get("false_positive", 5):
                handled_fp.append((a, old_val, new_val))

    # 3、漏报率已降低：上周>5%，本周<=5%
    handled_leak = []
    for a in new_map:
        if a in old_map:
            o, n = old_map[a], new_map[a]
            old_val = get_priority_value(o, "漏报率", cfg)
            new_val = get_priority_value(n, "漏报率", cfg)
            if old_val is not None and old_val > thresholds.get("false_positive", 5) and \
               new_val is not None and new_val <= thresholds.get("false_positive", 5):
                handled_leak.append((a, old_val, new_val))

    # 4、问题日志已填写：上周未完成，本周已完成
    handled_log = []
    for a in new_map:
        if a in old_map:
            old_stage = str(old_map[a].get("下一版本当前阶段") or "")
            new_stage = str(new_map[a].get("下一版本当前阶段") or "")
            if "填写问题日志" in old_stage and "填写问题日志" not in new_stage:
                handled_log.append(a)

    # 5、准确率已上传：上周未完成，本周已完成
    handled_upload = []
    for a in new_map:
        if a in old_map:
            old_stage = str(old_map[a].get("下一版本当前阶段") or "")
            new_stage = str(new_map[a].get("下一版本当前阶段") or "")
            if "准确率上传" in old_stage and "准确率上传" not in new_stage:
                handled_upload.append(a)

    # 6、任务日期已更新：上周久远，本周已更新
    handled_date = []
    for a in new_map:
        if a in old_map:
            old_d = parse_task_date(old_map[a].get("任务日期"))
            new_d = parse_task_date(new_map[a].get("任务日期"))
            if old_d is not None and (today - old_d).days > stale_days and \
               new_d is not None and (today - new_d).days <= stale_days:
                handled_date.append((a, old_d, new_d))

    handled_acc_display = filter_algorithm_items(
        handled_acc, algorithm_control_cfg, "handled"
    )
    handled_fp_display = filter_algorithm_items(
        handled_fp, algorithm_control_cfg, "handled"
    )
    handled_leak_display = filter_algorithm_items(
        handled_leak, algorithm_control_cfg, "handled"
    )
    handled_log_display = filter_algorithm_items(
        handled_log, algorithm_control_cfg, "handled"
    )
    handled_upload_display = filter_algorithm_items(
        handled_upload, algorithm_control_cfg, "handled"
    )
    handled_date_display = filter_algorithm_items(
        handled_date, algorithm_control_cfg, "handled"
    )
    report_context = {
        "today": today.strftime("%Y-%m-%d"),
        "stale_days": stale_days,
        "deadline": (today + timedelta(days=7)).strftime("%Y-%m-%d"),
        "low_acc_names": "、".join(name for name, _ in low_acc_display),
        "high_fp_names": "、".join(name for name, _ in high_fp_display),
        "idle_names": "、".join(idle_algos_display),
        "blur_names": "、".join(name for name, _, _ in blur_risk_items_display),
        "stale_names": "、".join(name for name, _, _ in stale_algos_display),
        "empty_point_names": "、".join(
            name for name, _ in empty_point_algos_display
        ),
        "no_data_names": "、".join(
            name for name, _, _ in no_data_algos_display
        ),
        "todo_log_names": "、".join(todo_log),
        "todo_upload_names": "、".join(todo_upload),
        "has_low_accuracy": bool(low_acc_display),
        "has_high_fp": bool(high_fp_display),
        "has_idle": bool(idle_algos_display),
        "has_blur": bool(blur_risk_items_display),
        "has_stale": bool(stale_algos_display),
        "has_empty_point": bool(empty_point_algos_display),
        "has_no_data": bool(no_data_algos_display),
        "has_todo_log": bool(todo_log),
        "has_todo_upload": bool(todo_upload),
        "low_acc_count": len(low_acc_display),
        "high_fp_count": len(high_fp_display),
        "idle_count": len(idle_algos_display),
        "blur_count": len(blur_risk_items_display),
        "stale_count": len(stale_algos_display),
        "empty_point_count": len(empty_point_algos_display),
        "no_data_count": len(no_data_algos_display),
        "todo_log_count": len(todo_log),
        "todo_upload_count": len(todo_upload),
    }

    # ========== 生成周报 ==========
    lines = []
    lines.append("【%s】周报" % today.strftime("%Y-%m-%d"))
    lines.append("")
    section_idx = 0

    # 一、站点
    if sections.get("site", True):
        lines.append("%s、站点：%s" % (CN_NUMS[section_idx], site_name))
        lines.append("项目名称：%s" % project_name_text)
        section_idx += 1
        lines.append("")

    # 二、算法运行实况
    if sections.get("overview", True):
        lines.append("%s、算法运行实况" % CN_NUMS[section_idx])
        section_idx += 1

        if overview_cfg.get("show_total_points", True):
            total_parts = [
                "本周算法点位 %d 个，涉及算法小类 %d 个" % (total_points, algo_count)
            ]
            if silent_algos:
                total_parts.append("静默任务有（%s）" % "、".join(silent_algos))
            lines.append("1、%s。" % "；".join(total_parts))

        if overview_cfg.get("show_avg_accuracy", True):
            accuracy_texts = []
            accuracy_suffix = "准确率" if accuracy_source == "raw" else "审核后准确率"
            for group_name, group_rows, avg_acc in accuracy_groups:
                if not group_rows:
                    continue
                parts = []
                if avg_acc is not None:
                    if group_name == "巡视":
                        parts.append(
                            "整体平均巡视%s %s"
                            % (accuracy_suffix, fmt_percent(avg_acc))
                        )
                    elif not patrol_rows:
                        parts.append(
                            "整体平均静默%s %s"
                            % (accuracy_suffix, fmt_percent(avg_acc))
                        )
                    else:
                        parts.append(
                            "静默%s %s" % (accuracy_suffix, fmt_percent(avg_acc))
                        )
                if parts:
                    accuracy_texts.append("，".join(parts))

            accuracy_summary = []
            if position_accuracy is not None:
                accuracy_summary.append(
                    "本站点当前算法点位准确率 %s" % fmt_percent(position_accuracy)
                )
            if patrol_rows:
                patrol_accuracy = average_percent(patrol_rows, accuracy_field)
                if patrol_accuracy is not None:
                    accuracy_summary.append(
                        "巡视算法类别达标率 %s" % fmt_percent(patrol_accuracy)
                    )
            if silent_rows:
                silent_accuracy = average_percent(silent_rows, accuracy_field)
                if silent_accuracy is not None:
                    accuracy_summary.append(
                        "静默算法类别达标率 %s" % fmt_percent(silent_accuracy)
                    )
            if accuracy_summary:
                lines.append("2、%s。" % "，".join(accuracy_summary))
                if overview_cfg.get("show_accuracy_details", True):
                    accuracy_detail_display = overview_cfg.get(
                        "accuracy_detail_display", "audited"
                    )
                    table_assets = create_accuracy_table_assets(
                        report_stem,
                        accuracy_display_rows,
                        accuracy_detail_display,
                        overview_cfg,
                        accuracy_source,
                        table_layout_cfg,
                        silent_substrings,
                    )
                    if table_assets:
                        generated_table_paths.append(table_assets["xlsx_path"])
                        if table_assets["png_path"]:
                            generated_table_paths.append(table_assets["png_path"])
                        if table_assets["png_path"]:
                            append_table_reference(
                                lines, table_assets["png_path"], "整体准确率明细"
                            )
            elif overview_cfg.get("show_empty_placeholder", False):
                lines.append("2、整体平均准确率：无数据。")

        item_no = 3
        if overview_cfg.get("show_low_accuracy", True):
            if low_acc_display:
                lines.append("%d、准确率低于 %.0f%% 的算法：" % (item_no, thresholds.get("accuracy", 95)))
                low_acc_table = create_sublist_table_asset(
                    report_stem,
                    "准确率低于阈值",
                    ["算法小类", accuracy_field_label],
                    [[name, fmt_percent(val)] for name, val in low_acc_display],
                    table_layout_cfg,
                )
                if low_acc_table:
                    generated_table_paths.append(low_acc_table)
                    append_table_reference(lines, low_acc_table, "准确率低于阈值")
                else:
                    for name, val in low_acc_display:
                        lines.append("   - %s：%s" % (name, fmt_percent(val)))
                item_no += 1
            elif overview_cfg.get("show_empty_placeholder", False):
                lines.append("%d、准确率低于 %.0f%% 的算法：无" % (item_no, thresholds.get("accuracy", 95)))
                item_no += 1

        if overview_cfg.get("show_high_false_positive", True):
            if high_fp_display:
                # 根据字段配置决定前缀（误报率/审核后误报率）
                fp_field_cfg = get_field_cfg(cfg, "误报率")
                fp_display = fp_field_cfg.get("display", "raw")
                fp_label = "审核后误报率" if fp_display == "audited" else "误报率"
                lines.append("%d、误报率高于 %.0f%% 的算法：" % (item_no, thresholds.get("false_positive", 5)))
                high_fp_table = create_sublist_table_asset(
                    report_stem,
                    "误报率高于阈值",
                    ["算法小类", table_field_label(fp_label)],
                    [[name, fmt_percent(val)] for name, val in high_fp_display],
                    table_layout_cfg,
                )
                if high_fp_table:
                    generated_table_paths.append(high_fp_table)
                    append_table_reference(lines, high_fp_table, "误报率高于阈值")
                else:
                    for name, val in high_fp_display:
                        lines.append("   - %s：%s %s" % (name, fp_label, fmt_percent(val)))
                item_no += 1
            elif overview_cfg.get("show_empty_placeholder", False):
                lines.append("%d、误报率高于 %.0f%% 的算法：无" % (item_no, thresholds.get("false_positive", 5)))
                item_no += 1

        if overview_cfg.get("show_blur", True):
            if blur_items_display:
                lines.append("%d、拍照模糊点位：共 %d 个，涉及 %d 个算法：" % (
                    item_no, blur_total_display, len(blur_items_display)
                ))
                blur_table = create_sublist_table_asset(
                    report_stem,
                    "拍照模糊点位",
                    ["算法小类", "拍照模糊数量", "审核后拍照模糊数量"],
                    [
                        [name, cnt, aud_cnt]
                        for name, cnt, aud_cnt in blur_items_display
                    ],
                    table_layout_cfg,
                )
                if blur_table:
                    generated_table_paths.append(blur_table)
                    append_table_reference(lines, blur_table, "拍照模糊点位")
                else:
                    for name, cnt, aud_cnt in blur_items_display:
                        lines.append("   - %s：%d 个点位模糊（审核后仍模糊 %d 个）" % (name, cnt, aud_cnt))
                item_no += 1
            elif overview_cfg.get("show_empty_placeholder", False):
                lines.append("%d、拍照模糊点位：无。" % item_no)
                item_no += 1

        if overview_cfg.get("show_idle", True):
            if idle_algos:
                lines.append("%d、本周无触发/长期挂起算法：%s" % (item_no, "、".join(idle_algos)))
                item_no += 1
            elif overview_cfg.get("show_empty_placeholder", False):
                lines.append("%d、本周无触发/长期挂起算法：无" % item_no)
                item_no += 1

        lines.append("")

    # 三、本周变化
    if sections.get("changes", True):
        # 判断是否有内容需要输出
        has_changes_content = False
        if changes_cfg.get("show_added", True) and (added or changes_cfg.get("show_empty_placeholder", False)):
            has_changes_content = True
        if changes_cfg.get("show_removed", True) and (removed or changes_cfg.get("show_empty_placeholder", False)):
            has_changes_content = True
        if changes_cfg.get("show_metric_changes", True):
            if metric_mode == "accuracy_trend":
                if accuracy_trend_assets or changes_cfg.get("show_empty_placeholder", False):
                    has_changes_content = True
            elif changed_display or changes_cfg.get("show_empty_placeholder", False):
                has_changes_content = True

        if has_changes_content:
            lines.append("%s、本周变化（对比上周）" % CN_NUMS[section_idx])
            section_idx += 1
            item_no = 1

            if changes_cfg.get("show_added", True):
                if added:
                    lines.append("%d、新增算法：%s" % (item_no, "、".join(added)))
                    item_no += 1
                elif changes_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、新增算法：无" % item_no)
                    item_no += 1

            if changes_cfg.get("show_removed", True):
                if removed:
                    lines.append("%d、移除算法：%s" % (item_no, "、".join(removed)))
                    item_no += 1
                elif changes_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、移除算法：无" % item_no)
                    item_no += 1

            if changes_cfg.get("show_metric_changes", True):
                if metric_mode == "accuracy_trend":
                    if accuracy_trend_assets:
                        lines.append("%d、算法准确率变化：" % item_no)
                        generated_table_paths.append(accuracy_trend_assets["xlsx_path"])
                        if accuracy_trend_assets.get("png_path"):
                            generated_table_paths.append(accuracy_trend_assets["png_path"])
                            append_table_reference(lines, accuracy_trend_assets["png_path"], "算法准确率变化")
                        item_no += 1
                    elif changes_cfg.get("show_empty_placeholder", False):
                        lines.append("%d、算法准确率变化：无" % item_no)
                        item_no += 1
                elif changed_display:
                    lines.append("%d、指标变化：" % item_no)
                    show_dash_for_unchanged = bool(
                        changes_cfg.get("show_dash_for_unchanged", True)
                    )
                    change_fields = list(compare_fields)
                    if not change_fields:
                        change_fields = [
                            field
                            for _, diffs in changed_display
                            for field, _, _ in diffs
                        ]
                        change_fields = list(dict.fromkeys(change_fields))
                    active_change_fields = []
                    for field in change_fields:
                        if any(
                            field in {diff_field for diff_field, _, _ in diffs}
                            for _, diffs in changed_display
                        ):
                            active_change_fields.append(field)
                    change_fields = active_change_fields
                    change_headers = ["算法小类"] + [
                        table_field_label(field) for field in change_fields
                    ]
                    change_rows = []
                    for name, diffs in changed_display:
                        diff_map = {
                            field: "%s → %s" % (old_value, new_value)
                            for field, old_value, new_value in diffs
                        }
                        current_row = new_map.get(name, {})
                        change_cells = []
                        for field in change_fields:
                            if field in diff_map:
                                change_cells.append(diff_map[field])
                            elif show_dash_for_unchanged:
                                change_cells.append("-")
                            else:
                                change_cells.append(
                                    format_unchanged_compare_value(field, current_row)
                                )
                        change_rows.append([name] + change_cells)
                    changes_table = create_sublist_table_asset(
                        report_stem,
                        "指标变化",
                        change_headers,
                        change_rows,
                        table_layout_cfg,
                    )
                    if changes_table:
                        generated_table_paths.append(changes_table)
                        append_table_reference(lines, changes_table, "指标变化")
                    else:
                        for name, diffs in changed:
                            lines.append("   - %s：" % name)
                            for index, (field, old_value, new_value) in enumerate(diffs):
                                suffix = "；" if index < len(diffs) - 1 else ""
                                lines.append("      %s：%s → %s%s" % (
                                    field, old_value, new_value, suffix
                                ))
                    item_no += 1
                elif changes_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、指标变化：无" % item_no)
                    item_no += 1

            lines.append("")

    # 四、已处理（对比上周已解决的事项）
    if sections.get("handled", False):
        has_handled_content = False
        if handled_cfg.get("show_accuracy_improved", True) and (handled_acc_display or handled_cfg.get("show_empty_placeholder", False)):
            has_handled_content = True
        if handled_cfg.get("show_fp_reduced", True) and (handled_fp_display or handled_cfg.get("show_empty_placeholder", False)):
            has_handled_content = True
        if handled_cfg.get("show_leak_reduced", True) and (handled_leak_display or handled_cfg.get("show_empty_placeholder", False)):
            has_handled_content = True
        if handled_cfg.get("show_problem_log_done", True) and (handled_log_display or handled_cfg.get("show_empty_placeholder", False)):
            has_handled_content = True
        if handled_cfg.get("show_upload_done", True) and (handled_upload_display or handled_cfg.get("show_empty_placeholder", False)):
            has_handled_content = True
        if handled_cfg.get("show_date_updated", True) and (handled_date_display or handled_cfg.get("show_empty_placeholder", False)):
            has_handled_content = True

        if has_handled_content:
            lines.append("%s、已处理（对比上周）" % CN_NUMS[section_idx])
            section_idx += 1
            item_no = 1

            if handled_cfg.get("show_accuracy_improved", True):
                if handled_acc_display:
                    lines.append("%d、准确率已达标：" % item_no)
                    handled_acc_table = create_sublist_table_asset(
                        report_stem,
                        "准确率已达标",
                        ["算法小类", accuracy_field_label],
                        [
                            [a, "%s → %s" % (fmt_percent(old_val), fmt_percent(new_val))]
                            for a, old_val, new_val in handled_acc_display
                        ],
                        table_layout_cfg,
                    )
                    if handled_acc_table:
                        generated_table_paths.append(handled_acc_table)
                        append_table_reference(lines, handled_acc_table, "准确率已达标")
                    else:
                        for a, old_val, new_val in handled_acc_display:
                            lines.append("   - %s：%s → %s" % (a, fmt_percent(old_val), fmt_percent(new_val)))
                    item_no += 1
                elif handled_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、准确率已达标：无" % item_no)
                    item_no += 1

            if handled_cfg.get("show_fp_reduced", True):
                if handled_fp_display:
                    lines.append("%d、误报率已降低：" % item_no)
                    handled_fp_table = create_sublist_table_asset(
                        report_stem,
                        "误报率已降低",
                        ["算法小类", "误报率"],
                        [
                            [a, "%s → %s" % (fmt_percent(old_val), fmt_percent(new_val))]
                            for a, old_val, new_val in handled_fp_display
                        ],
                        table_layout_cfg,
                    )
                    if handled_fp_table:
                        generated_table_paths.append(handled_fp_table)
                        append_table_reference(lines, handled_fp_table, "误报率已降低")
                    else:
                        for a, old_val, new_val in handled_fp_display:
                            lines.append("   - %s：%s → %s" % (a, fmt_percent(old_val), fmt_percent(new_val)))
                    item_no += 1
                elif handled_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、误报率已降低：无" % item_no)
                    item_no += 1

            if handled_cfg.get("show_leak_reduced", True):
                if handled_leak_display:
                    lines.append("%d、漏报率已降低：" % item_no)
                    handled_leak_table = create_sublist_table_asset(
                        report_stem,
                        "漏报率已降低",
                        ["算法小类", "漏报率"],
                        [
                            [a, "%s → %s" % (fmt_percent(old_val), fmt_percent(new_val))]
                            for a, old_val, new_val in handled_leak_display
                        ],
                        table_layout_cfg,
                    )
                    if handled_leak_table:
                        generated_table_paths.append(handled_leak_table)
                        append_table_reference(lines, handled_leak_table, "漏报率已降低")
                    else:
                        for a, old_val, new_val in handled_leak_display:
                            lines.append("   - %s：%s → %s" % (a, fmt_percent(old_val), fmt_percent(new_val)))
                    item_no += 1
                elif handled_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、漏报率已降低：无" % item_no)
                    item_no += 1

            if handled_cfg.get("show_problem_log_done", True):
                if handled_log_display:
                    lines.append("%d、问题日志已填写完成：%s" % (item_no, "、".join(handled_log_display)))
                    item_no += 1
                elif handled_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、问题日志已填写完成：无" % item_no)
                    item_no += 1

            if handled_cfg.get("show_upload_done", True):
                if handled_upload_display:
                    lines.append("%d、准确率已上传：%s" % (item_no, "、".join(handled_upload_display)))
                    item_no += 1
                elif handled_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、准确率已上传：无" % item_no)
                    item_no += 1

            if handled_cfg.get("show_date_updated", True):
                if handled_date_display:
                    lines.append("%d、任务日期已更新：" % item_no)
                    handled_date_table = create_sublist_table_asset(
                        report_stem,
                        "任务日期已更新",
                        ["算法小类", "任务日期"],
                        [
                            [a, "%s → %s" % (fmt_date(old_d), fmt_date(new_d))]
                            for a, old_d, new_d in handled_date_display
                        ],
                        table_layout_cfg,
                    )
                    if handled_date_table:
                        generated_table_paths.append(handled_date_table)
                        append_table_reference(lines, handled_date_table, "任务日期已更新")
                    else:
                        for a, old_d, new_d in handled_date_display:
                            lines.append("   - %s：%s → %s" % (a, fmt_date(old_d), fmt_date(new_d)))
                    item_no += 1
                elif handled_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、任务日期已更新：无" % item_no)
                    item_no += 1

            lines.append("")

    # 五、问题及风险
    # 下一步计划先暂存，待问题及风险输出后再追加，以保持报告顺序。
    next_steps_lines = []
    if sections.get("next_steps", True):
        lines_for_next_steps = lines
        lines = next_steps_lines
        next_section_idx = section_idx + (
            1 if sections.get("risks", True) else 0
        )
        lines.append("%s、下一步计划" % CN_NUMS[next_section_idx])
        item_no = 1

        next_steps_templates = next_steps_cfg.get("templates")
        if isinstance(next_steps_templates, dict) and next_steps_templates:
            if next_steps_cfg.get("show_fp_optimization", True):
                template = next_steps_templates.get("fp_optimization")
                if high_fp and template:
                    lines.append("%d、%s" % (
                        item_no,
                        render_template_text(template, report_context),
                    ))
                    item_no += 1
                elif high_fp:
                    fp_list = "、".join(a for a, _ in high_fp)
                    lines.append("%d、针对误报率较高的算法（%s）继续负样本采集与模型优化。" % (item_no, fp_list))
                    item_no += 1
                elif next_steps_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、本周无误报率较高的算法，持续监控。" % item_no)
                    item_no += 1

            if next_steps_cfg.get("show_problem_log_deadline", False):
                template = next_steps_templates.get("problem_log_deadline")
                if todo_log and template:
                    lines.append("%d、%s" % (
                        item_no,
                        render_template_text(template, report_context),
                    ))
                    item_no += 1
                elif todo_log:
                    deadline = today + timedelta(days=7)
                    lines.append("%d、推动工程组在 %s 前完成问题日志填写。" % (item_no, deadline.strftime("%Y-%m-%d")))
                    item_no += 1

            if next_steps_cfg.get("show_idle_evaluation", True):
                template = next_steps_templates.get("idle_evaluation")
                if idle_algos and template:
                    lines.append("%d、%s" % (
                        item_no,
                        render_template_text(template, report_context),
                    ))
                    item_no += 1
                elif idle_algos:
                    lines.append("%d、对正在使用但目前准确率为0%的算法（%s）评估继续投入或替换模型。" % (item_no, "、".join(idle_algos)))
                    item_no += 1
                elif next_steps_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、无正在使用但目前准确率为0%的算法。" % item_no)
                    item_no += 1
        elif isinstance(next_steps_templates, list):
            item_no = append_configured_texts(
                lines, next_steps_cfg, report_context, item_no
            )
            if item_no == 1 and next_steps_cfg.get("show_empty_placeholder", False):
                lines.append("%d、无。" % item_no)
                item_no += 1
        else:
            if next_steps_cfg.get("show_fp_optimization", True):
                if high_fp:
                    fp_list = "、".join(a for a, _ in high_fp)
                    lines.append("%d、针对误报率较高的算法（%s）继续负样本采集与模型优化。" % (item_no, fp_list))
                    item_no += 1
                elif next_steps_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、本周无误报率较高的算法，持续监控。" % item_no)
                    item_no += 1

            if next_steps_cfg.get("show_problem_log_deadline", False):
                if todo_log:
                    deadline = today + timedelta(days=7)
                    lines.append("%d、推动工程组在 %s 前完成问题日志填写。" % (item_no, deadline.strftime("%Y-%m-%d")))
                    item_no += 1

            if next_steps_cfg.get("show_idle_evaluation", True):
                if idle_algos:
                    lines.append("%d、对正在使用但目前准确率为0%的算法（%s）评估继续投入或替换模型。" % (item_no, "、".join(idle_algos)))
                    item_no += 1
                elif next_steps_cfg.get("show_empty_placeholder", False):
                    lines.append("%d、无正在使用但目前准确率为0%的算法。" % item_no)
                    item_no += 1

        lines.append("")
        lines = lines_for_next_steps

    # 问题及风险
    if sections.get("risks", True):
        lines.append("%s、问题及风险" % CN_NUMS[section_idx])
        section_idx += 1
        detailed = risks_cfg.get("detailed", True)
        risk_no = 1
        risk_templates = risks_cfg.get("templates", {})
        if not isinstance(risk_templates, dict):
            risk_templates = {}

        if risks_cfg.get("show_blur_risk", True) and blur_risk_items_display:
            if detailed:
                blur_template = risk_templates.get("blur", {}) or {}
                blur_intro = blur_template.get(
                    "detailed",
                    "图片异常风险：存在不满足算法要求的点位图片(模糊、黑白、遮掩等)，进而影响识别准确率：",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(blur_intro, report_context),
                ))
                blur_risk_table = create_sublist_table_asset(
                    report_stem,
                    "图片异常风险",
                    ["算法小类", "图片异常数量"],
                    [
                        [a, blur_cnt, blur_aud_cnt]
                        for a, blur_cnt, blur_aud_cnt in blur_risk_items_display
                    ],
                    table_layout_cfg,
                )
                if blur_risk_table:
                    generated_table_paths.append(blur_risk_table)
                    append_table_reference(lines, blur_risk_table, "模糊风险")
                else:
                    for a, blur_cnt, blur_aud_cnt in blur_risk_items_display:
                        lines.append(
                            "   - %s（拍照模糊 %d 个，审核后仍模糊 %d 个）"
                                % (a, blur_cnt, blur_aud_cnt)
                        )
            else:
                blur_template = risk_templates.get("blur", {}) or {}
                blur_summary = blur_template.get(
                    "summary",
                    "部分算法存在拍照模糊点位，可能影响识别准确率，需检查拍摄环境及设备状态。",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(blur_summary, report_context),
                ))
            risk_no += 1

        if risks_cfg.get("show_point_empty_risk", True) and empty_point_algos_display:
            if detailed:
                empty_point_template = risk_templates.get("point_empty", {}) or {}
                empty_point_intro = empty_point_template.get(
                    "detailed",
                    "点位数据风险：算法点位数量为空，影响准确率统计，将评估继续投入或就删除算法展开讨论：",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(empty_point_intro, report_context),
                ))
                empty_point_table = create_sublist_table_asset(
                    report_stem,
                    "点位数量为空风险",
                    ["算法小类", "点位数量"],
                    [
                        [
                            name,
                            0 if is_blank_or_dash(point_count)
                            else parse_int(point_count),
                        ]
                        for name, point_count in empty_point_algos_display
                    ],
                    table_layout_cfg,
                    compact_columns=False,
                )
                if empty_point_table:
                    generated_table_paths.append(empty_point_table)
                    append_table_reference(
                        lines, empty_point_table, "点位数量为空风险"
                    )
                else:
                    for name, point_count in empty_point_algos_display:
                        lines.append(
                            "   - %s：点位数量 %s"
                            % (
                                name,
                                0 if is_blank_or_dash(point_count)
                                else parse_int(point_count),
                            )
                        )
            else:
                empty_point_template = risk_templates.get("point_empty", {}) or {}
                empty_point_summary = empty_point_template.get(
                    "summary",
                    "部分算法点位数量为空或为“-”，暂无法纳入准确率统计，需补充点位数据。",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(empty_point_summary, report_context),
                ))
            risk_no += 1

        if risks_cfg.get("show_no_data_risk", True) and no_data_algos_display:
            if detailed:
                no_data_template = risk_templates.get("no_data", {}) or {}
                no_data_intro = no_data_template.get(
                    "detailed",
                    "算法无数据风险：以下算法点位数量为 0，暂不纳入本周统计：",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(no_data_intro, report_context),
                ))
                no_data_table = create_sublist_table_asset(
                    report_stem,
                    "算法无数据风险",
                    ["算法小类", "建立日期", "点位数量"],
                    [
                        [name, fmt_table_date(build_date), point_count]
                        for name, build_date, point_count in no_data_algos_display
                    ],
                    table_layout_cfg,
                    compact_columns=False,
                )
                if no_data_table:
                    generated_table_paths.append(no_data_table)
                    append_table_reference(lines, no_data_table, "算法无数据风险")
                else:
                    for name, build_date, point_count in no_data_algos_display:
                        lines.append(
                            "   - %s：建立日期 %s，点位数量 %d"
                            % (name, fmt_table_date(build_date), point_count)
                        )
            else:
                no_data_template = risk_templates.get("no_data", {}) or {}
                no_data_summary = no_data_template.get(
                    "summary",
                    "部分算法点位数量为 0，暂不纳入本周统计，需确认是否已配置有效点位。",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(no_data_summary, report_context),
                ))
            risk_no += 1

        if risks_cfg.get("show_stale_risk", True) and stale_algos_display:
            if detailed:
                stale_template = risk_templates.get("stale", {}) or {}
                stale_intro = stale_template.get(
                    "detailed",
                    "数据失效风险：以下算法任务日期久远（超过 {stale_days} 天），存在数据失效风险：",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(stale_intro, report_context),
                ))
                stale_table = create_sublist_table_asset(
                    report_stem,
                    "数据失效风险",
                    ["算法小类", "点位数量", "任务日期"],
                    [
                        [a, pts, fmt_date(d)]
                        for a, d, pts in stale_algos_display
                    ],
                    table_layout_cfg,
                )
                if stale_table:
                    generated_table_paths.append(stale_table)
                    append_table_reference(lines, stale_table, "数据失效风险")
                else:
                    for a, d, pts in stale_algos_display:
                        lines.append("   - %s（%d 个点位，最后任务日期 %s）" % (a, pts, fmt_date(d)))
            else:
                stale_template = risk_templates.get("stale", {}) or {}
                stale_summary = stale_template.get(
                    "summary",
                    "部分算法任务日期久远，存在数据失效风险，需确认是否继续运行。",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(stale_summary, report_context),
                ))
            risk_no += 1

        if risks_cfg.get("show_idle_risk", True) and idle_algos_display:
            if detailed:
                idle_template = risk_templates.get("idle", {}) or {}
                idle_intro = idle_template.get(
                    "detailed",
                    "数据质量风险：以下算法长期无有效准确率数据，需关注数据质量与算法运行情况：",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(idle_intro, report_context),
                ))
                idle_risk_rows = []
                for a in idle_algos_display:
                    r = new_map.get(a, {})
                    pts = parse_int(r.get("点位数量"))
                    acc = get_priority_value(r, accuracy_field, cfg)
                    idle_risk_rows.append([a, pts, fmt_percent(acc)])
                idle_risk_table = create_sublist_table_asset(
                    report_stem,
                    "数据质量风险",
                    ["算法小类", "点位数量", accuracy_field_label],
                    idle_risk_rows,
                    table_layout_cfg,
                )
                if idle_risk_table:
                    generated_table_paths.append(idle_risk_table)
                    append_table_reference(lines, idle_risk_table, "数据质量风险")
                else:
                    for a, pts, acc in idle_risk_rows:
                        lines.append("   - %s（%d 个点位，准确率 %s）" % (a, pts, acc))
            else:
                idle_template = risk_templates.get("idle", {}) or {}
                idle_summary = idle_template.get(
                    "summary",
                    "部分算法长期无有效准确率数据，需关注数据质量与算法运行情况。",
                )
                lines.append("%d、%s" % (
                    risk_no,
                    render_template_text(idle_summary, report_context),
                ))
            risk_no += 1

        if risk_no == 1:
            lines.append("1、无。")

    if sections.get("next_steps", True):
        lines.extend(next_steps_lines)
        section_idx += 1

    report = "\n".join(lines)
    print(report)

    # 保存到文件
    out_path = report_stem + ".md"
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(report)
    print("\n[已保存] %s" % out_path)
    share_path = os.path.join(output_dir, "clipboard_share.html")
    with open(share_path, "w", encoding="utf-8") as f:
        f.write(report_to_share_html(
            report,
            output_dir,
            site_name,
            today.strftime("%Y-%m-%d"),
            read_lan_uploads_metadata(data_temp_dir),
        ))
    print("[已生成局域网复制页] %s" % share_path)
    full_report_image_path = None
    preview_image = create_report_preview_image(report, output_dir)
    if preview_image is not None:
        full_report_image_path = report_stem + "_整页周报.png"
        preview_image.save(full_report_image_path, "PNG")
        print("[已生成整页周报图片] %s" % full_report_image_path)
    elif pure_image_clipboard:
        print("[整页周报图片生成失败] 未安装 Pillow 或图片渲染失败")
    for asset_path in generated_table_paths:
        if asset_path.lower().endswith(".xlsx"):
            print("[已生成准确率表格] %s" % asset_path)
        else:
            print("[已生成表格截图] %s" % asset_path)
    if deleted_outputs:
        print("[已清理过期周报] %d 个文件" % len(deleted_outputs))
    for path, exc in cleanup_failures:
        print("[清理过期周报失败] %s: %s" % (path, exc))

    # 复制到剪贴板
    if run_options.get("no_clipboard"):
        print("[剪贴板复制已跳过]")
    else:
        clipboard_status = copy_rich_to_clipboard(
            report, output_dir, pure_image=pure_image_clipboard
        )
        if clipboard_status == "rich":
            if pure_image_clipboard:
                print("[已复制纯图片到剪贴板]")
            else:
                print("[已复制文字+表格图片到剪贴板]")
        elif clipboard_status == "text":
            print("[已复制普通文本到剪贴板，富文本图片复制失败]")
        else:
            print("[剪贴板复制失败]")

    share_url = None
    if os.environ.get("RUST_PORTAL_NO_LAN") != "1":
        share_url = start_lan_clipboard_server(
            output_dir,
            clipboard_cfg,
            session_dir=lan_session_dir,
        )
    if share_url:
        print("[局域网复制页] %s" % share_url)
    elif clipboard_cfg.get("lan_share", False):
        print("[局域网复制页启动失败，请检查 8765-8774 端口占用或防火墙]")

    return 0


if __name__ == "__main__":
    if len(sys.argv) >= 5 and sys.argv[1] == "--lan-share-server":
        run_lan_share_server(
            sys.argv[2],
            int(sys.argv[3]),
            sys.argv[4],
            sys.argv[5] if len(sys.argv) >= 6 else None,
        )
        sys.exit(0)
    sys.exit(main())
