from __future__ import annotations

import copy
from collections import Counter
from dataclasses import dataclass
import errno
import html
import hashlib
import ipaddress
import json
import logging
import os
import re
import shutil
import socket
import sys
import threading
import webbrowser
import zipfile
from datetime import datetime
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import parse_qs, quote, unquote, urlparse

from annotation_db import AnnotationDatabase


STATUS_NAMES = {
    0: "未标注",
    1: "对",
    2: "错",
    3: "问题图片",
}
SAVE_PROBLEM_KEYS = ("save_problem_images", "save_blur_images", "keep_problem_images")
RECOGNITION_RED_KEY = "recognition_red_keywords"
POINT_NAME_RED_KEY = "point_name_red_keywords"
POINT_NAME_END_OPEN_KEY = "End_open"
POINT_NAME_END_CHARACTERS_KEY = "End_characters"
POINT_NAME_END_CHARACTER_COLOR_KEY = "End_character_color"
DEFAULT_ZOOM_KEY = "default_zoom"
AUTO_CLICK_TIMES_KEY = "auto_click_times_ms"
SHOW_AUTO_CLICK_KEY = "show_auto_click"
SHOW_MARK_ALL_CORRECT_KEY = "show_mark_all_correct"
SHOW_MARK_SAME_RECOGNITION_KEY = "show_mark_same_recognition"
WEB_HOTKEYS_KEY = "web_hotkeys"
NEW_KEYWORDS_KEY = "newkeywords"
DEFAULT_WRONG_KEY = "default_wrong_keywords"
DEFAULT_PROBLEM_KEY = "default_problem_keywords"
EXTRA_XLSX_RECOGNITION_ENABLED_KEY = "extra_xlsx_recognition_enabled"
EXTRA_XLSX_RECOGNITION_KEY = "extra_xlsx_recognition_keywords"
SHOW_ALGORITHM_NORMAL_KEY = "show_algorithm_normal"
CHECKLIST_IMAGE_WIDTH_ENABLED_KEY = "checklist_image_width_enabled"
ALGO_MAPPING_RULES_KEY = "mapping_rules"
RESULT2_XLSX_PATH_KEY = "result2_xlsx_path"
RESULT2_ALGORITHM_MAPPING_RULES_KEY = "result2_algorithm_mapping_rules"
RESULT2_ALGORITHM_COLUMN_INDEX = 8
CHECKLIST_TABLE_REPORT_TITLE = "清单表格："
CHECKLIST_IMAGE_DPI = 96
CHECKLIST_IMAGE_MIN_WIDTH_PX = 8 / 2.54 * CHECKLIST_IMAGE_DPI
CHECKLIST_IMAGE_TARGET_WIDTH_PX = 8.41 / 2.54 * CHECKLIST_IMAGE_DPI
SOLUTION_RULES_KEY = "Assess_the_situation"
SOLUTION_OPTIONAL_KEY = "solution_optional"
SOLUTION_OPTIONAL_SHEET_NAME = "solution_optional"
RESPONSIBLE_RULES_KEY = "responsible_rules"
RESPONSIBLE_VALUES = ("算法调试", "现场调试", "误报无需调试")
SESSION_STATE_DIRNAME = ".review_state"
SESSION_BACKUP_DIRNAME = "images"
OUTER_INNER_PATTERN = re.compile(r"^(accuracy|image_result)_.+\.zip$", re.IGNORECASE)
DEFAULT_WEB_HOTKEYS = {
    "previous": "ArrowLeft",
    "next": "ArrowRight",
    "correct": "0",
    "problem": "ArrowUp",
    "wrong": "ArrowDown",
    "auto_click_toggle": "2",
}
POINT_NAME_END_COLORS = {
    "green": "#57d38c",
    "red": "#f06d73",
    "amber": "#f2bf67",
    "yellow": "#f7d774",
    "blue": "#65a9e9",
    "cyan": "#62d6e8",
    "purple": "#c59bff",
    "pink": "#ff9dca",
    "orange": "#ffb56b",
    "white": "#eef3f8",
}


HTML_PAGE = r"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>巡视点位确认</title>
  <style>
    :root {
      color-scheme: dark;
      --bg: #10151b;
      --panel: #19222c;
      --panel-strong: #202c38;
      --line: #334150;
      --text: #eef3f8;
      --muted: #9eacba;
      --green: #57d38c;
      --red: #f06d73;
      --amber: #f2bf67;
      --blue: #65a9e9;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      min-width: 960px;
      height: 100vh;
      overflow: hidden;
      background: var(--bg);
      color: var(--text);
      font: 13px/1.4 "Microsoft YaHei", "Segoe UI", sans-serif;
    }
    button, select { font: inherit; }
    button {
      border: 1px solid var(--line);
      border-radius: 6px;
      background: var(--panel-strong);
      color: var(--text);
      cursor: pointer;
      padding: 8px 12px;
    }
    button:hover:not(:disabled) { border-color: var(--blue); }
    button:disabled { cursor: not-allowed; opacity: .45; }
    .app {
      display: grid;
      grid-template-rows: 54px minmax(0, 1fr) 58px;
      gap: 10px;
      height: 100vh;
      padding: 10px;
    }
    .topbar, .bottombar, .info-panel, .image-panel {
      border: 1px solid var(--line);
      background: var(--panel);
    }
    .topbar {
      display: flex;
      align-items: center;
      gap: 10px;
      padding: 8px 12px;
      overflow: hidden;
    }
    .brand {
      flex: 0 0 auto;
      font-size: 16px;
      font-weight: 700;
      white-space: nowrap;
    }
    .top-stats {
      display: flex;
      gap: 8px;
      flex: 1 1 auto;
      min-width: 0;
      overflow-x: auto;
    }
    .status-summary {
      display: flex;
      gap: 8px;
      flex: 0 0 auto;
    }
    .stat-chip {
      min-width: 116px;
      padding: 5px 10px;
      border: 0;
      border-left: 3px solid var(--line);
      background: #151d25;
      color: var(--text);
      cursor: pointer;
      text-align: left;
      white-space: nowrap;
    }
    .stat-chip strong { display: block; font-size: 13px; overflow: hidden; text-overflow: ellipsis; }
    .stat-chip span { color: var(--muted); font-size: 12px; }
    .stat-chip.done { border-left-color: var(--green); }
    .stat-chip.active { outline: 1px solid var(--blue); }
    .summary-chip {
      min-width: 92px;
      padding: 5px 10px;
      border-left: 3px solid var(--line);
      background: #151d25;
      white-space: nowrap;
    }
    .summary-chip strong { display: block; font-size: 13px; }
    .summary-chip span { color: var(--muted); font-size: 12px; }
    .summary-chip.status-1 { border-left-color: var(--green); }
    .summary-chip.status-2 { border-left-color: var(--red); }
    .summary-chip.status-3 { border-left-color: var(--amber); }
    .workspace {
      display: grid;
      grid-template-columns: minmax(150px, 14%) minmax(0, 86%);
      gap: 10px;
      min-height: 0;
    }
    .info-panel {
      display: flex;
      min-height: 0;
      flex-direction: column;
      padding: 8px;
      overflow: auto;
    }
    .section-title {
      margin: 0 0 10px;
      color: var(--muted);
      font-size: 11px;
      letter-spacing: 0;
      text-transform: uppercase;
    }
    .meta {
      display: grid;
      gap: 5px;
      margin: 0 0 10px;
    }
    .meta-row {
      display: grid;
      grid-template-columns: 50px minmax(0, 1fr);
      gap: 6px;
      padding-bottom: 5px;
      border-bottom: 1px solid #293540;
    }
    .meta-row dt { color: var(--muted); }
    .meta-row dd { margin: 0; word-break: break-word; }
    .recognition {
      margin: 0;
      padding: 8px;
      border: 1px solid #394b5d;
      background: #141c24;
      color: #f7e7a6;
      white-space: pre-wrap;
      word-break: break-word;
    }
    .recognition .hit { color: var(--red); font-weight: 700; }
    .labels {
      display: flex;
      flex-wrap: wrap;
      gap: 6px;
      margin-bottom: 16px;
    }
    .label {
      border-radius: 4px;
      padding: 4px 7px;
      color: var(--muted);
      background: #27333f;
    }
    .label.current { color: var(--text); border: 1px solid var(--blue); }
    .label.status-1 { color: var(--green); }
    .label.status-2 { color: var(--red); }
    .label.status-3 { color: var(--amber); }
    .auto-click-row {
      display: grid;
      gap: 8px;
      margin: -8px 0 14px;
    }
    .auto-click-primary,
    .auto-click-status {
      display: flex;
      align-items: center;
      gap: 6px;
    }
    .auto-click-primary {
      flex-wrap: nowrap;
    }
    .auto-click-status {
      justify-content: flex-start;
    }
    .auto-click-toggle {
      min-height: 0;
      border: 1px solid var(--line);
      text-align: center;
    }
    .auto-click-toggle.active {
      border-color: var(--red);
      background: #3a1c22;
      color: var(--red);
    }
    .auto-click-select {
      min-width: 78px;
      height: 27px;
      border: 1px solid var(--line);
      border-radius: 4px;
      background: var(--panel-strong);
      color: var(--text);
      padding: 2px 5px;
    }
    .status-actions {
      display: grid;
      gap: 8px;
      margin-top: auto;
    }
    .status-actions button { min-height: 40px; text-align: left; }
    .status-actions button.active.status-1 { border-color: var(--green); color: var(--green); }
    .status-actions button.active.status-2 { border-color: var(--red); color: var(--red); }
    .status-actions button.active.status-3 { border-color: var(--amber); color: var(--amber); }
    .save-row {
      display: flex;
      align-items: center;
      gap: 8px;
      margin-top: 10px;
    }
    .save-row button { flex: 1; }
    .save-state { color: var(--muted); font-size: 12px; }
    .image-panel {
      display: grid;
      grid-template-rows: 38px minmax(0, 1fr);
      min-width: 0;
      min-height: 0;
    }
    .image-toolbar {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 8px;
      padding: 5px 8px;
      border-bottom: 1px solid var(--line);
      color: var(--muted);
    }
    .image-recognition {
      position: relative;
      flex: 1 1 auto;
      max-height: 32px;
      overflow: auto;
      text-align: center;
      font-size: 15px;
      font-weight: 700;
      line-height: 1.2;
    }
    .toolbar-point-name,
    .image-recognition-text {
      position: absolute;
      top: 50%;
      transform: translateY(-50%);
    }
    .toolbar-point-name {
      left: 8px;
      width: calc(33.333% - 8px);
      overflow: hidden;
      color: var(--text);
      font-size: 12px;
      font-weight: 400;
      line-height: 1.2;
      text-align: left;
      white-space: nowrap;
      z-index: 1;
    }
    .toolbar-point-name .point-name-hit { color: var(--red); font-weight: 700; }
    .image-recognition-text {
      left: calc(33.333% + 8px);
      right: 8px;
      max-height: 24px;
      overflow: auto;
      text-align: center;
      z-index: 2;
    }
    .image-toolbar .tools { display: flex; gap: 6px; }
    .image-toolbar button { min-width: 34px; padding: 5px 8px; }
    .image-viewport {
      display: flex;
      min-height: 0;
      align-items: center;
      justify-content: center;
      overflow: auto;
      background:
        linear-gradient(45deg, #151d25 25%, transparent 25%) 0 0 / 24px 24px,
        linear-gradient(-45deg, #151d25 25%, transparent 25%) 0 0 / 24px 24px,
        linear-gradient(45deg, transparent 75%, #151d25 75%) 0 0 / 24px 24px,
        linear-gradient(-45deg, transparent 75%, #151d25 75%) 0 0 / 24px 24px,
        #111820;
    }
    #mainImage {
      display: block;
      max-width: 100%;
      max-height: 100%;
      width: 100%;
      height: 100%;
      object-fit: contain;
      transform-origin: center center;
      user-select: none;
    }
    .empty {
      color: var(--muted);
      text-align: center;
    }
    .bottombar {
      display: flex;
      align-items: center;
      gap: 10px;
      padding: 8px 10px;
      min-width: 0;
    }
    .category-select {
      min-width: 270px;
      max-width: 34vw;
      padding: 8px;
      border: 1px solid var(--line);
      border-radius: 6px;
      background: var(--panel-strong);
      color: var(--text);
    }
    .counter {
      color: var(--muted);
      white-space: nowrap;
    }
    .nav-actions { display: flex; gap: 7px; margin-left: auto; }
    .export {
      border-color: var(--green);
      color: var(--green);
      font-weight: 700;
    }
    .export.ready { background: #173c2a; }
    .notice {
      position: fixed;
      right: 18px;
      bottom: 78px;
      max-width: 420px;
      padding: 10px 14px;
      border: 1px solid var(--line);
      border-radius: 6px;
      background: #1c2731;
      color: var(--text);
      box-shadow: 0 8px 28px #0008;
    }
    .notice.error { border-color: var(--red); }
    @media (max-width: 1100px) {
      body { min-width: 0; }
      .app { padding: 6px; gap: 6px; }
      .workspace { grid-template-columns: minmax(160px, 22%) minmax(0, 78%); }
      .category-select { min-width: 210px; }
    }
  </style>
</head>
<body>
  <main class="app">
    <header class="topbar">
      <div class="brand">巡视点位确认</div>
      <div id="topStats" class="top-stats"></div>
      <div id="statusSummary" class="status-summary"></div>
    </header>
    <section class="workspace">
      <aside class="info-panel">
        <p class="section-title">点位信息</p>
        <dl class="meta">
          <div class="meta-row"><dt>类别</dt><dd id="algoType">-</dd></div>
          <div class="meta-row"><dt>点位编码</dt><dd id="pointCode">-</dd></div>
          <div class="meta-row"><dt>图片</dt><dd id="imageName">-</dd></div>
        </dl>
        <div id="labels" class="labels"></div>
        <div id="autoClickControls" class="auto-click-row">
          <div class="auto-click-primary">
            <button id="autoClickToggle" class="label auto-click-toggle" type="button">按时点击</button>
            <select id="autoClickSelect" class="auto-click-select" aria-label="按时点击时间"></select>
          </div>
          <div class="auto-click-status">
            <select id="autoClickStatusSelect" class="auto-click-select" aria-label="按时点击结果">
              <option value="1">对</option>
              <option value="2">错</option>
              <option value="3">问题图片</option>
            </select>
          </div>
        </div>
        <div class="status-actions">
          <button id="markAllCorrectButton" class="status-1" type="button" hidden>一键为对</button>
          <button data-status="1" class="status-1">对</button>
          <button data-status="2" class="status-2">错</button>
          <button data-status="3" class="status-3">问题图片</button>
        </div>
        <div class="save-row">
          <button id="saveButton">保存当前</button>
          <span id="saveState" class="save-state"></span>
        </div>
      </aside>
      <section class="image-panel">
        <div class="image-toolbar">
          <span id="imageCounter">-</span>
          <div id="recognition" class="recognition image-recognition">
            <span id="pointName" class="toolbar-point-name">-</span>
            <span id="recognitionText" class="image-recognition-text">-</span>
          </div>
          <div class="tools">
            <button id="zoomOut" title="缩小">−</button>
            <button id="zoomReset" title="重置缩放">100%</button>
            <button id="zoomIn" title="放大">+</button>
          </div>
        </div>
        <div id="imageViewport" class="image-viewport">
          <img id="mainImage" alt="点位图片">
          <div id="emptyImage" class="empty" hidden>图片不存在</div>
        </div>
      </section>
    </section>
    <footer class="bottombar">
      <select id="categorySelect" class="category-select" aria-label="类别"></select>
      <span id="categoryCounter" class="counter">类别点位 -</span>
      <span id="remainingCounter" class="counter">剩余 -</span>
      <button id="firstUnmarkedButton" title="跳到当前类别最左侧未标注图片" hidden>跳到未标注</button>
      <button id="wrongMarkedButton" title="跳到当前类别已标注为错的图片" hidden>跳到标注为错:0/0</button>
      <button id="sameRecognitionUnmarkedButton" title="跳到当前识别结果类别最左侧图片" hidden>分析类别跳转</button>
      <button id="sameRecognitionMarkButton" title="批量标注当前识别结果类别中未确认的点位" hidden>批量标注</button>
      <select id="sameRecognitionMarkSelect" aria-label="批量标注结果" hidden>
        <option value="1">对</option>
        <option value="2">错</option>
        <option value="3">问题图片</option>
      </select>
      <div class="nav-actions">
        <button id="previousButton" title="上一张">上一张</button>
        <button id="nextButton" title="下一张">下一张</button>
        <button id="exportButton" class="export" hidden>导出压缩</button>
      </div>
    </footer>
  </main>
  <div id="notice" class="notice" hidden></div>
  <script>
    const app = {
      categories: [],
      categoryIndex: 0,
      imageIndex: 0,
      draftStatus: 0,
      savedStatus: 0,
      zoom: 1,
      defaultZoom: 1,
      recognitionRedKeywords: [],
      pointNameRedKeywords: [],
      pointNameEndOpen: false,
      pointNameEndCharacters: 0,
      pointNameEndCharacterColor: "",
      exitAfterExport: false,
      defaultWrongKeywords: [],
      defaultProblemKeywords: [],
      autoClickTimes: [],
      showAutoClick: false,
      showMarkAllCorrect: false,
      showMarkSameRecognition: false,
      webHotkeys: {},
      autoClickEnabled: false,
      autoClickDelay: 0,
      autoClickStatus: 1,
      autoClickTimer: null,
      saving: false,
      noticeTimer: null
    };

    const $ = (id) => document.getElementById(id);
    const currentCategory = () => app.categories[app.categoryIndex];
    const currentItem = () => currentCategory()?.images[app.imageIndex];
    const hotkeyMatches = (event, action) => app.webHotkeys[action] && event.key === app.webHotkeys[action];
    const statusText = (value) => ({0: "未标注", 1: "对", 2: "错", 3: "问题图片"}[value] || "未标注");
    const suggestedStatus = (item) => {
      if (!item || item.auditStatus !== 0) return item ? item.auditStatus : 0;
      const recognition = String(item.recognition || "");
      if (app.defaultProblemKeywords.some(keyword => keyword && recognition.includes(keyword))) return 3;
      return app.defaultWrongKeywords.some(keyword => keyword && recognition.includes(keyword)) ? 2 : 0;
    };

    function clearAutoClickTimer() {
      if (app.autoClickTimer) {
        clearTimeout(app.autoClickTimer);
        app.autoClickTimer = null;
      }
    }

    function scheduleAutoClickTimer() {
      clearAutoClickTimer();
      const item = currentItem();
      if (!app.autoClickEnabled || app.autoClickDelay <= 0 || app.saving || !item || item.auditStatus !== 0 || allDone()) return;
      app.autoClickTimer = setTimeout(() => {
        app.autoClickTimer = null;
        const activeItem = currentItem();
        if (!app.autoClickEnabled || app.saving || !activeItem || activeItem.auditStatus !== 0 || allDone()) return;
        saveCurrent(true, app.autoClickStatus).catch(error => showNotice(error.message, true));
      }, app.autoClickDelay);
    }

    function refreshAutoClickUI() {
      $("autoClickControls").hidden = !app.showAutoClick;
      if (!app.showAutoClick) {
        app.autoClickEnabled = false;
        clearAutoClickTimer();
      }
      $("autoClickToggle").classList.toggle("active", app.autoClickEnabled);
      $("autoClickToggle").setAttribute("aria-pressed", String(app.autoClickEnabled));
      $("autoClickSelect").disabled = !app.autoClickTimes.length;
      $("autoClickSelect").value = String(app.autoClickDelay || (app.autoClickTimes[0] ?? 0));
      $("autoClickStatusSelect").disabled = !app.showAutoClick;
      $("autoClickStatusSelect").value = String(app.autoClickStatus);
    }

    function toggleAutoClick() {
      if (!app.showAutoClick) return;
      app.autoClickEnabled = !app.autoClickEnabled;
      refreshAutoClickUI();
      if (app.autoClickEnabled) {
        scheduleAutoClickTimer();
      } else {
        clearAutoClickTimer();
      }
    }

    function loadItemState() {
      const item = currentItem();
      app.draftStatus = item ? suggestedStatus(item) : 0;
      app.savedStatus = item ? item.auditStatus : 0;
      app.zoom = app.defaultZoom;
      const image = $("mainImage");
      if (image) image.style.transformOrigin = "center center";
    }

    function showNotice(message, error = false) {
      const node = $("notice");
      node.textContent = message;
      node.className = error ? "notice error" : "notice";
      node.hidden = false;
      clearTimeout(app.noticeTimer);
      app.noticeTimer = setTimeout(() => node.hidden = true, 2600);
    }

    function statusValue(value) {
      const number = Number(value);
      return [0, 1, 2, 3].includes(number) ? number : 0;
    }

    function categoryDone(category) {
      return category.images.every(item => item.auditStatus !== 0);
    }

    function allDone() {
      return app.categories.length > 0 && app.categories.every(categoryDone);
    }

    function firstUnmarkedIndex(category) {
      return category ? category.images.findIndex(item => item.auditStatus === 0) : -1;
    }

    function wrongMarkedIndices(category) {
      if (!category) return [];
      return category.images
        .map((item, index) => (statusValue(item.auditStatus) === 2 ? index : -1))
        .filter(index => index >= 0);
    }

    function wrongMarkedState(category, imageIndex) {
      const indices = wrongMarkedIndices(category);
      const total = indices.length;
      if (!total) {
        return { total: 0, display: 0, targetIndex: -1 };
      }
      const currentPosition = indices.indexOf(imageIndex);
      if (currentPosition >= 0) {
        return {
          total,
          display: currentPosition + 1,
          targetIndex: indices[(currentPosition + 1) % total],
        };
      }
      const nextPosition = indices.findIndex(index => index > imageIndex);
      const targetPosition = nextPosition >= 0 ? nextPosition : 0;
      return {
        total,
        display: targetPosition + 1,
        targetIndex: indices[targetPosition],
      };
    }

    function firstSameRecognitionIndex(category, item) {
      if (!category || !item) return -1;
      const recognition = String(item.recognition || "");
      return category.images.findIndex(candidate => String(candidate.recognition || "") === recognition);
    }

    function firstSameRecognitionUnmarkedIndex(category, item) {
      if (!category || !item) return -1;
      const recognition = String(item.recognition || "");
      return category.images.findIndex(candidate =>
        candidate.auditStatus === 0 && String(candidate.recognition || "") === recognition
      );
    }

    function applyState(state) {
      app.categories = state.categories || [];
      app.defaultZoom = Number(state.defaultZoom || 1);
      app.recognitionRedKeywords = Array.isArray(state.recognitionRedKeywords) ? state.recognitionRedKeywords : [];
      app.pointNameRedKeywords = Array.isArray(state.pointNameRedKeywords) ? state.pointNameRedKeywords : [];
      app.pointNameEndOpen = state.pointNameEndOpen === true;
      app.pointNameEndCharacters = Number(state.pointNameEndCharacters || 0);
      app.pointNameEndCharacterColor = String(state.pointNameEndCharacterColor || "");
      app.exitAfterExport = state.exitAfterExport === true;
      app.defaultWrongKeywords = Array.isArray(state.defaultWrongKeywords) ? state.defaultWrongKeywords : [];
      app.defaultProblemKeywords = Array.isArray(state.defaultProblemKeywords) ? state.defaultProblemKeywords : [];
      app.showAutoClick = state.showAutoClick === true;
      app.showMarkAllCorrect = state.showMarkAllCorrect === true;
      app.showMarkSameRecognition = state.showMarkSameRecognition === true;
      app.webHotkeys = state.webHotkeys && typeof state.webHotkeys === "object" ? state.webHotkeys : {};
      app.autoClickTimes = Array.isArray(state.autoClickTimesMs)
        ? state.autoClickTimesMs.map(value => Number(value)).filter(value => Number.isFinite(value) && value > 0)
        : [];
      if (!app.autoClickTimes.length) app.autoClickTimes = [1000, 2000, 3000];
      if (!app.autoClickDelay || !app.autoClickTimes.includes(app.autoClickDelay)) {
        app.autoClickDelay = app.autoClickTimes[0];
      }
      app.categoryIndex = Math.max(0, Math.min(
        Number(state.currentCategoryIndex ?? app.categoryIndex),
        Math.max(app.categories.length - 1, 0)
      ));
      const category = currentCategory();
      app.imageIndex = category ? Math.max(0, Math.min(
        Number(state.currentImageIndex ?? app.imageIndex),
        Math.max(category.images.length - 1, 0)
      )) : 0;
      loadItemState();
      renderAutoClickOptions();
      refreshAutoClickUI();
      scheduleAutoClickTimer();
      render();
    }

    function renderAutoClickOptions() {
      const select = $("autoClickSelect");
      if (!select) return;
      select.innerHTML = "";
      app.autoClickTimes.forEach(value => {
        const option = document.createElement("option");
        option.value = String(value);
        option.textContent = `${value} ms`;
        select.appendChild(option);
      });
      select.value = String(app.autoClickDelay);
    }

    function renderTopStats() {
      $("topStats").innerHTML = app.categories.map((category, index) => {
        const done = category.images.filter(item => item.auditStatus !== 0).length;
        return `<button type="button" data-category-index="${index}" class="stat-chip ${categoryDone(category) ? "done" : ""} ${index === app.categoryIndex ? "active" : ""}">
          <strong>${escapeHtml(category.algoType)}</strong><span>${done}/${category.images.length}</span>
        </button>`;
      }).join("");
    }

    function renderStatusSummary() {
      const category = currentCategory();
      if (!category) {
        $("statusSummary").innerHTML = "";
        return;
      }
      const counts = {
        1: category.correct ?? category.images.filter(item => item.auditStatus === 1).length,
        2: category.inaccuracy ?? category.images.filter(item => item.auditStatus === 2).length,
        3: category.blur ?? category.images.filter(item => item.auditStatus === 3).length
      };
      $("statusSummary").innerHTML = `
        <div class="summary-chip status-1"><strong>对</strong><span>${counts[1]}</span></div>
        <div class="summary-chip status-2"><strong>错</strong><span>${counts[2]}</span></div>
        <div class="summary-chip status-3"><strong>问题</strong><span>${counts[3]}</span></div>
      `;
    }

    function renderCategoryOptions() {
      const select = $("categorySelect");
      select.innerHTML = "";
      app.categories.forEach((category, index) => {
        const option = document.createElement("option");
        option.value = index;
        option.textContent = `${category.algoType} (${category.images.filter(item => item.auditStatus !== 0).length}/${category.images.length})`;
        option.style.color = categoryDone(category) ? "#57d38c" : "#eef3f8";
        select.appendChild(option);
      });
      select.value = String(app.categoryIndex);
    }

    function render() {
      const category = currentCategory();
      const item = currentItem();
      if (!category) return;
      if (!item) {
        $("pointName").innerHTML = "(该类别无点位)";
        $("pointName").title = "";
        $("algoType").textContent = category.algoType;
        $("pointCode").textContent = "-";
        $("imageName").textContent = "-";
        $("recognitionText").innerHTML = "-";
        $("imageCounter").textContent = "0 / 0";
        $("categoryCounter").textContent = "类别点位 0/0";
        $("remainingCounter").textContent = "剩余 0";
        $("firstUnmarkedButton").hidden = true;
        $("wrongMarkedButton").hidden = true;
        $("sameRecognitionUnmarkedButton").hidden = true;
        $("sameRecognitionMarkButton").hidden = true;
        $("sameRecognitionMarkSelect").hidden = true;
        $("previousButton").disabled = app.categoryIndex === 0;
        $("nextButton").disabled = app.categoryIndex === app.categories.length - 1;
        $("saveButton").disabled = true;
        $("markAllCorrectButton").hidden = !app.showMarkAllCorrect;
        $("markAllCorrectButton").disabled = true;
        $("saveState").textContent = "无点位";
        document.querySelectorAll("[data-status]").forEach(button => button.disabled = true);
        $("labels").innerHTML = `<span class="label current">已完成</span>`;
        $("mainImage").hidden = true;
        $("emptyImage").hidden = false;
        renderTopStats();
        renderStatusSummary();
        renderCategoryOptions();
        $("exportButton").hidden = !allDone();
        $("exportButton").classList.toggle("ready", allDone());
        refreshAutoClickUI();
        return;
      }
      $("pointName").innerHTML = renderPointName(item.pointName || "(未命名点位)");
      $("pointName").title = String(item.pointName || "");
      $("algoType").textContent = category.algoType;
      $("pointCode").textContent = item.pointCode || "-";
      $("imageName").textContent = item.imageName || "-";
      $("recognitionText").innerHTML = renderRecognition(item.recognition ?? "-");
      $("imageCounter").textContent = `${app.imageIndex + 1} / ${category.images.length}`;
      $("categoryCounter").textContent = `类别点位 ${category.completed}/${category.images.length}`;
      const remaining = category.images.length - category.completed;
      const wrongMarked = wrongMarkedState(category, app.imageIndex);
      const sameRecognitionIndex = firstSameRecognitionIndex(category, item);
      const sameRecognitionUnmarkedIndex = firstSameRecognitionUnmarkedIndex(category, item);
      $("remainingCounter").textContent = `剩余 ${remaining}`;
      $("firstUnmarkedButton").hidden = remaining === 0;
      $("firstUnmarkedButton").disabled = app.saving || remaining === 0;
      $("wrongMarkedButton").hidden = wrongMarked.total === 0;
      $("wrongMarkedButton").disabled = app.saving || wrongMarked.total === 0;
      $("wrongMarkedButton").textContent = `跳到标注为错:${wrongMarked.display}/${wrongMarked.total}`;
      $("sameRecognitionUnmarkedButton").hidden = sameRecognitionIndex < 0;
      $("sameRecognitionUnmarkedButton").disabled = app.saving || sameRecognitionIndex < 0;
      $("sameRecognitionMarkButton").hidden = !app.showMarkSameRecognition || remaining === 0;
      $("sameRecognitionMarkSelect").hidden = !app.showMarkSameRecognition || remaining === 0;
      $("sameRecognitionMarkButton").disabled = app.saving || sameRecognitionUnmarkedIndex < 0;
      $("sameRecognitionMarkSelect").disabled = app.saving || sameRecognitionUnmarkedIndex < 0;
      $("previousButton").disabled = app.categoryIndex === 0 && app.imageIndex === 0;
      $("nextButton").disabled = app.categoryIndex === app.categories.length - 1 &&
        app.imageIndex === category.images.length - 1;
      $("saveButton").disabled = app.saving || app.draftStatus === app.savedStatus;
      $("markAllCorrectButton").hidden = !app.showMarkAllCorrect;
      $("markAllCorrectButton").disabled = app.saving || category.images.length === 0;
      $("saveState").textContent = app.saving ? "保存中" : (app.draftStatus === app.savedStatus ? "已保存" : "待保存");
      document.querySelectorAll("[data-status]").forEach(button => {
        const status = Number(button.dataset.status);
        button.disabled = app.saving;
        button.classList.toggle("active", status === app.draftStatus);
      });
      $("labels").innerHTML = `
        <span class="label current">${escapeHtml(statusText(app.draftStatus))}</span>
        <span class="label status-${item.auditStatus}">已确认: ${escapeHtml(statusText(item.auditStatus))}</span>
      `;
      const image = $("mainImage");
      image.src = `/image/${encodeURIComponent(category.algoType)}/${encodeURIComponent(item.imageName)}`;
      image.alt = item.pointName || item.imageName || "点位图片";
      image.hidden = false;
      $("emptyImage").hidden = true;
      setZoom(app.zoom);
      renderTopStats();
      renderStatusSummary();
      renderCategoryOptions();
      $("exportButton").hidden = !allDone();
      $("exportButton").classList.toggle("ready", allDone());
      refreshAutoClickUI();
    }

    function setZoom(value, anchor = null) {
      app.zoom = Math.max(.25, Math.min(4, value));
      const image = $("mainImage");
      if (anchor && image && !image.hidden) {
        const rect = image.getBoundingClientRect();
        if (rect.width > 0 && rect.height > 0) {
          const x = Math.max(0, Math.min(100, ((anchor.clientX - rect.left) / rect.width) * 100));
          const y = Math.max(0, Math.min(100, ((anchor.clientY - rect.top) / rect.height) * 100));
          image.style.transformOrigin = `${x}% ${y}%`;
        }
      }
      image.style.transform = `scale(${app.zoom})`;
      $("zoomReset").textContent = `${Math.round(app.zoom * 100)}%`;
    }

    function renderRecognition(value) {
      const text = String(value ?? "-");
      const keywords = app.recognitionRedKeywords
        .map(item => String(item || ""))
        .filter(Boolean)
        .sort((left, right) => right.length - left.length);
      if (!keywords.length) return escapeHtml(text);
      return renderKeywordSegments(text, keywords, "hit");
    }

    function renderPointName(value) {
      const text = String(value ?? "-");
      const keywords = app.pointNameRedKeywords
        .map(item => String(item || ""))
        .filter(Boolean)
        .sort((left, right) => right.length - left.length);
      const endCharacters = app.pointNameEndOpen ? Math.max(0, Math.floor(Number(app.pointNameEndCharacters || 0))) : 0;
      const endColor = pointNameEndColor();
      const visible = pointNameVisibleText(text);
      if (!keywords.length && !endCharacters && !visible.prefix) return escapeHtml(visible.text);
      return renderPointNameSegments(visible.text, visible.prefix, keywords, endCharacters, endColor);
    }

    function renderKeywordSegments(text, keywords, className) {
      const pattern = new RegExp(keywords.map(escapeRegExp).join("|"), "g");
      const parts = [];
      let offset = 0;
      text.replace(pattern, (match, index) => {
        parts.push(escapeHtml(text.slice(offset, index)));
        parts.push(`<span class="${className}">${escapeHtml(match)}</span>`);
        offset = index + match.length;
        return match;
      });
      parts.push(escapeHtml(text.slice(offset)));
      return parts.join("");
    }

    function pointNameVisibleText(text) {
      const limit = pointNameDisplayLimit();
      if (text.length <= limit) return {prefix: "", text};
      return {prefix: "...", text: text.slice(-limit)};
    }

    function pointNameDisplayLimit() {
      const node = $("pointName");
      const width = node?.getBoundingClientRect?.().width || node?.clientWidth || 0;
      if (!width) return 36;
      return Math.max(12, Math.floor(width / 12));
    }

    function renderPointNameSegments(text, prefix, keywords, endCharacters, endColor) {
      const endStart = endCharacters > 0 ? Math.max(0, text.length - endCharacters) : text.length;
      const pattern = keywords.length ? new RegExp(keywords.map(escapeRegExp).join("|"), "g") : null;
      const parts = [];
      const endStyle = endColor ? ` style="color: ${escapeHtml(endColor)}"` : "";

      const appendPlain = (plainStart, plainEnd) => {
        if (plainEnd <= plainStart) return;
        const segment = text.slice(plainStart, plainEnd);
        if (!segment) return;
        if (!endColor || plainEnd <= endStart) {
          parts.push(escapeHtml(segment));
          return;
        }
        if (plainStart >= endStart) {
          parts.push(`<span class="point-name-end"${endStyle}>${escapeHtml(segment)}</span>`);
          return;
        }
        const prefixText = text.slice(plainStart, endStart);
        const suffixText = text.slice(endStart, plainEnd);
        if (prefixText) parts.push(escapeHtml(prefixText));
        if (suffixText) parts.push(`<span class="point-name-end"${endStyle}>${escapeHtml(suffixText)}</span>`);
      };

      let offset = 0;
      if (pattern) {
        for (const match of text.matchAll(pattern)) {
          const index = match.index ?? 0;
          appendPlain(offset, index);
          parts.push(`<span class="point-name-hit">${escapeHtml(match[0])}</span>`);
          offset = index + match[0].length;
        }
      }
      appendPlain(offset, text.length);
      return `${escapeHtml(prefix)}${parts.join("")}`;
    }

    function pointNameEndColor() {
      const color = String(app.pointNameEndCharacterColor || "").trim().toLowerCase();
      const namedColor = {
        green: "#57d38c",
        red: "#f06d73",
        amber: "#f2bf67",
        yellow: "#f7d774",
        blue: "#65a9e9",
        cyan: "#62d6e8",
        purple: "#c59bff",
        pink: "#ff9dca",
        orange: "#ffb56b",
        white: "#eef3f8",
      }[color];
      if (namedColor) return namedColor;
      if (/^#[0-9a-f]{3,8}$/i.test(color)) return color;
      if (/^(?:rgb|hsl)a?\([^()]+\)$/i.test(color)) return color;
      return "";
    }

    function escapeRegExp(value) {
      return String(value).replace(/[.*+?^${}()|[\]\\]/g, "\\$&");
    }

    function escapeHtml(value) {
      return String(value ?? "").replace(/[&<>"']/g, character => ({
        "&": "&amp;", "<": "&lt;", ">": "&gt;", '"': "&quot;", "'": "&#39;"
      }[character]));
    }

    async function goTo(categoryIndex, imageIndex) {
      if (app.saving) return;
      const response = await fetch("/api/navigate", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({categoryIndex, imageIndex})
      });
      const result = await response.json();
      if (!response.ok) throw new Error(result.error || "跳转失败");
      applyState(result);
    }

    async function move(delta) {
      const category = currentCategory();
      if (!category) return;
      let nextCategory = app.categoryIndex;
      let nextImage = app.imageIndex + delta;
      if (nextImage < 0) {
        nextCategory -= 1;
        nextImage = nextCategory >= 0 ? app.categories[nextCategory].images.length - 1 : 0;
      } else if (nextImage >= category.images.length) {
        nextCategory += 1;
        nextImage = nextCategory < app.categories.length ? 0 : category.images.length - 1;
      }
      if (nextCategory < 0 || nextCategory >= app.categories.length) return;
      await goTo(nextCategory, nextImage);
    }

    async function saveCurrent(advance = false, auditStatus = app.draftStatus) {
      const category = currentCategory();
      if (!category || app.saving) return;
      app.saving = true;
      app.draftStatus = auditStatus;
      render();
      try {
        const response = await fetch("/api/mark", {
          method: "POST",
          headers: {"Content-Type": "application/json"},
          body: JSON.stringify({
            algoType: category.algoType,
            index: app.imageIndex,
            auditStatus,
            advance
          })
        });
        const result = await response.json();
        if (!response.ok) throw new Error(result.error || "保存失败");
        app.saving = false;
        applyState(result);
        if (!advance) showNotice("当前点位已保存");
      } catch (error) {
        app.saving = false;
        render();
        throw error;
      }
    }

    async function markCurrentCategoryCorrect() {
      const category = currentCategory();
      if (!category || app.saving) return;
      app.saving = true;
      render();
      try {
        const response = await fetch("/api/mark-category-correct", {
          method: "POST",
          headers: {"Content-Type": "application/json"},
          body: JSON.stringify({algoType: category.algoType})
        });
        const result = await response.json();
        if (!response.ok) throw new Error(result.error || "一键为对失败");
        app.saving = false;
        applyState(result);
      } catch (error) {
        app.saving = false;
        render();
        throw error;
      }
    }

    async function markSameRecognition() {
      const category = currentCategory();
      const item = currentItem();
      if (!category || !item || app.saving) return;
      const auditStatus = Number($("sameRecognitionMarkSelect").value);
      if (![1, 2, 3].includes(auditStatus)) return;
      app.saving = true;
      render();
      try {
        const response = await fetch("/api/mark-recognition", {
          method: "POST",
          headers: {"Content-Type": "application/json"},
          body: JSON.stringify({
            algoType: category.algoType,
            recognition: String(item.recognition || ""),
            auditStatus
          })
        });
        const result = await response.json();
        if (!response.ok) throw new Error(result.error || "批量标注失败");
        app.saving = false;
        applyState(result);
      } catch (error) {
        app.saving = false;
        render();
        throw error;
      }
    }

    async function exportPackage() {
      if (!allDone()) return;
      const button = $("exportButton");
      button.disabled = true;
      button.textContent = "导出中...";
      try {
        const response = await fetch("/api/export", {method: "POST"});
        const result = await response.json();
        if (!response.ok) throw new Error(result.error || "导出失败");
        showNotice(`导出完成: ${result.package}`);
        button.textContent = "已完成";
        setTimeout(() => {
          if (app.exitAfterExport) {
            window.close();
            return;
          }
          window.location.href = "/";
        }, 500);
      } catch (error) {
        button.disabled = false;
        button.textContent = "导出压缩";
        showNotice(error.message, true);
      }
    }

    $("topStats").addEventListener("click", event => {
      const button = event.target.closest("[data-category-index]");
      if (!button || app.saving) return;
      goTo(Number(button.dataset.categoryIndex), 0).catch(error => showNotice(error.message, true));
    });
    $("categorySelect").addEventListener("change", event => {
      if (app.saving) {
        event.target.value = String(app.categoryIndex);
        return;
      }
      goTo(Number(event.target.value), 0).catch(error => showNotice(error.message, true));
    });
    document.querySelectorAll("[data-status]").forEach(button => {
      button.addEventListener("click", () => {
        saveCurrent(true, Number(button.dataset.status)).catch(error => showNotice(error.message, true));
      });
    });
    $("saveButton").addEventListener("click", () => saveCurrent().catch(error => showNotice(error.message, true)));
    $("markAllCorrectButton").addEventListener("click", () => {
      markCurrentCategoryCorrect().catch(error => showNotice(error.message, true));
    });
    $("firstUnmarkedButton").addEventListener("click", () => {
      const index = firstUnmarkedIndex(currentCategory());
      if (index >= 0) goTo(app.categoryIndex, index).catch(error => showNotice(error.message, true));
    });
    $("wrongMarkedButton").addEventListener("click", () => {
      const state = wrongMarkedState(currentCategory(), app.imageIndex);
      if (state.targetIndex >= 0) {
        goTo(app.categoryIndex, state.targetIndex).catch(error => showNotice(error.message, true));
      }
    });
    $("sameRecognitionUnmarkedButton").addEventListener("click", () => {
      const index = firstSameRecognitionIndex(currentCategory(), currentItem());
      if (index >= 0) goTo(app.categoryIndex, index).catch(error => showNotice(error.message, true));
    });
    $("sameRecognitionMarkButton").addEventListener("click", () => {
      markSameRecognition().catch(error => showNotice(error.message, true));
    });
    $("previousButton").addEventListener("click", () => move(-1).catch(error => showNotice(error.message, true)));
    $("nextButton").addEventListener("click", () => move(1).catch(error => showNotice(error.message, true)));
    $("exportButton").addEventListener("click", () => exportPackage().catch(error => showNotice(error.message, true)));
    $("autoClickToggle").addEventListener("click", toggleAutoClick);
    $("autoClickSelect").addEventListener("change", event => {
      const value = Number(event.target.value);
      if (!Number.isFinite(value) || value <= 0) return;
      app.autoClickDelay = value;
      refreshAutoClickUI();
      if (app.autoClickEnabled) scheduleAutoClickTimer();
    });
    $("autoClickStatusSelect").addEventListener("change", event => {
      const value = Number(event.target.value);
      if (![1, 2, 3].includes(value)) return;
      app.autoClickStatus = value;
      refreshAutoClickUI();
      if (app.autoClickEnabled) scheduleAutoClickTimer();
    });
    $("zoomOut").addEventListener("click", () => setZoom(app.zoom - .25));
    $("zoomIn").addEventListener("click", () => setZoom(app.zoom + .25));
    $("zoomReset").addEventListener("click", () => {
      $("mainImage").style.transformOrigin = "center center";
      setZoom(app.defaultZoom);
    });
    $("imageViewport").addEventListener("wheel", event => {
      event.preventDefault();
      setZoom(app.zoom + (event.deltaY < 0 ? .1 : -.1), event);
    }, {passive: false});
    document.addEventListener("keydown", event => {
      const tagName = event.target?.tagName;
      if (tagName === "INPUT" || tagName === "SELECT" || tagName === "TEXTAREA" || app.saving) return;
      if (hotkeyMatches(event, "previous")) {
        event.preventDefault();
        if (app.autoClickEnabled) scheduleAutoClickTimer();
        move(-1).catch(error => showNotice(error.message, true));
      } else if (hotkeyMatches(event, "next")) {
        event.preventDefault();
        if (app.autoClickEnabled) scheduleAutoClickTimer();
        move(1).catch(error => showNotice(error.message, true));
      } else if (hotkeyMatches(event, "correct")) {
        event.preventDefault();
        if (app.autoClickEnabled) scheduleAutoClickTimer();
        saveCurrent(true, 1).catch(error => showNotice(error.message, true));
      } else if (hotkeyMatches(event, "problem")) {
        event.preventDefault();
        if (app.autoClickEnabled) scheduleAutoClickTimer();
        saveCurrent(true, 3).catch(error => showNotice(error.message, true));
      } else if (hotkeyMatches(event, "wrong")) {
        event.preventDefault();
        if (app.autoClickEnabled) scheduleAutoClickTimer();
        saveCurrent(true, 2).catch(error => showNotice(error.message, true));
      } else if (hotkeyMatches(event, "auto_click_toggle") && !event.repeat) {
        event.preventDefault();
        toggleAutoClick();
      }
    });
    $("mainImage").addEventListener("error", () => {
      $("mainImage").hidden = true;
      $("emptyImage").hidden = false;
    });
    $("mainImage").addEventListener("load", () => {
      if (app.zoom === app.defaultZoom) {
        $("mainImage").style.transformOrigin = "center center";
        setZoom(app.defaultZoom);
      }
    });

    fetch("/api/state")
      .then(response => response.json().then(result => {
        if (!response.ok) throw new Error(result.error || "读取失败");
        return result;
      }))
      .then(result => {
        applyState(result);
      })
      .catch(error => showNotice(error.message, true));
  </script>
</body>
</html>
"""


IMPORT_HTML_PAGE = r"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>导入离线包</title>
  <style>
    :root {
      color-scheme: dark;
      --bg: #10151b;
      --panel: #19222c;
      --panel-strong: #202c38;
      --line: #334150;
      --text: #eef3f8;
      --muted: #9eacba;
      --green: #57d38c;
      --red: #f06d73;
      --blue: #65a9e9;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      min-height: 100vh;
      background: var(--bg);
      color: var(--text);
      font: 14px/1.5 "Microsoft YaHei", "Segoe UI", sans-serif;
    }
    button, input { font: inherit; }
    button, .file-button {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-height: 38px;
      border: 1px solid var(--line);
      border-radius: 6px;
      background: var(--panel-strong);
      color: var(--text);
      cursor: pointer;
      padding: 8px 14px;
      text-decoration: none;
    }
    button:hover:not(:disabled), .file-button:hover { border-color: var(--blue); }
    button:disabled { cursor: not-allowed; opacity: .45; }
    .flow {
      display: grid;
      grid-template-rows: 54px minmax(0, 1fr);
      min-height: 100vh;
      padding: 10px;
      gap: 10px;
    }
    .topbar, .import-panel {
      border: 1px solid var(--line);
      background: var(--panel);
    }
    .topbar {
      display: flex;
      align-items: center;
      justify-content: space-between;
      padding: 8px 12px;
    }
    .brand { font-size: 16px; font-weight: 700; }
    .import-panel {
      display: grid;
      align-content: center;
      justify-items: center;
      gap: 18px;
      min-height: 0;
      padding: 28px;
    }
    .import-box {
      width: min(720px, 100%);
      display: grid;
      gap: 14px;
    }
    h1 { margin: 0; font-size: 24px; letter-spacing: 0; }
    .selected {
      min-height: 44px;
      padding: 11px 12px;
      border: 1px solid var(--line);
      background: #141c24;
      color: var(--muted);
      word-break: break-all;
    }
    .actions { display: flex; flex-wrap: wrap; gap: 10px; }
    .primary {
      border-color: var(--green);
      background: #173c2a;
      color: var(--green);
      font-weight: 700;
    }
    .notice {
      min-height: 24px;
      color: var(--muted);
      white-space: pre-wrap;
      word-break: break-word;
    }
    .notice.error { color: var(--red); }
    input[type="file"] { position: absolute; left: -9999px; }
  </style>
</head>
<body>
  <main class="flow">
    <header class="topbar">
      <div class="brand">巡视点位确认</div>
      <div>导入离线包</div>
    </header>
    <section class="import-panel">
      <div class="import-box">
        <h1>导入离线包</h1>
        <input id="packageInput" type="file" accept=".zip,application/zip,application/x-zip-compressed">
        <div id="selectedFile" class="selected">未选择文件</div>
        <div class="actions">
          <label class="file-button" for="packageInput">选择压缩包</label>
          <button id="uploadButton" class="primary" type="button" disabled>导入离线包</button>
          <button id="startButton" class="primary" type="button" hidden>开始标注</button>
        </div>
        <div id="notice" class="notice"></div>
      </div>
    </section>
  </main>
  <script>
    const input = document.getElementById("packageInput");
    const selectedFile = document.getElementById("selectedFile");
    const uploadButton = document.getElementById("uploadButton");
    const startButton = document.getElementById("startButton");
    const notice = document.getElementById("notice");
    let currentFile = null;

    function setNotice(message, isError = false) {
      notice.textContent = message || "";
      notice.classList.toggle("error", isError);
    }

    input.addEventListener("change", () => {
      currentFile = input.files && input.files[0] ? input.files[0] : null;
      selectedFile.textContent = currentFile ? `${currentFile.name} (${Math.ceil(currentFile.size / 1024)} KB)` : "未选择文件";
      uploadButton.disabled = !currentFile;
      startButton.hidden = true;
      setNotice("");
    });

    uploadButton.addEventListener("click", async () => {
      if (!currentFile) return;
      uploadButton.disabled = true;
      uploadButton.textContent = "导入中...";
      startButton.hidden = true;
      setNotice("正在导入离线包");
      try {
        const response = await fetch(`/api/import?filename=${encodeURIComponent(currentFile.name)}`, {
          method: "POST",
          headers: {"Content-Type": "application/octet-stream"},
          body: currentFile
        });
        const result = await response.json();
        if (!response.ok) throw new Error(result.error || "导入失败");
        setNotice(`导入完成: ${result.package || currentFile.name}`);
        startButton.hidden = false;
      } catch (error) {
        uploadButton.disabled = false;
        setNotice(error.message, true);
      } finally {
        uploadButton.textContent = "导入离线包";
      }
    });

    startButton.addEventListener("click", () => {
      window.location.href = "/";
    });
  </script>
</body>
</html>
"""


def render_export_page(files: ExportedFiles) -> str:
    package_name = html.escape(files.package.name)
    checklist_name = html.escape(files.checklist.name)
    package_size = html.escape(format_file_size(files.package))
    checklist_size = html.escape(format_file_size(files.checklist))
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>导出文件</title>
  <style>
    :root {{
      color-scheme: dark;
      --bg: #10151b;
      --panel: #19222c;
      --panel-strong: #202c38;
      --line: #334150;
      --text: #eef3f8;
      --muted: #9eacba;
      --green: #57d38c;
      --blue: #65a9e9;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      min-height: 100vh;
      background: var(--bg);
      color: var(--text);
      font: 14px/1.5 "Microsoft YaHei", "Segoe UI", sans-serif;
    }}
    button, a {{ font: inherit; }}
    .flow {{
      display: grid;
      grid-template-rows: 54px minmax(0, 1fr);
      min-height: 100vh;
      padding: 10px;
      gap: 10px;
    }}
    .topbar, .export-panel {{
      border: 1px solid var(--line);
      background: var(--panel);
    }}
    .topbar {{
      display: flex;
      align-items: center;
      justify-content: space-between;
      padding: 8px 12px;
    }}
    .brand {{ font-size: 16px; font-weight: 700; }}
    .export-panel {{
      display: grid;
      align-content: center;
      justify-items: center;
      min-height: 0;
      padding: 28px;
    }}
    .export-box {{
      width: min(760px, 100%);
      display: grid;
      gap: 14px;
    }}
    h1 {{ margin: 0; font-size: 24px; letter-spacing: 0; }}
    .file-row {{
      display: grid;
      grid-template-columns: minmax(0, 1fr) auto;
      align-items: center;
      gap: 12px;
      padding: 12px;
      border: 1px solid var(--line);
      background: #141c24;
    }}
    .file-name {{ font-weight: 700; word-break: break-all; }}
    .file-size {{ color: var(--muted); }}
    .download, .reset {{
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-height: 38px;
      border: 1px solid var(--green);
      border-radius: 6px;
      background: #173c2a;
      color: var(--green);
      cursor: pointer;
      padding: 8px 14px;
      font-weight: 700;
      text-decoration: none;
      white-space: nowrap;
    }}
    .download:hover, .reset:hover {{ border-color: var(--blue); }}
    .reset {{
      position: fixed;
      right: 18px;
      bottom: 18px;
    }}
  </style>
</head>
<body>
  <main class="flow">
    <header class="topbar">
      <div class="brand">巡视点位确认</div>
      <div>导出文件</div>
    </header>
    <section class="export-panel">
      <div class="export-box">
        <h1>导出文件</h1>
        <div class="file-row">
          <div>
            <div class="file-name">{package_name}</div>
            <div class="file-size">{package_size}</div>
          </div>
          <a class="download" href="/download/package">下载完整压缩包</a>
        </div>
        <div class="file-row">
          <div>
            <div class="file-name">{checklist_name}</div>
            <div class="file-size">{checklist_size}</div>
          </div>
          <a class="download" href="/download/checklist">下载清单</a>
        </div>
      </div>
    </section>
  </main>
  <button id="resetButton" class="reset" type="button">重新导入</button>
  <script>
    document.getElementById("resetButton").addEventListener("click", async () => {{
      await fetch("/api/reset", {{method: "POST"}});
      window.location.href = "/";
    }});
  </script>
</body>
</html>
"""


def safe_extract(zip_path: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    destination_resolved = destination.resolve()
    with zipfile.ZipFile(zip_path) as archive:
        for info in archive.infolist():
            raw_name = info.filename.replace("\\", "/")
            if not raw_name or raw_name == "/":
                continue
            pure = PurePosixPath(raw_name)
            if pure.is_absolute() or any(part == ".." for part in pure.parts):
                raise ValueError(f"ZIP 条目路径不安全: {zip_path} -> {info.filename}")
            target = (destination / Path(*pure.parts)).resolve()
            if not target.is_relative_to(destination_resolved):
                raise ValueError(f"ZIP 条目越界: {zip_path} -> {info.filename}")
            if info.is_dir() or raw_name.endswith("/"):
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with archive.open(info) as source, target.open("wb") as output:
                shutil.copyfileobj(source, output)


def load_json(path: Path) -> Any:
    with path.open("r", encoding="utf-8-sig") as file:
        return json.load(file)


def write_json(path: Path, value: Any) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as file:
        json.dump(value, file, ensure_ascii=False, indent=2)
        file.write("\n")


def strip_yaml_scalar(value: str) -> str:
    value = value.strip()
    if len(value) >= 2 and value[0] == value[-1] and value[0] in ("'", '"'):
        return value[1:-1].strip()
    return value


def load_save_problem_images(config_path: Path) -> bool:
    if not config_path.exists():
        logging.warning("未找到 config.yaml，m 模式默认保留问题图片: %s", config_path)
        return True

    keys = "|".join(re.escape(key) for key in SAVE_PROBLEM_KEYS)
    key_pattern = re.compile(rf"^\s*({keys})\s*:\s*(.*?)\s*$", re.IGNORECASE)
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0]
        match = key_pattern.match(line)
        if not match:
            continue
        value = strip_yaml_scalar(match.group(2)).lower()
        if value in ("y", "yes", "true", "1"):
            logging.info("m 模式配置：保留问题图片到 accuracy imagesData")
            return True
        if value in ("n", "no", "false", "0"):
            logging.info("m 模式配置：不保留问题图片到 accuracy imagesData")
            return False
        raise ValueError(f"config.yaml 的 {match.group(1)} 只能是 y/n: {value}")

    logging.info("config.yaml 未配置 save_problem_images，m 模式默认保留问题图片")
    return True


def load_yaml_bool(config_path: Path, key: str, default: bool) -> bool:
    if not config_path.exists():
        return default

    key_pattern = re.compile(rf"^\s*{re.escape(key)}\s*:\s*(.*?)\s*$", re.IGNORECASE)
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0]
        match = key_pattern.match(line)
        if not match:
            continue
        value = strip_yaml_scalar(match.group(1)).lower()
        if value in ("y", "yes", "true", "1"):
            return True
        if value in ("n", "no", "false", "0"):
            return False
        raise ValueError(f"config.yaml 的 {key} 只能是 y/n: {value}")
    return default


def load_yaml_yes(config_path: Path, key: str) -> bool:
    if not config_path.exists():
        return False

    key_pattern = re.compile(rf"^\s*{re.escape(key)}\s*:\s*(.*?)\s*$", re.IGNORECASE)
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0]
        match = key_pattern.match(line)
        if not match:
            continue
        return strip_yaml_scalar(match.group(1)).lower() == "yes"
    return False


def load_yaml_scalar(config_path: Path, key: str, default: str = "") -> str:
    if not config_path.exists():
        return default

    key_pattern = re.compile(rf"^\s*{re.escape(key)}\s*:\s*(.*?)\s*$", re.IGNORECASE)
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0]
        match = key_pattern.match(line)
        if match:
            return strip_yaml_scalar(match.group(1))
    return default


def load_yaml_list(config_path: Path, key: str, keep_empty: bool = False) -> list[str]:
    if not config_path.exists():
        return []

    values: list[str] = []
    in_list = False
    key_pattern = re.compile(rf"^\s*{re.escape(key)}\s*:\s*(.*?)\s*$", re.IGNORECASE)
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0].rstrip()
        if not line.strip():
            continue
        match = key_pattern.match(line)
        if match:
            in_list = True
            inline_raw = match.group(1).strip()
            inline_value = strip_yaml_scalar(inline_raw)
            if inline_value or (keep_empty and inline_raw):
                values.append(inline_value)
            continue
        if in_list:
            item_match = re.match(r"^\s*-\s*(.*?)\s*$", line)
            if item_match:
                item_raw = item_match.group(1).strip()
                value = strip_yaml_scalar(item_raw)
                if value or (keep_empty and item_raw):
                    values.append(value)
                continue
            if re.match(r"^\S[^:]*\s*:", line):
                break
    return values


def load_point_name_end_config(config_path: Path) -> tuple[bool, int, str]:
    end_open = load_yaml_bool(config_path, POINT_NAME_END_OPEN_KEY, False)
    raw_characters = load_yaml_scalar(config_path, POINT_NAME_END_CHARACTERS_KEY, "0")
    try:
        end_characters = max(0, int(raw_characters or 0))
    except ValueError as exc:
        raise ValueError(f"config.yaml 的 {POINT_NAME_END_CHARACTERS_KEY} 必须是整数: {raw_characters}") from exc
    end_color = load_yaml_scalar(config_path, POINT_NAME_END_CHARACTER_COLOR_KEY, "").strip()
    if end_open and end_characters <= 0:
        logging.info("config.yaml 的 %s 已启用但字符数为 0", POINT_NAME_END_OPEN_KEY)
    if end_color and end_color.lower() not in POINT_NAME_END_COLORS and not end_color.startswith("#"):
        logging.info("config.yaml 的 %s 使用自定义颜色: %s", POINT_NAME_END_CHARACTER_COLOR_KEY, end_color)
    return end_open, end_characters, end_color


def load_default_zoom(config_path: Path) -> float:
    if not config_path.exists():
        return 1.0

    key_pattern = re.compile(rf"^\s*{re.escape(DEFAULT_ZOOM_KEY)}\s*:\s*(.*?)\s*$", re.IGNORECASE)
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0]
        match = key_pattern.match(line)
        if not match:
            continue
        value = strip_yaml_scalar(match.group(1))
        try:
            zoom = float(value)
        except ValueError as exc:
            raise ValueError(f"config.yaml 的 {DEFAULT_ZOOM_KEY} 必须是数字: {value}") from exc
        if zoom <= 0:
            raise ValueError(f"config.yaml 的 {DEFAULT_ZOOM_KEY} 必须大于 0: {value}")
        return max(0.25, min(4.0, zoom))
    return 1.0


def load_yaml_int_list(config_path: Path, key: str, defaults: list[int]) -> list[int]:
    values: list[int] = []
    for item in load_yaml_list(config_path, key):
        try:
            value = int(item)
        except ValueError as exc:
            raise ValueError(f"config.yaml 的 {key} 必须是整数毫秒: {item}") from exc
        if value <= 0:
            raise ValueError(f"config.yaml 的 {key} 必须大于 0: {item}")
        values.append(value)
    return values or defaults


def load_yaml_mapping(config_path: Path, key: str) -> dict[str, str]:
    if not config_path.exists():
        return {}

    values: dict[str, str] = {}
    in_block = False
    base_indent = 0
    key_pattern = re.compile(rf"^(\s*){re.escape(key)}\s*:\s*(.*?)\s*$", re.IGNORECASE)
    item_pattern = re.compile(r"^(\s*)([^:#]+?)\s*:\s*(.*?)\s*$")
    for raw_line in config_path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.split("#", 1)[0].rstrip()
        if not line.strip():
            continue
        match = key_pattern.match(line)
        if match:
            in_block = True
            base_indent = len(match.group(1))
            continue
        if not in_block:
            continue
        item_match = item_pattern.match(line)
        if not item_match:
            continue
        indent = len(item_match.group(1))
        if indent <= base_indent:
            break
        name = item_match.group(2).strip()
        value = strip_yaml_scalar(item_match.group(3))
        if name and value:
            values[name] = value
    return values


def _normalized_config_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def load_solution_rules(config_path: Path) -> list[tuple[str, str, str]]:
    if not config_path.exists():
        logging.info("未找到解决方案配置，清单解决方案留空: %s", config_path)
        return []

    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("读取解决方案配置需要 PyYAML") from exc

    data = yaml.safe_load(config_path.read_text(encoding="utf-8-sig")) or {}
    if not isinstance(data, dict):
        raise ValueError("config.yaml 根节点必须是对象")

    rules_root = next(
        (value for key, value in data.items() if _normalized_config_key(key) == _normalized_config_key(SOLUTION_RULES_KEY)),
        None,
    )
    if rules_root is None:
        logging.info("config.yaml 未配置 %s，清单解决方案留空", SOLUTION_RULES_KEY)
        return []

    rules: list[tuple[str, str, str]] = []
    situation_keys = {"situation", "assesssituation", "judgment", "judgement"}
    recognition_keys = {
        "recognitionresult",
        "recognitioresult",
        "result",
        "recognition",
    }
    resolve_keys = {"resolve", "solution", "resolution"}

    def walk(node: Any, situation: str = "", recognition: str = "") -> None:
        if isinstance(node, list):
            for item in node:
                walk(item, situation, recognition)
            return
        if not isinstance(node, dict):
            return

        current_situation = situation
        current_recognition = recognition
        resolved = ""
        for key, value in node.items():
            normalized_key = _normalized_config_key(key)
            if normalized_key in situation_keys and not isinstance(value, (dict, list)):
                current_situation = str(value or "").strip()
            elif normalized_key in recognition_keys:
                if isinstance(value, (dict, list)):
                    walk(value, current_situation, current_recognition)
                else:
                    current_recognition = str(value or "").strip()
            elif normalized_key in resolve_keys and not isinstance(value, (dict, list)):
                resolved = str(value or "").strip()

        if current_situation and current_recognition and resolved:
            rules.append((current_situation, current_recognition, resolved))

        for key, value in node.items():
            normalized_key = _normalized_config_key(key)
            if normalized_key in situation_keys or normalized_key in recognition_keys or normalized_key in resolve_keys:
                continue
            if isinstance(value, (dict, list)):
                walk(value, current_situation, current_recognition)

    walk(rules_root)
    logging.info("读取解决方案配置 %s 条: %s", len(rules), config_path)
    return rules


def load_solution_options(config_path: Path) -> list[str]:
    options = load_yaml_list(config_path, SOLUTION_OPTIONAL_KEY)
    if not config_path.exists():
        logging.info("未找到解决方案下拉配置，清单解决方案不启用下拉: %s", config_path)
    elif options:
        logging.info("读取解决方案下拉配置 %s 条: %s", len(options), config_path)
    else:
        logging.info("config.yaml 未配置 %s，清单解决方案不启用下拉", SOLUTION_OPTIONAL_KEY)
    return options


def resolve_solution(
    rules: list[tuple[str, str, str]],
    situation: str,
    recognition: str,
) -> str:
    situation = situation.strip()
    recognition = recognition.strip()
    for rule_situation, rule_recognition, solution in rules:
        if rule_situation.strip() == situation and rule_recognition.strip() == recognition:
            return solution
    if situation == "算法不准" and recognition == "未识别":
        return "算法涉入"
    return ""


def load_responsible_rules(config_path: Path) -> list[tuple[str, str]]:
    if not config_path.exists():
        logging.info("未找到负责人配置，清单负责人留空: %s", config_path)
        return []

    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("读取负责人配置需要 PyYAML") from exc

    data = yaml.safe_load(config_path.read_text(encoding="utf-8-sig")) or {}
    if not isinstance(data, dict):
        raise ValueError("config.yaml 根节点必须是对象")
    rules_root = next(
        (
            value
            for key, value in data.items()
            if _normalized_config_key(key) in {
                _normalized_config_key(RESPONSIBLE_RULES_KEY),
                "responsible",
                "owner_rules",
            }
        ),
        None,
    )
    if rules_root is None:
        logging.info("config.yaml 未配置 %s，清单负责人留空", RESPONSIBLE_RULES_KEY)
        return []

    rules: list[tuple[str, str]] = []

    def add_rule(keyword: Any, owner: Any) -> None:
        keyword_text = str(keyword or "").strip()
        owner_text = str(owner or "").strip()
        if not keyword_text or not owner_text:
            return
        if owner_text not in RESPONSIBLE_VALUES:
            raise ValueError(
                f"负责人只能是 {sorted(RESPONSIBLE_VALUES)}，收到: {owner_text}"
            )
        rules.append((keyword_text, owner_text))

    if isinstance(rules_root, dict):
        for keyword, owner in rules_root.items():
            if isinstance(owner, dict):
                add_rule(keyword, owner.get("responsible") or owner.get("owner"))
            else:
                add_rule(keyword, owner)
    elif isinstance(rules_root, list):
        for item in rules_root:
            if not isinstance(item, dict):
                continue
            keyword = item.get("keyword") or item.get("recognition") or item.get("match")
            owner = item.get("responsible") or item.get("owner")
            add_rule(keyword, owner)
    else:
        raise ValueError(f"config.yaml 的 {RESPONSIBLE_RULES_KEY} 必须是对象或数组")

    logging.info("读取负责人配置 %s 条: %s", len(rules), config_path)
    return rules


def resolve_responsible(
    rules: list[tuple[str, str]],
    recognition: str,
) -> str:
    for keyword, owner in rules:
        if keyword in recognition:
            return owner
    return ""


def summarize_algorithm_counts(rows: list[dict[str, Any]], label: str) -> str:
    counts = Counter(
        str(record.get("algoType") or "").strip() or "未命名算法"
        for record in rows
    )
    if not counts:
        return f"{label}:{{无}}"
    parts = [
        f"{algo_name}:{count}"
        for algo_name, count in sorted(counts.items(), key=lambda item: (-item[1], item[0]))
    ]
    return f"{label}:{{" + ";".join(parts) + "}"


def _format_archive_time(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        try:
            timestamp = float(value)
            if timestamp > 10_000_000_000:
                timestamp /= 1000
            return datetime.fromtimestamp(timestamp).strftime("%Y-%m-%d %H:%M:%S")
        except (OverflowError, OSError, ValueError):
            return ""
    text = str(value).strip()
    if not text:
        return ""
    if re.fullmatch(r"\d{14}", text):
        return f"{text[:4]}-{text[4:6]}-{text[6:8]} {text[8:10]}:{text[10:12]}:{text[12:]}"
    if re.fullmatch(r"\d{8}", text):
        return f"{text[:4]}-{text[4:6]}-{text[6:]}"
    return text


def load_archive_audit_time(roots: list[Path], fallback: str) -> str:
    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("读取压缩包 YAML 需要 PyYAML") from exc

    exact_keys = {
        "audit_time",
        "audittime",
        "review_time",
        "reviewtime",
        "assessment_time",
        "assessmenttime",
        "审算时间",
        "审核时间",
    }
    candidate_keys = ("audit", "review", "assessment", "审算", "审核", "审查")

    def find_value(node: Any) -> str:
        if isinstance(node, dict):
            for key, value in node.items():
                normalized = _normalized_config_key(key)
                key_text = str(key).lower()
                if normalized in exact_keys or (
                    any(token in normalized or token in key_text for token in candidate_keys)
                    and ("time" in normalized or "时间" in key_text or "date" in normalized)
                ):
                    result = _format_archive_time(value)
                    if result:
                        return result
            for value in node.values():
                result = find_value(value)
                if result:
                    return result
        elif isinstance(node, list):
            for value in node:
                result = find_value(value)
                if result:
                    return result
        return ""

    for root in roots:
        for path in sorted(root.rglob("*"), key=lambda item: str(item)):
            if not path.is_file() or path.suffix.lower() not in (".yaml", ".yml"):
                continue
            try:
                value = yaml.safe_load(path.read_text(encoding="utf-8-sig"))
            except (OSError, UnicodeError, yaml.YAMLError) as exc:
                logging.warning("读取压缩包 YAML 失败: %s (%s)", path, exc)
                continue
            result = find_value(value)
            if result:
                logging.info("从压缩包 YAML 读取审算时间: %s -> %s", path, result)
                return result

    logging.warning("压缩包内未找到审算时间 YAML，使用回退值: %s", fallback)
    return fallback


def load_web_hotkeys(config_path: Path) -> dict[str, str]:
    hotkeys = DEFAULT_WEB_HOTKEYS.copy()
    configured = load_yaml_mapping(config_path, WEB_HOTKEYS_KEY)
    for action, key in configured.items():
        if action in hotkeys:
            hotkeys[action] = key
    return hotkeys


def load_algo_type_labels(path: Path) -> dict[str, str]:
    if not path.is_file():
        logging.info("未找到算法映射配置，清单使用原始 algoType: %s", path)
        return {}
    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("读取算法映射配置需要 PyYAML") from exc

    data = yaml.safe_load(path.read_text(encoding="utf-8-sig")) or {}
    if not isinstance(data, dict):
        raise ValueError("config.yaml 根节点必须是对象")

    rules = next(
        (
            value
            for key, value in data.items()
            if _normalized_config_key(key) == _normalized_config_key(ALGO_MAPPING_RULES_KEY)
        ),
        None,
    )
    if rules is None:
        logging.info("config.yaml 未配置 %s，清单使用原始 algoType", ALGO_MAPPING_RULES_KEY)
        return {}
    if not isinstance(rules, list):
        raise ValueError(f"config.yaml 的 {ALGO_MAPPING_RULES_KEY} 必须是数组")

    labels: dict[str, str] = {}
    for rule in rules:
        if not isinstance(rule, dict):
            raise ValueError(f"config.yaml 的 {ALGO_MAPPING_RULES_KEY} 项必须是对象")
        target = str(rule.get("target") or "").strip()
        matches = rule.get("matches") or []
        if isinstance(matches, str):
            matches = [matches]
        if not isinstance(matches, list):
            raise ValueError("config.yaml 的 mapping_rules.matches 必须是数组")
        if not target:
            continue
        for match in matches:
            algo_id = str(match or "").strip()
            if algo_id:
                labels[algo_id] = target

    logging.info("读取算法映射 %s 条: %s", len(labels), path)
    return labels


def load_result2_xlsx_path(config_path: Path) -> Path | None:
    if not config_path.is_file():
        return None
    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("读取 result-2.xlsx 配置需要 PyYAML") from exc

    data = yaml.safe_load(config_path.read_text(encoding="utf-8-sig")) or {}
    if not isinstance(data, dict):
        raise ValueError("config.yaml 根节点必须是对象")
    configured = next(
        (
            value
            for key, value in data.items()
            if _normalized_config_key(key) in {
                _normalized_config_key(RESULT2_XLSX_PATH_KEY),
                "result2xlsx",
                "result2xlsxpath",
                "resultxlsxpath",
            }
        ),
        None,
    )
    if configured is None:
        return None
    path = Path(str(configured).strip()).expanduser()
    if not path.is_absolute():
        path = config_path.parent / path
    return path


def load_result2_algorithm_aliases(config_path: Path) -> dict[str, set[str]]:
    if not config_path.is_file():
        return {}
    try:
        import yaml
    except ImportError as exc:
        raise RuntimeError("读取 result-2.xlsx 算法映射需要 PyYAML") from exc

    data = yaml.safe_load(config_path.read_text(encoding="utf-8-sig")) or {}
    if not isinstance(data, dict):
        raise ValueError("config.yaml 根节点必须是对象")
    rules = next(
        (
            value
            for key, value in data.items()
            if _normalized_config_key(key) in {
                _normalized_config_key(RESULT2_ALGORITHM_MAPPING_RULES_KEY),
                "result2algorithmmapping",
                "result2algorithmaliases",
                "algorithmcompatibility",
            }
        ),
        None,
    )
    if rules is None:
        return {}
    if not isinstance(rules, list):
        raise ValueError(f"config.yaml 的 {RESULT2_ALGORITHM_MAPPING_RULES_KEY} 必须是数组")

    aliases: dict[str, set[str]] = {}
    for rule in rules:
        if not isinstance(rule, dict):
            raise ValueError(f"config.yaml 的 {RESULT2_ALGORITHM_MAPPING_RULES_KEY} 项必须是对象")
        target = str(rule.get("target") or "").strip()
        matches = rule.get("matches") or []
        if isinstance(matches, str):
            matches = [matches]
        if not target or not isinstance(matches, list):
            continue
        aliases.setdefault(normalize_algorithm_name(target), set()).update(
            str(match).strip() for match in matches if str(match).strip()
        )
    return aliases


def normalize_algorithm_name(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip()


def result2_algorithm_names(
    result2_path: Path | None,
    station_name: str,
) -> list[str] | None:
    if result2_path is None:
        logging.info("未配置 result-2.xlsx 路径，跳过算法类别比对")
        return None
    if not result2_path.is_file():
        logging.warning("result-2.xlsx 不存在，跳过算法类别比对: %s", result2_path)
        return None
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 result-2.xlsx 需要 openpyxl") from exc

    station_key = normalize_algorithm_name(station_name)
    if not station_key:
        logging.warning("站点名称为空，跳过 result-2.xlsx 算法类别比对")
        return None
    names: list[str] = []
    seen_names: set[str] = set()
    matched_rows = 0
    workbook = load_workbook(result2_path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            algorithm_column_index = RESULT2_ALGORITHM_COLUMN_INDEX
            for row in sheet.iter_rows(values_only=True):
                header_index = next(
                    (
                        index
                        for index, value in enumerate(row)
                        if normalize_algorithm_name(value) == "算法小类"
                    ),
                    None,
                )
                if header_index is not None:
                    algorithm_column_index = header_index
                    continue
                values = [str(value or "") for value in row]
                if not any(station_key in normalize_algorithm_name(value) for value in values):
                    continue
                matched_rows += 1
                if len(row) <= algorithm_column_index:
                    continue
                algorithm = normalize_algorithm_name(row[algorithm_column_index])
                if algorithm and algorithm not in seen_names:
                    names.append(algorithm)
                    seen_names.add(algorithm)
    finally:
        workbook.close()
    if not matched_rows:
        logging.warning("result-2.xlsx 未找到对应站点: %s", station_name)
        return None
    logging.info(
        "读取 result-2.xlsx 算法类别: station=%s count=%s path=%s",
        station_name,
        len(names),
        result2_path,
    )
    return names


def checklist_table_point_names(workbook_path: Path | None) -> list[str]:
    if workbook_path is None or not workbook_path.is_file():
        return []
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取清单 xlsx 需要 openpyxl") from exc

    point_names: list[str] = []
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            point_column_index: int | None = None
            for row in sheet.iter_rows(values_only=True):
                values = list(row)
                header_index = next(
                    (
                        index
                        for index, value in enumerate(values)
                        if str(value or "").strip() == "点位名称"
                    ),
                    None,
                )
                if header_index is not None:
                    point_column_index = header_index
                    continue
                if point_column_index is not None:
                    if len(values) <= point_column_index:
                        continue
                    point_name = str(values[point_column_index] or "").strip()
                    if point_name:
                        point_names.append(point_name)
                    continue
                if (
                    len(values) > 1
                    and isinstance(values[0], (int, float))
                    and not isinstance(values[0], bool)
                ):
                    point_name = str(values[1] or "").strip()
                    if point_name and point_name != "点位名称":
                        point_names.append(point_name)
    finally:
        workbook.close()

    logging.info("读取清单表格点位: count=%s path=%s", len(point_names), workbook_path)
    return point_names


def checklist_algorithm_records(
    accuracy: dict[str, Any],
    algo_type_labels: dict[str, str],
) -> list[dict[str, str]]:
    statistics = accuracy.get("data", {}).get("statistic") or []
    records: list[dict[str, str]] = []
    seen_algo_types: set[str] = set()
    for stat in statistics:
        algo_type = str(stat.get("algoType") or "").strip()
        if not algo_type or algo_type in seen_algo_types:
            continue
        seen_algo_types.add(algo_type)
        label = algo_type_labels.get(algo_type, algo_type)
        records.append({"label": label, "algoType": algo_type})
    return records


def formatted_checklist_algorithm(record: dict[str, str]) -> str:
    label = str(record.get("label") or "").strip()
    algo_type = str(record.get("algoType") or "").strip()
    if label and algo_type and label != algo_type:
        return f"{label}（{algo_type}）"
    return label or algo_type


def compare_checklist_and_result2_algorithms(
    checklist_records: list[dict[str, str]],
    result2_algorithms: list[str],
    aliases: dict[str, set[str]],
) -> tuple[list[str], list[str]]:
    checklist_candidates: list[tuple[dict[str, str], set[str]]] = []
    for record in checklist_records:
        target = normalize_algorithm_name(record.get("label"))
        candidates = {target}
        candidates.update(normalize_algorithm_name(value) for value in aliases.get(target, set()))
        checklist_candidates.append((record, {candidate for candidate in candidates if candidate}))

    result2_only: list[str] = []
    for result_algorithm in result2_algorithms:
        if not any(result_algorithm in candidates for _, candidates in checklist_candidates):
            result2_only.append(result_algorithm)

    checklist_only: list[str] = []
    for record, candidates in checklist_candidates:
        if not any(
            result_algorithm in candidates
            for result_algorithm in result2_algorithms
        ):
            checklist_only.append(formatted_checklist_algorithm(record))
    return result2_only, checklist_only


def checklist_algorithm_names(
    accuracy: dict[str, Any],
    algo_type_labels: dict[str, str],
) -> list[str]:
    algorithm_names: list[str] = []
    seen_algorithm_names: set[str] = set()
    for record in checklist_algorithm_records(accuracy, algo_type_labels):
        algorithm_name = record["label"]
        if algorithm_name not in seen_algorithm_names:
            seen_algorithm_names.add(algorithm_name)
            algorithm_names.append(algorithm_name)
    return algorithm_names


def find_single_file(root: Path, name: str) -> Path:
    matches = [path for path in root.rglob(name) if path.is_file()]
    if not matches:
        raise FileNotFoundError(f"未找到 {name}: {root}")
    if len(matches) > 1:
        raise ValueError(f"找到多个 {name}: {matches}")
    return matches[0]


def zip_dir_contents(source_dir: Path, target_zip: Path) -> None:
    target_zip.parent.mkdir(parents=True, exist_ok=True)
    if target_zip.exists():
        target_zip.unlink()
    with zipfile.ZipFile(target_zip, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(source_dir.rglob("*"), key=lambda item: item.relative_to(source_dir).as_posix()):
            if path.is_file():
                archive.write(path, path.relative_to(source_dir).as_posix())


def nonconflicting_path(path: Path) -> Path:
    if not path.exists():
        return path
    counter = 1
    while True:
        candidate = path.with_name(f"{path.stem} -{counter}{path.suffix}")
        if not candidate.exists():
            return candidate
        counter += 1


def format_file_size(path: Path) -> str:
    try:
        size = path.stat().st_size
    except OSError:
        return "-"
    units = ("B", "KB", "MB", "GB")
    value = float(size)
    for unit in units:
        if value < 1024 or unit == units[-1]:
            if unit == "B":
                return f"{int(value)} {unit}"
            return f"{value:.1f} {unit}"
        value /= 1024
    return f"{size} B"


def safe_upload_filename(filename: str) -> str:
    name = Path(filename or "offline_package.zip").name
    stem = sanitize_stem(Path(name).stem, "offline_package", max_length=100)
    return f"{stem}.zip"


def is_complete_outer_zip(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            names = sorted(
                info.filename
                for info in archive.infolist()
                if not info.is_dir() and info.filename
            )
    except (OSError, zipfile.BadZipFile):
        return False
    return len(names) == 2 and all(OUTER_INNER_PATTERN.match(name) for name in names)


def prepare_package_for_review(
    package_path: Path,
    output_dir: Path,
    config_path: Path,
    force_rebuild: bool = False,
    database_path: Path | None = None,
) -> Path:
    if not force_rebuild and is_complete_outer_zip(package_path):
        logging.info("m 模式上传完整外层 ZIP，直接进入标注: %s", package_path)
        return package_path

    from rebuild_sources import load_bool_config, load_keywords, process_outer_zip

    run_dt = datetime.now().replace(microsecond=0)
    new_task_id = run_dt.strftime("%Y%m%d%H%M%S")
    logging.info("m 模式上传原始离线包，先执行 Y 模式整理: %s", package_path)
    result = process_outer_zip(
        outer_zip=package_path,
        output_dir=output_dir,
        keywords=load_keywords(config_path),
        empty_recognition_last=load_bool_config(config_path, "empty_recognition_last", True),
        cluster_by_recognition=load_bool_config(config_path, "cluster_by_recognition", False),
        new_task_id=new_task_id,
        run_dt=run_dt,
        split_limit=None,
        annotation_db=AnnotationDatabase(database_path) if database_path else None,
    )
    if len(result.output_zips) != 1:
        raise ValueError(f"上传离线包整理后输出数量异常: {len(result.output_zips)}")
    logging.info("m 模式上传离线包整理完成: %s", result.output_zips[0])
    return result.output_zips[0]


def find_input_package(output_dir: Path, package: Path | None) -> Path:
    if package is not None:
        candidate = package.expanduser().resolve()
        if candidate.is_dir():
            zips = [path for path in candidate.glob("*.zip") if is_complete_outer_zip(path)]
            if not zips:
                raise FileNotFoundError(f"目录中没有完整外层 ZIP: {candidate}")
            return max(zips, key=lambda path: path.stat().st_mtime_ns)
        if not candidate.is_file() or not is_complete_outer_zip(candidate):
            raise ValueError(f"不是完整外层 ZIP（应只包含 accuracy_*.zip 和 image_result_*.zip）: {candidate}")
        return candidate

    candidates = [
        path
        for path in output_dir.rglob("*.zip")
        if path.is_file() and is_complete_outer_zip(path)
    ]
    if not candidates:
        raise FileNotFoundError(
            f"output 中没有完整外层 ZIP: {output_dir}；请先执行 Y 模式，或使用 --package 指定"
        )
    return max(candidates, key=lambda path: path.stat().st_mtime_ns)


def choose_local_package(output_dir: Path, package: Path | None) -> Path:
    if package is not None:
        return find_input_package(output_dir, package)

    candidates = sorted(
        (
            path
            for path in output_dir.rglob("*.zip")
            if path.is_file() and is_complete_outer_zip(path)
        ),
        key=lambda path: path.stat().st_mtime_ns,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError(
            "output 中没有完整外层 ZIP；请先执行 Y 模式，或使用 --package 指定"
        )
    if not sys.stdin.isatty():
        return candidates[0]

    print("本地模式选择离线包：")
    for index, path in enumerate(candidates, 1):
        print(f"  {index}. {path}")
    selection = input("请选择序号，回车默认最新 [1] ").strip()
    index = 1 if not selection else int(selection)
    if index < 1 or index > len(candidates):
        raise ValueError(f"本地模式离线包序号超出范围: {index}")
    selected = candidates[index - 1]
    logging.info("m 模式本地模式选择离线包: %s", selected)
    return selected


def find_inner_zips(root: Path) -> tuple[Path, Path]:
    entries = [path for path in root.iterdir() if path.is_file() and path.suffix.lower() == ".zip"]
    accuracy = [path for path in entries if path.name.lower().startswith("accuracy_")]
    image = [path for path in entries if path.name.lower().startswith("image_result_")]
    if len(accuracy) != 1 or len(image) != 1:
        raise ValueError(f"完整外层 ZIP 内层包数量错误: {root}")
    return accuracy[0], image[0]


def status_value(value: Any) -> int:
    try:
        value = int(value)
    except (TypeError, ValueError):
        return 0
    return value if value in STATUS_NAMES else 0


def percentage_value(numerator: int, denominator: int) -> float | int:
    if not denominator:
        return 0
    return int(numerator * 100 / denominator)


def sanitize_stem(value: str, fallback: str, max_length: int = 120) -> str:
    text = value.strip() or fallback
    text = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "-", text)
    text = re.sub(r"\s+", " ", text)
    text = text.strip(" ._-")
    if not text:
        text = fallback
    return text[:max_length].rstrip(" ._-") or fallback


def unique_filename(stem: str, suffix: str, used: set[str]) -> str:
    suffix = suffix if suffix.startswith(".") else f".{suffix}"
    candidate = f"{stem}{suffix}"
    counter = 2
    while candidate.lower() in used:
        candidate = f"{stem}_{counter}{suffix}"
        counter += 1
    used.add(candidate.lower())
    return candidate


def selected_filename_for_item(item: dict[str, Any], original_name: str) -> str:
    suffix = Path(original_name).suffix.lower() or ".jpg"
    recognition = str(item.get("recognition") or "").strip()
    point_name = str(item.get("pointName") or "").strip()
    point_code = str(item.get("pointCode") or "").strip()
    fallback = Path(original_name).stem
    parts = [part for part in (recognition, point_name or point_code) if part]
    stem = sanitize_stem("-".join(parts), fallback)
    return f"{stem}{suffix}"


def session_package_key(input_package: Path) -> str:
    return hashlib.sha1(str(input_package.resolve()).encode("utf-8")).hexdigest()[:16]


def reviewed_output_stem(input_package: Path) -> str:
    return input_package.stem.replace("未审核", "")


def attachment_output_stem(input_package: Path) -> str:
    text = reviewed_output_stem(input_package)
    text = re.sub(r"^\d+[_-]", "", text)
    text = re.sub(r"(?i)[_-]?m[_-]?\d{14}$", "", text)
    text = re.sub(r"[_-]?\d{14}$", "", text)
    text = re.sub(r"(巡检|巡视|静默).*?$", "", text)
    text = text.replace("_", "-").strip(" -_*")
    return sanitize_stem(text, reviewed_output_stem(input_package))


def dated_attachment_output_stem(input_package: Path) -> str:
    return f"{datetime.now():%m%d}-{attachment_output_stem(input_package)}"


def calculate_statistics(stat: dict[str, Any], report_time: int | None = None) -> dict[str, Any]:
    images = stat.get("imagesData")
    if not isinstance(images, list):
        raise ValueError(f"imagesData 不是数组: {stat.get('algoType', '')}")

    correct = sum(1 for item in images if status_value(item.get("auditStatus")) == 1)
    inaccuracy = sum(1 for item in images if status_value(item.get("auditStatus")) == 2)
    blur = sum(1 for item in images if status_value(item.get("auditStatus")) == 3)
    detected = correct + inaccuracy
    total = len(images)

    stat["pointTotal"] = total
    stat["correct"] = correct
    stat["inaccuracy"] = inaccuracy
    stat["blur"] = blur
    stat["accuracyRate"] = percentage_value(correct, detected)
    stat["detectionRate"] = percentage_value(detected, total)
    if report_time is not None:
        stat["latestReportTime"] = report_time
    return stat


def prepare_accuracy_for_export(
    accuracy: dict[str, Any],
    report_time: int,
    save_problem_images: bool,
) -> dict[str, Any]:
    exported = copy.deepcopy(accuracy)
    statistics = exported.get("data", {}).get("statistic")
    if not isinstance(statistics, list):
        raise ValueError("accuracy.json 缺少 data.statistic 数组")

    for stat in statistics:
        if not isinstance(stat, dict):
            raise ValueError("statistic 条目不是对象")
        calculate_statistics(stat, report_time)
        if not save_problem_images:
            images = stat.get("imagesData")
            if not isinstance(images, list):
                raise ValueError(f"imagesData 不是数组: {stat.get('algoType', '')}")
            stat["imagesData"] = [
                item
                for item in images
                if status_value(item.get("auditStatus")) in (1, 2)
            ]
    return exported


def build_state(accuracy: dict[str, Any]) -> list[dict[str, Any]]:
    statistics = accuracy.get("data", {}).get("statistic")
    if not isinstance(statistics, list):
        raise ValueError("accuracy.json 缺少 data.statistic 数组")

    categories: list[dict[str, Any]] = []
    for stat in statistics:
        if not isinstance(stat, dict):
            raise ValueError("statistic 条目不是对象")
        algo_type = str(stat.get("algoType") or "").strip()
        images = stat.get("imagesData")
        if not algo_type or not isinstance(images, list):
            raise ValueError(f"statistic 数据无效: {stat}")
        normalized_images: list[dict[str, Any]] = []
        for item in images:
            if not isinstance(item, dict):
                raise ValueError(f"imagesData 条目不是对象: {algo_type}")
            normalized = copy.deepcopy(item)
            normalized["auditStatus"] = status_value(normalized.get("auditStatus"))
            normalized_images.append(normalized)
        completed = sum(1 for item in normalized_images if item["auditStatus"] != 0)
        correct = sum(1 for item in normalized_images if item["auditStatus"] == 1)
        inaccuracy = sum(1 for item in normalized_images if item["auditStatus"] == 2)
        blur = sum(1 for item in normalized_images if item["auditStatus"] == 3)
        detected = correct + inaccuracy
        total = len(normalized_images)
        categories.append({
            "algoType": algo_type,
            "images": normalized_images,
            "completed": completed,
            "pointTotal": total,
            "correct": correct,
            "inaccuracy": inaccuracy,
            "blur": blur,
            "accuracyRate": percentage_value(correct, detected),
            "detectionRate": percentage_value(detected, total),
            "latestReportTime": stat.get("latestReportTime"),
        })
    return categories


def apply_mark(accuracy: dict[str, Any], algo_type: str, index: int, audit_status: int) -> dict[str, Any]:
    categories = accuracy.get("data", {}).get("statistic")
    if not isinstance(categories, list):
        raise ValueError("accuracy.json 缺少 data.statistic 数组")
    if audit_status not in (0, 1, 2, 3):
        raise ValueError(f"auditStatus 无效: {audit_status}")

    for stat in categories:
        if not isinstance(stat, dict) or str(stat.get("algoType") or "") != algo_type:
            continue
        images = stat.get("imagesData")
        if not isinstance(images, list) or index < 0 or index >= len(images):
            raise IndexError(f"点位索引无效: {algo_type}/{index}")
        images[index]["auditStatus"] = audit_status
        calculate_statistics(stat)
        state_category = build_state(accuracy)[
            next(i for i, item in enumerate(categories) if item is stat)
        ]
        return state_category
    raise KeyError(f"不存在的 algoType: {algo_type}")


def all_marked(accuracy: dict[str, Any]) -> bool:
    categories = build_state(accuracy)
    return bool(categories) and all(item["auditStatus"] != 0 for category in categories for item in category["images"])


def export_reviewed_package(
    input_package: Path,
    output_root: Path,
    accuracy: dict[str, Any],
    save_problem_images: bool,
) -> Path:
    timestamp = datetime.now().replace(microsecond=0)
    stamp = timestamp.strftime("%Y%m%d%H%M%S")
    output_stem = reviewed_output_stem(input_package)
    output_dir = nonconflicting_path(output_root / f"{output_stem}_m_{stamp}")
    output_dir.mkdir(parents=True)
    report_time = int(timestamp.timestamp() * 1000)

    with zipfile.ZipFile(input_package) as outer:
        outer_names = [info.filename for info in outer.infolist() if not info.is_dir()]
        if len(outer_names) != 2:
            raise ValueError("输入外层 ZIP 必须只包含两个内层 ZIP")
        outer_extract = output_dir / "_outer"
        safe_extract(input_package, outer_extract)

    accuracy_zip, image_zip = find_inner_zips(outer_extract)
    accuracy_root = output_dir / "_accuracy"
    image_root = output_dir / "_image"
    safe_extract(accuracy_zip, accuracy_root)
    safe_extract(image_zip, image_root)

    accuracy_json = find_single_file(accuracy_root, "accuracy.json")
    find_single_file(image_root, "image.json")
    exported_accuracy = prepare_accuracy_for_export(
        accuracy,
        report_time=report_time,
        save_problem_images=save_problem_images,
    )
    write_json(accuracy_json, exported_accuracy)

    image_tree = image_root / "image"
    validate_reviewed_package(exported_accuracy, image_tree)

    rebuilt_accuracy = output_dir / accuracy_zip.name
    rebuilt_image = output_dir / image_zip.name
    zip_dir_contents(accuracy_root, rebuilt_accuracy)
    zip_dir_contents(image_root, rebuilt_image)

    output_package = nonconflicting_path(output_dir / f"{output_stem}_m_{stamp}.zip")
    with zipfile.ZipFile(output_package, "w", compression=zipfile.ZIP_DEFLATED) as outer:
        outer.write(rebuilt_accuracy, rebuilt_accuracy.name)
        outer.write(rebuilt_image, rebuilt_image.name)
    with zipfile.ZipFile(output_package) as archive:
        names = sorted(info.filename for info in archive.infolist() if not info.is_dir())
    if sorted([rebuilt_accuracy.name, rebuilt_image.name]) != names:
        raise ValueError(f"新外层 ZIP 条目不符合预期: {names}")

    shutil.rmtree(outer_extract, ignore_errors=True)
    shutil.rmtree(accuracy_root, ignore_errors=True)
    shutil.rmtree(image_root, ignore_errors=True)
    logging.info("m 模式产出完成: %s", output_package)
    return output_package


def validate_reviewed_package(accuracy: dict[str, Any], image_root: Path) -> None:
    categories = accuracy.get("data", {}).get("statistic")
    if not isinstance(categories, list):
        raise ValueError("导出校验失败: data.statistic 不是数组")

    for stat in categories:
        algo_type = str(stat.get("algoType") or "")
        images = stat.get("imagesData") or []
        if not isinstance(images, list):
            raise ValueError(f"导出校验失败: imagesData 不是数组: {algo_type}")
        image_dir = image_root / algo_type
        disk_names = {path.name for path in image_dir.iterdir() if path.is_file()} if image_dir.is_dir() else set()
        json_names = [str(item.get("imageName") or "") for item in images]
        if len(json_names) != len(set(json_names)):
            raise ValueError(f"导出校验失败: imageName 重复: {algo_type}")
        missing = sorted(set(json_names) - disk_names)
        if missing:
            raise ValueError(f"导出校验失败: 图片缺失: {algo_type} -> {missing[:5]}")


@dataclass(frozen=True)
class ReviewSessionOptions:
    config_path: Path
    database_path: Path | None
    save_problem_images: bool
    recognition_red_keywords: list[str]
    point_name_red_keywords: list[str]
    point_name_end_open: bool
    point_name_end_characters: int
    point_name_end_character_color: str
    default_zoom: float
    auto_click_times_ms: list[int]
    show_auto_click: bool
    show_mark_all_correct: bool
    show_mark_same_recognition: bool
    exit_after_export: bool
    web_hotkeys: dict[str, str]
    newkeywords: list[str]
    default_wrong_keywords: list[str]
    default_problem_keywords: list[str]
    extra_xlsx_recognition_enabled: bool
    extra_xlsx_recognition_keywords: list[str]
    show_algorithm_normal: bool
    checklist_image_width_enabled: bool
    algo_type_labels: dict[str, str]
    solution_rules: list[tuple[str, str, str]]
    solution_options: list[str]
    responsible_rules: list[tuple[str, str]]

    def session_kwargs(self) -> dict[str, Any]:
        return {
            "config_path": self.config_path,
            "database_path": self.database_path,
            "save_problem_images": self.save_problem_images,
            "recognition_red_keywords": self.recognition_red_keywords,
            "point_name_red_keywords": self.point_name_red_keywords,
            "point_name_end_open": self.point_name_end_open,
            "point_name_end_characters": self.point_name_end_characters,
            "point_name_end_character_color": self.point_name_end_character_color,
            "default_zoom": self.default_zoom,
            "auto_click_times_ms": self.auto_click_times_ms,
            "show_auto_click": self.show_auto_click,
            "show_mark_all_correct": self.show_mark_all_correct,
            "show_mark_same_recognition": self.show_mark_same_recognition,
            "exit_after_export": self.exit_after_export,
            "web_hotkeys": self.web_hotkeys,
            "newkeywords": self.newkeywords,
            "default_wrong_keywords": self.default_wrong_keywords,
            "default_problem_keywords": self.default_problem_keywords,
            "extra_xlsx_recognition_enabled": self.extra_xlsx_recognition_enabled,
            "extra_xlsx_recognition_keywords": self.extra_xlsx_recognition_keywords,
            "show_algorithm_normal": self.show_algorithm_normal,
            "checklist_image_width_enabled": self.checklist_image_width_enabled,
            "algo_type_labels": self.algo_type_labels,
            "solution_rules": self.solution_rules,
            "solution_options": self.solution_options,
            "responsible_rules": self.responsible_rules,
        }


@dataclass(frozen=True)
class ExportedFiles:
    package: Path
    checklist: Path


class ReviewSession:
    def __init__(
        self,
        input_package: Path,
        output_root: Path,
        temp_root: Path,
        config_path: Path,
        save_problem_images: bool,
        recognition_red_keywords: list[str],
        point_name_red_keywords: list[str],
        point_name_end_open: bool,
        point_name_end_characters: int,
        point_name_end_character_color: str,
        default_zoom: float,
        auto_click_times_ms: list[int],
        show_auto_click: bool,
        show_mark_all_correct: bool,
        show_mark_same_recognition: bool,
        exit_after_export: bool,
        web_hotkeys: dict[str, str],
        newkeywords: list[str],
        default_wrong_keywords: list[str],
        default_problem_keywords: list[str],
        extra_xlsx_recognition_enabled: bool,
        extra_xlsx_recognition_keywords: list[str],
        show_algorithm_normal: bool,
        checklist_image_width_enabled: bool,
        algo_type_labels: dict[str, str],
        solution_rules: list[tuple[str, str, str]],
        solution_options: list[str],
        responsible_rules: list[tuple[str, str]],
        database_path: Path | None = None,
    ) -> None:
        self.input_package = input_package
        self.output_root = output_root
        self.config_path = config_path
        self.save_problem_images = save_problem_images
        self.recognition_red_keywords = recognition_red_keywords
        self.point_name_red_keywords = point_name_red_keywords
        self.point_name_end_open = point_name_end_open
        self.point_name_end_characters = point_name_end_characters
        self.point_name_end_character_color = point_name_end_character_color
        self.default_zoom = default_zoom
        self.auto_click_times_ms = auto_click_times_ms
        self.show_auto_click = show_auto_click
        self.show_mark_all_correct = show_mark_all_correct
        self.show_mark_same_recognition = show_mark_same_recognition
        self.exit_after_export = exit_after_export
        self.web_hotkeys = web_hotkeys
        self.newkeywords = newkeywords
        self.default_wrong_keywords = default_wrong_keywords
        self.default_problem_keywords = default_problem_keywords
        self.extra_xlsx_recognition_enabled = extra_xlsx_recognition_enabled
        self.extra_xlsx_recognition_keywords = extra_xlsx_recognition_keywords
        self.show_algorithm_normal = show_algorithm_normal
        self.checklist_image_width_enabled = checklist_image_width_enabled
        self.algo_type_labels = algo_type_labels
        self.solution_rules = solution_rules
        self.solution_options = solution_options
        self.responsible_rules = responsible_rules
        self.annotation_db = AnnotationDatabase(database_path) if database_path else None
        self.temp_root = temp_root
        self.state_dir = self.output_root / SESSION_STATE_DIRNAME
        self.backup_dir = nonconflicting_path(self.output_root / f"{datetime.now():%m%d}-images-{attachment_output_stem(self.input_package)}")
        self.outer_root = temp_root / "outer"
        self.accuracy_root = temp_root / "accuracy"
        self.image_root = temp_root / "image_result"
        safe_extract(input_package, self.outer_root)
        self.accuracy_zip, self.image_zip = find_inner_zips(self.outer_root)
        safe_extract(self.accuracy_zip, self.accuracy_root)
        safe_extract(self.image_zip, self.image_root)
        self.accuracy_path = find_single_file(self.accuracy_root, "accuracy.json")
        self.image_json_path = find_single_file(self.image_root, "image.json")
        self.image_tree = self.image_root / "image"
        self.accuracy = load_json(self.accuracy_path)
        data = self.accuracy.get("data", {})
        fallback_audit_time = _format_archive_time(data.get("taskRunTime")) or "-"
        self.archive_audit_time = load_archive_audit_time(
            [self.accuracy_root, self.image_root],
            fallback_audit_time,
        )
        self.category_index = 0
        self.image_index = 0
        self.backed_up_images: set[str] = set()
        self.backup_names: set[str] = set()
        self.backup_records: dict[str, dict[str, str]] = {}
        self.session_key = f"{session_package_key(self.input_package)}_{os.getpid()}_{int(datetime.now().timestamp() * 1000)}"
        self.state_path = self._build_state_path()
        self.state_dir.mkdir(parents=True, exist_ok=True)
        self.backup_dir.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.exported: Path | None = None
        self.exported_checklist: Path | None = None
        restored_state = self._restore_state()
        self._apply_review_sorting(preserve_position=restored_state)
        if self.annotation_db:
            applied = self.annotation_db.apply_to_accuracy(self.accuracy)
            if applied:
                self._clamp_position()
                logging.info("m/h 模式复用历史标注: %s 条", applied)
        validate_reviewed_package(self.accuracy, self.image_tree)

    def _apply_review_sorting(self, preserve_position: bool = False) -> None:
        anchor: tuple[str, str] | None = None
        if preserve_position:
            try:
                statistics = self.accuracy["data"]["statistic"]
                stat = statistics[self.category_index]
                item = stat["imagesData"][self.image_index]
                anchor = (
                    str(stat.get("algoType") or ""),
                    str(item.get("imageName") or ""),
                )
            except (AttributeError, IndexError, KeyError, TypeError):
                anchor = None

        from rebuild_sources import load_bool_config, load_keywords, sort_accuracy_by_config

        sort_accuracy_by_config(
            self.accuracy,
            keywords=load_keywords(self.config_path),
            empty_recognition_last=load_bool_config(
                self.config_path,
                "empty_recognition_last",
                True,
            ),
            cluster_by_recognition=load_bool_config(
                self.config_path,
                "cluster_by_recognition",
                False,
            ),
        )

        if anchor:
            statistics = self.accuracy.get("data", {}).get("statistic") or []
            for category_index, stat in enumerate(statistics):
                if not isinstance(stat, dict) or str(stat.get("algoType") or "") != anchor[0]:
                    continue
                for image_index, item in enumerate(stat.get("imagesData") or []):
                    if isinstance(item, dict) and str(item.get("imageName") or "") == anchor[1]:
                        self.category_index = category_index
                        self.image_index = image_index
                        self._clamp_position()
                        return
        self._clamp_position()

    def _build_state_path(self) -> Path:
        return self.state_dir / f"{self.session_key}.json"

    def _current_categories(self) -> list[dict[str, Any]]:
        return build_state(self.accuracy)

    def _clamp_position(self, categories: list[dict[str, Any]] | None = None) -> None:
        categories = categories if categories is not None else self._current_categories()
        if not categories:
            self.category_index = 0
            self.image_index = 0
            return
        self.category_index = max(0, min(self.category_index, len(categories) - 1))
        category = categories[self.category_index]
        images = category["images"]
        self.image_index = max(0, min(self.image_index, len(images) - 1))

    def _position_payload(self) -> dict[str, int]:
        return {"categoryIndex": self.category_index, "imageIndex": self.image_index}

    def _persist_state(self) -> None:
        payload = {
            "inputPackage": str(self.input_package.resolve()),
            "categoryIndex": self.category_index,
            "imageIndex": self.image_index,
            "backedUpImages": sorted(self.backed_up_images),
            "backupRecords": self.backup_records,
            "accuracy": self.accuracy,
        }
        write_json(self.state_path, payload)

    def _restore_state(self) -> bool:
        if not self.state_path.is_file():
            self.backup_names = {
                path.name.lower()
                for path in self.backup_dir.glob("*")
                if path.is_file()
            }
            return False
        try:
            payload = load_json(self.state_path)
            if str(Path(payload.get("inputPackage") or "").resolve()) != str(self.input_package.resolve()):
                raise ValueError("input package mismatch")
            accuracy = payload.get("accuracy")
            if not isinstance(accuracy, dict):
                raise ValueError("accuracy state missing")
            self.accuracy = accuracy
            self.category_index = int(payload.get("categoryIndex", 0))
            self.image_index = int(payload.get("imageIndex", 0))
            self.backed_up_images = {str(item) for item in payload.get("backedUpImages") or []}
            records = payload.get("backupRecords")
            self.backup_records = records if isinstance(records, dict) else {}
            self.backup_names = {
                path.name.lower()
                for path in self.backup_dir.glob("*")
                if path.is_file()
            }
            logging.info("m 模式恢复快照: %s", self.state_path)
            return True
        except Exception as exc:
            logging.warning("m 模式快照恢复失败，改用原始包: %s", exc)
            self.backed_up_images = set()
            self.backup_records = {}
            self.backup_names = {
                path.name.lower()
                for path in self.backup_dir.glob("*")
                if path.is_file()
            }
            return False

    def _backup_problem_image(self, algo_type: str, item: dict[str, Any]) -> str | None:
        image_name = str(item.get("imageName") or "").strip()
        if not image_name:
            return None
        key = f"{algo_type}/{image_name}"
        record = self.backup_records.get(key)
        if record:
            filename = str(record.get("filename") or "")
            if filename and (self.backup_dir / filename).is_file():
                return filename
        source_image = self.image_tree / algo_type / image_name
        if not source_image.is_file():
            raise FileNotFoundError(f"问题图片源文件不存在: {source_image}")
        backup_name = selected_filename_for_item(item, image_name)
        final_name = unique_filename(Path(backup_name).stem, Path(backup_name).suffix, self.backup_names)
        while (self.backup_dir / final_name).exists():
            final_name = unique_filename(Path(final_name).stem, Path(final_name).suffix, self.backup_names)
        shutil.copy2(source_image, self.backup_dir / final_name)
        self.backed_up_images.add(key)
        logging.info("m 模式备份问题图片: %s -> %s", source_image, self.backup_dir / final_name)
        return final_name

    def _ensure_problem_image_record(self, algo_type: str, item: dict[str, Any], audit_status: int) -> dict[str, str] | None:
        image_name = str(item.get("imageName") or "").strip()
        if not image_name:
            return None
        filename = self._backup_problem_image(algo_type, item)
        if not filename:
            return None
        key = f"{algo_type}/{image_name}"
        record = {
            "filename": filename,
            "pointName": str(item.get("pointName") or item.get("pointCode") or Path(filename).stem),
            "status": STATUS_NAMES.get(audit_status, str(audit_status)),
            "recognition": str(item.get("recognition") or ""),
        }
        self.backup_records[key] = record
        return record

    def _write_problem_images_workbook(self) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.drawing.image import Image as ExcelImage
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
            from openpyxl.worksheet.datavalidation import DataValidation
            from openpyxl.workbook.defined_name import DefinedName
            from PIL import Image as PILImage
        except ImportError as exc:
            raise RuntimeError("生成问题图片 xlsx 需要 openpyxl 和 Pillow") from exc

        workbook_path = nonconflicting_path(self.output_root / f"{dated_attachment_output_stem(self.input_package)}-清单.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "核算报告"
        thin_black = Side(style="thin", color="000000")
        cell_border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
        image_column_border = Border(left=thin_black, right=thin_black, top=thin_black)
        title_fill = PatternFill("solid", fgColor="1F4E78")
        white_fill = PatternFill("solid", fgColor="FFFFFF")
        section_fill = PatternFill("solid", fgColor="F2F2F2")
        centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
        wrapped = Alignment(horizontal="left", vertical="center", wrap_text=True)
        responsible_validation = DataValidation(
            type="list",
            formula1=f'"{",".join(RESPONSIBLE_VALUES)}"',
            allow_blank=True,
        )
        sheet.add_data_validation(responsible_validation)
        solution_validation = None
        if self.solution_options:
            option_sheet = workbook.create_sheet(SOLUTION_OPTIONAL_SHEET_NAME)
            for option_row, option in enumerate(self.solution_options, 1):
                option_sheet.cell(option_row, 1, option)
            option_sheet.sheet_state = "hidden"
            option_range = f"'{option_sheet.title}'!$A$1:$A${len(self.solution_options)}"
            workbook.defined_names.add(
                DefinedName("solution_options", attr_text=option_range)
            )
            solution_validation = DataValidation(
                type="list",
                formula1="=solution_options",
                allow_blank=True,
            )
            sheet.add_data_validation(solution_validation)

        sheet.column_dimensions["A"].width = 8
        sheet.column_dimensions["B"].width = 40
        sheet.column_dimensions["C"].width = 24
        sheet.column_dimensions["D"].width = 16
        sheet.column_dimensions["E"].width = 16
        sheet.column_dimensions["F"].width = 30
        sheet.column_dimensions["G"].width = 28
        sheet.column_dimensions["H"].width = 43
        image_display_width = 320
        image_display_height = 180
        sheet.column_dimensions["I"].width = 43
        sheet.freeze_panes = None
        sheet.sheet_view.showGridLines = False
        header_row_height = 28.20

        statistics = self.accuracy.get("data", {}).get("statistic") or []
        data = self.accuracy.get("data", {})
        all_items = [
            item
            for stat in statistics
            for item in (stat.get("imagesData") or [])
            if isinstance(item, dict)
        ]
        total_count = len(all_items)
        online_count = sum(
            1 for item in all_items if "不在线" in str(item.get("recognition") or "")
        )
        failure_count = sum(
            1
            for item in all_items
            if any(keyword in str(item.get("recognition") or "") for keyword in ("失败", "为空"))
        )
        abnormal_point_count = sum(
            1
            for item in all_items
            if status_value(item.get("auditStatus")) not in (0, 1)
        )

        def percentage_text(count: int) -> str:
            return f"{count * 100 / total_count:.2f}%" if total_count else "0.00%"

        def format_report_date(value: Any) -> str:
            text = str(value or "").strip()
            if re.fullmatch(r"\d{8}", text):
                return f"{text[:4]}-{text[4:6]}-{text[6:]}"
            return text or "-"

        def format_report_time(value: Any) -> str:
            text = str(value or "").strip()
            if re.fullmatch(r"\d{14}", text):
                return (
                    f"{text[:4]}-{text[4:6]}-{text[6:8]} "
                    f"{text[8:10]}:{text[10:12]}:{text[12:]}"
                )
            return text or "-"

        station_name = str(self.accuracy.get("stationName") or "变电站").strip()
        task_name = str(data.get("taskName") or attachment_output_stem(self.input_package)).strip()
        algorithm_names = checklist_algorithm_names(self.accuracy, self.algo_type_labels)
        algorithm_summary = "、".join(algorithm_names) if algorithm_names else "无"
        task_summary = (
            f"{task_name}；共 {len(algorithm_names)} 个算法，分别为（{algorithm_summary}）"
        )
        report_date = format_report_date(data.get("statisticDate"))
        task_run_time = format_report_time(data.get("taskRunTime"))
        report_time = self.archive_audit_time

        sheet.merge_cells("A1:I1")
        sheet["A1"] = f"{station_name}核算报告"
        sheet["A1"].font = Font(bold=True, color="FFFFFF", size=16)
        sheet["A1"].fill = title_fill
        sheet["A1"].alignment = centered
        sheet.row_dimensions[1].height = header_row_height

        def write_info_row(row: int, left_label: str, left_value: str, right_label: str = "", right_value: str = "") -> None:
            sheet.cell(row, 1, left_label)
            sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            sheet.cell(row, 2, left_value)
            sheet.cell(row, 4, right_label)
            sheet.merge_cells(start_row=row, start_column=5, end_row=row, end_column=9)
            sheet.cell(row, 5, right_value)
            for column in range(1, 10):
                cell = sheet.cell(row, column)
                cell.border = cell_border
                cell.alignment = wrapped if column in (2, 5) else centered
                if column in (1, 4):
                    cell.font = Font(bold=True, color="000000")
                if column == 1:
                    cell.fill = white_fill

        write_info_row(2, "变电站", station_name, "异常点位", str(abnormal_point_count))
        write_info_row(3, "巡视日期", report_date, "巡视任务", task_summary)
        write_info_row(
            4,
            "巡视设备在线率",
            f"包含“不在线”点位 {online_count} 个，占全部点位 {percentage_text(online_count)}",
            "巡视结果失败率",
            f"包含“失败/为空”点位 {failure_count} 个，占全部点位 {percentage_text(failure_count)}",
        )
        write_info_row(5, "审核人", "-", "审核时间", report_time)
        write_info_row(6, "巡视开始时间", task_run_time, "巡视结束时间", report_time)

        inaccurate_rows: list[dict[str, Any]] = []
        problem_rows: list[dict[str, Any]] = []
        normal_rows: list[dict[str, Any]] = []
        for stat in statistics:
            algo_type = str(stat.get("algoType") or "").strip()
            for item in stat.get("imagesData") or []:
                if not isinstance(item, dict):
                    continue
                audit_status = status_value(item.get("auditStatus"))
                recognition = str(item.get("recognition") or "")
                extra_match = self.extra_xlsx_recognition_enabled and any(
                    not recognition.strip() if keyword == "" else keyword in recognition
                    for keyword in self.extra_xlsx_recognition_keywords
                )
                is_failure = any(keyword in recognition for keyword in ("失败", "为空"))
                is_normal = self.show_algorithm_normal and audit_status == 1
                if audit_status == 1 and not self.show_algorithm_normal:
                    continue
                if audit_status != 3 and audit_status != 2 and not is_failure and not extra_match and not is_normal:
                    continue
                record: dict[str, Any] = {
                    "algoType": self.algo_type_labels.get(algo_type, algo_type),
                    "pointName": str(item.get("pointName") or item.get("pointCode") or item.get("imageName") or ""),
                    "recognition": recognition.strip() or "无分析结果",
                    "auditStatus": audit_status,
                    "item": item,
                    "algoKey": algo_type,
                }
                if audit_status in (2, 3):
                    record = self._ensure_problem_image_record(algo_type, item, audit_status)
                    if not record:
                        continue
                    record = {
                        "algoType": self.algo_type_labels.get(algo_type, algo_type),
                        "pointName": record["pointName"],
                        "recognition": recognition.strip() or "无分析结果",
                        "auditStatus": audit_status,
                        "imagePath": self.backup_dir / record["filename"],
                        "imageLabel": record["filename"],
                    }
                else:
                    image_name = str(item.get("imageName") or "").strip()
                    if not image_name:
                        continue
                    image_path = self.image_tree / algo_type / image_name
                    if not image_path.is_file():
                        logging.warning("清单跳过缺失图片: %s", image_path)
                        continue
                    record["imagePath"] = image_path
                    record["imageLabel"] = image_name

                record["situation"] = {
                    1: "算法正常",
                    2: "算法不准",
                    3: "问题图片",
                }.get(audit_status, "未标注")
                record["solution"] = resolve_solution(
                    self.solution_rules,
                    record["situation"],
                    record["recognition"],
                )
                record["responsible"] = resolve_responsible(
                    self.responsible_rules,
                    record["recognition"],
                )
                record["auditTime"] = self.archive_audit_time
                if audit_status == 3:
                    problem_rows.append(record)
                elif audit_status == 1:
                    normal_rows.append(record)
                else:
                    inaccurate_rows.append(record)

        def write_section(row: int, title: str, rows: list[dict[str, Any]]) -> int:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=9)
            title_cell = sheet.cell(row, 1, title)
            title_cell.font = Font(bold=True, color="FF0000", size=16)
            title_cell.fill = section_fill
            title_cell.alignment = centered
            sheet.row_dimensions[row + 1].height = 21
            row += 2
            headers = ["序号", "点位名称", "算法类型", "审算时间", "识别结果", "判断情况", "负责人", "解决方案", "图片"]
            for column, header in enumerate(headers, 1):
                cell = sheet.cell(row, column, header)
                cell.font = Font(bold=True, color="000000")
                cell.fill = white_fill
                cell.alignment = centered
                cell.border = image_column_border if column == 9 else cell_border
            row += 1
            first_data_row = row
            for sequence, record in enumerate(rows, 1):
                image_path = Path(record["imagePath"])
                current_algo_type = str(record["algoType"])
                sheet.cell(row, 1, sequence)
                sheet.cell(row, 2, record["pointName"])
                sheet.cell(row, 3, current_algo_type)
                sheet.cell(row, 4, record["auditTime"])
                sheet.cell(row, 5, record["recognition"])
                sheet.cell(row, 6, record["situation"])
                sheet.cell(row, 7, record["responsible"])
                sheet.cell(row, 8, record["solution"])
                sheet.cell(row, 9, record["imageLabel"])
                sheet.cell(row, 9).hyperlink = image_path.resolve().as_uri()
                sheet.cell(row, 9).style = "Hyperlink"
                sheet.row_dimensions[row].height = image_display_height * 0.75
                for column in range(1, 10):
                    cell = sheet.cell(row, column)
                    cell.alignment = centered if column in (1, 3, 4, 5, 6, 7, 8, 9) else wrapped
                    cell.border = image_column_border if column == 9 else cell_border
                    cell.fill = white_fill
                    if column == 1:
                        cell.font = Font(color="000000")
                try:
                    with PILImage.open(image_path) as source:
                        width, height = source.size
                    image = ExcelImage(image_path)
                    scale = min(image_display_width / width, image_display_height / height)
                    image.width = int(width * scale)
                    image.height = int(height * scale)
                    if (
                        self.checklist_image_width_enabled
                        and image.width < CHECKLIST_IMAGE_MIN_WIDTH_PX
                    ):
                        scale = CHECKLIST_IMAGE_TARGET_WIDTH_PX / image.width
                        image.width = round(image.width * scale)
                        image.height = round(image.height * scale)
                        sheet.row_dimensions[row].height = max(
                            sheet.row_dimensions[row].height or 0,
                            image.height * 0.75,
                        )
                    sheet.add_image(image, f"I{row}")
                except (OSError, ValueError) as exc:
                    logging.warning("清单嵌入图片失败，保留图片链接: %s (%s)", image_path, exc)
                row += 1
            if row > first_data_row:
                responsible_validation.add(f"G{first_data_row}:G{row - 1}")
                if solution_validation is not None:
                    solution_validation.add(f"H{first_data_row}:H{row - 1}")
            return row

        problem_count = len(problem_rows)
        inaccurate_count = len(inaccurate_rows)
        summary = (
            f"总点位{total_count}个，问题图片{problem_count}个，"
            f"算法不准{inaccurate_count}个。"
        )
        conclusion = (
            summarize_algorithm_counts(inaccurate_rows, "一、算法不准")
            + ";"
            + summarize_algorithm_counts(problem_rows, "二、问题图片")
            + ";三、解决方案涉及\"采集\"字样的点位如有需要请及时采集图片，用以优化更新模型、"
            "涉及现场调试还请及时处理或反馈无法处理缘由，一同推进算法优化进程。"
        )
        sheet.merge_cells("A7:A7")
        sheet.cell(7, 1, "核算报告")
        sheet.merge_cells("B7:I7")
        sheet.cell(7, 2, summary)
        sheet.merge_cells("A8:A8")
        sheet.cell(8, 1, "核算结论")
        sheet.merge_cells("B8:I8")
        sheet.cell(8, 2, conclusion)
        for row in (7, 8):
            for column in range(1, 10):
                cell = sheet.cell(row, column)
                cell.border = cell_border
                cell.alignment = wrapped if column == 2 else centered
                cell.fill = white_fill
                cell.font = Font(bold=(column == 1), color="000000")
        for row in range(1, 9):
            sheet.row_dimensions[row].height = header_row_height

        row_index = 9
        if inaccurate_rows:
            row_index = write_section(row_index, "算法不准", inaccurate_rows)
        if problem_rows:
            row_index = write_section(row_index, "问题图片", problem_rows)
        if normal_rows:
            write_section(row_index, "算法正常", normal_rows)

        workbook.save(workbook_path)
        logging.info(
            "m 模式生成核算报告清单: %s（算法不准=%s，问题图片=%s，算法正常=%s）",
            workbook_path,
            inaccurate_count,
            problem_count,
            len(normal_rows),
        )
        return workbook_path

    def _compare_result2_algorithm_names(self) -> tuple[list[str], list[str]] | None:
        result2_path = load_result2_xlsx_path(self.config_path)
        if result2_path is None:
            return None
        result2_algorithms = result2_algorithm_names(
            result2_path,
            str(
                self.accuracy.get("stationName")
                or self.accuracy.get("data", {}).get("stationName")
                or ""
            ),
        )
        if result2_algorithms is None:
            return None
        aliases = load_result2_algorithm_aliases(self.config_path)
        return compare_checklist_and_result2_algorithms(
            checklist_algorithm_records(self.accuracy, self.algo_type_labels),
            result2_algorithms,
            aliases,
        )

    def _write_newkeywords_report(
        self,
        report_path: Path | None = None,
        ensure_exists: bool = False,
    ) -> Path | None:
        report_path = report_path or nonconflicting_path(
            self.output_root / f"{dated_attachment_output_stem(self.input_package)}-清单.txt"
        )
        statistics = self.accuracy.get("data", {}).get("statistic") or []
        lines: list[str] = []
        checklist_points = checklist_table_point_names(report_path.with_suffix(".xlsx"))
        if checklist_points:
            lines.append(f"类型 => {CHECKLIST_TABLE_REPORT_TITLE}")
            lines.extend(checklist_points)
            lines.append("")

        for keyword in self.newkeywords:
            title = '""' if keyword == "" else keyword
            seen: set[str] = set()
            matches: list[str] = []
            for stat in statistics:
                images = stat.get("imagesData") or []
                if not isinstance(images, list):
                    continue
                for item in images:
                    if not isinstance(item, dict):
                        continue
                    recognition = str(item.get("recognition") or "")
                    matched = not recognition.strip() if keyword == "" else keyword in recognition
                    if not matched:
                        continue
                    point_name = str(item.get("pointName") or item.get("pointCode") or item.get("imageName") or "").strip()
                    if point_name and point_name not in seen:
                        matches.append(point_name)
                        seen.add(point_name)
            if matches:
                lines.append(f"类型 => {title}")
                lines.extend(matches)
                lines.append("")

        algorithm_differences = self._compare_result2_algorithm_names()
        if algorithm_differences is not None:
            result2_only, checklist_only = algorithm_differences
            if lines and lines[-1]:
                lines.append("")
            lines.append("本次核算中autoedge未被覆盖的算法有{")
            lines.extend(result2_only)
            lines.append("}")
            lines.append("")
            lines.append("本次核算中autoedge应该新增的算法有{")
            lines.extend(checklist_only)
            lines.append("}")

        if not lines:
            if ensure_exists:
                report_path.write_text("", encoding="utf-8")
                logging.info("生成空清单文本: %s", report_path)
                return report_path
            logging.info("清单文本无命中内容，不生成报告")
            return None

        report_path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
        logging.info("生成清单文本报告: %s", report_path)
        return report_path

    def _next_position(self, category_index: int, image_index: int, categories: list[dict[str, Any]] | None = None) -> tuple[int, int]:
        categories = categories if categories is not None else self._current_categories()
        if not categories:
            return 0, 0
        next_category = category_index
        next_image = image_index + 1
        if next_image >= len(categories[category_index]["images"]):
            next_category += 1
            while next_category < len(categories) and not categories[next_category]["images"]:
                next_category += 1
            if next_category >= len(categories):
                next_category = len(categories) - 1
                next_image = len(categories[next_category]["images"]) - 1
            else:
                next_image = 0
        return next_category, next_image

    def _set_position(self, category_index: int, image_index: int) -> None:
        self.category_index = category_index
        self.image_index = image_index
        self._clamp_position()

    def state(self) -> dict[str, Any]:
        with self.lock:
            categories = self._current_categories()
            self._clamp_position(categories)
            return {
                "package": self.input_package.name,
                "categories": categories,
                "allMarked": all_marked(self.accuracy),
                "currentCategoryIndex": self.category_index,
                "currentImageIndex": self.image_index,
                "recognitionRedKeywords": self.recognition_red_keywords,
                "pointNameRedKeywords": self.point_name_red_keywords,
                "pointNameEndOpen": self.point_name_end_open,
                "pointNameEndCharacters": self.point_name_end_characters,
                "pointNameEndCharacterColor": self.point_name_end_character_color,
                "defaultWrongKeywords": self.default_wrong_keywords,
                "defaultProblemKeywords": self.default_problem_keywords,
                "defaultZoom": self.default_zoom,
                "autoClickTimesMs": self.auto_click_times_ms,
                "showAutoClick": self.show_auto_click,
                "showMarkAllCorrect": self.show_mark_all_correct,
                "showMarkSameRecognition": self.show_mark_same_recognition,
                "exitAfterExport": self.exit_after_export,
                "webHotkeys": self.web_hotkeys,
            }

    def navigate(self, category_index: int, image_index: int) -> dict[str, Any]:
        with self.lock:
            categories = self._current_categories()
            if not categories:
                raise ValueError("accuracy.json 中没有可浏览的类别")
            self._set_position(category_index, image_index)
            self._persist_state()
            return self.state()

    def mark(self, algo_type: str, index: int, audit_status: int, advance: bool = True) -> dict[str, Any]:
        with self.lock:
            categories = self.accuracy.get("data", {}).get("statistic")
            if not isinstance(categories, list):
                raise ValueError("accuracy.json 缺少 data.statistic 数组")
            if audit_status not in (0, 1, 2, 3):
                raise ValueError(f"auditStatus 无效: {audit_status}")

            matched_index: int | None = None
            for stat_index, stat in enumerate(categories):
                if not isinstance(stat, dict) or str(stat.get("algoType") or "") != algo_type:
                    continue
                images = stat.get("imagesData")
                if not isinstance(images, list) or index < 0 or index >= len(images):
                    raise IndexError(f"点位索引无效: {algo_type}/{index}")
                images[index]["auditStatus"] = audit_status
                calculate_statistics(stat)
                matched_index = stat_index
                if audit_status in (2, 3):
                    self._ensure_problem_image_record(algo_type, images[index], audit_status)
                else:
                    image_name = str(images[index].get("imageName") or "").strip()
                    if image_name:
                        self.backup_records.pop(f"{algo_type}/{image_name}", None)
                break
            if matched_index is None:
                raise KeyError(f"不存在的 algoType: {algo_type}")

            if advance:
                categories_state = self._current_categories()
                self._set_position(*self._next_position(matched_index, index, categories_state))
            else:
                self._set_position(matched_index, index)
            self._persist_state()
            return self.state()

    def mark_category_correct(self, algo_type: str) -> dict[str, Any]:
        with self.lock:
            categories = self.accuracy.get("data", {}).get("statistic")
            if not isinstance(categories, list):
                raise ValueError("accuracy.json 缺少 data.statistic 数组")

            matched_index: int | None = None
            for stat_index, stat in enumerate(categories):
                if not isinstance(stat, dict) or str(stat.get("algoType") or "") != algo_type:
                    continue
                images = stat.get("imagesData")
                if not isinstance(images, list):
                    raise ValueError(f"imagesData 不是数组: {algo_type}")
                for item in images:
                    if isinstance(item, dict):
                        item["auditStatus"] = 1
                calculate_statistics(stat)
                matched_index = stat_index
                break
            if matched_index is None:
                raise KeyError(f"不存在的 algoType: {algo_type}")

            categories_state = self._current_categories()
            next_category = matched_index + 1
            if next_category < len(categories_state):
                self._set_position(next_category, 0)
            else:
                self._set_position(matched_index, 0)
            self._persist_state()
            return self.state()

    def mark_same_recognition(self, algo_type: str, recognition: str, audit_status: int) -> dict[str, Any]:
        with self.lock:
            if audit_status not in (1, 2, 3):
                raise ValueError(f"auditStatus 无效: {audit_status}")
            categories = self.accuracy.get("data", {}).get("statistic")
            if not isinstance(categories, list):
                raise ValueError("accuracy.json 缺少 data.statistic 数组")

            matched_index: int | None = None
            marked = 0
            for stat_index, stat in enumerate(categories):
                if not isinstance(stat, dict) or str(stat.get("algoType") or "") != algo_type:
                    continue
                images = stat.get("imagesData")
                if not isinstance(images, list):
                    raise ValueError(f"imagesData 不是数组: {algo_type}")
                for item in images:
                    if not isinstance(item, dict):
                        continue
                    if status_value(item.get("auditStatus")) != 0:
                        continue
                    if str(item.get("recognition") or "") != recognition:
                        continue
                    item["auditStatus"] = audit_status
                    marked += 1
                    if audit_status in (2, 3):
                        self._ensure_problem_image_record(algo_type, item, audit_status)
                    else:
                        image_name = str(item.get("imageName") or "").strip()
                        if image_name:
                            self.backup_records.pop(f"{algo_type}/{image_name}", None)
                calculate_statistics(stat)
                matched_index = stat_index
                break
            if matched_index is None:
                raise KeyError(f"不存在的 algoType: {algo_type}")
            if marked == 0:
                raise ValueError("当前识别结果类别没有未确认点位")

            self._set_position(matched_index, self.image_index)
            self._persist_state()
            return self.state()

    def export(self) -> ExportedFiles:
        with self.lock:
            if not all_marked(self.accuracy):
                raise ValueError("仍有未标注点位，不能产出压缩包")
            if self.annotation_db:
                saved = self.annotation_db.upsert_marked_accuracy(self.accuracy)
                logging.info("m/h 模式保存标注到 SQLite: %s 条", saved)
            self.exported = export_reviewed_package(
                self.input_package,
                self.output_root,
                self.accuracy,
                self.save_problem_images,
            )
            workbook_path = self._write_problem_images_workbook()
            self.exported_checklist = workbook_path
            self._write_newkeywords_report(report_path=workbook_path.with_suffix(".txt"))
            try:
                self.state_path.unlink(missing_ok=True)
            except Exception:
                logging.warning("m 模式导出后清理快照失败: %s", self.state_path)
            return ExportedFiles(package=self.exported, checklist=workbook_path)


class ReviewWorkspace:
    def __init__(
        self,
        rebuild_output_root: Path,
        review_output_root: Path,
        temp_root: Path,
        options: ReviewSessionOptions,
    ) -> None:
        self.rebuild_output_root = rebuild_output_root
        self.review_output_root = review_output_root
        self.temp_root = temp_root
        self.options = options
        self.session: ReviewSession | None = None
        self.exported: ExportedFiles | None = None
        self.source_package: Path | None = None
        self.prepared_package: Path | None = None
        self.lock = threading.RLock()
        self.upload_root = self.temp_root / "uploads"
        self.upload_root.mkdir(parents=True, exist_ok=True)

    def _session_kwargs(self) -> dict[str, Any]:
        return self.options.session_kwargs()

    def _cleanup_session(self) -> None:
        if not self.session:
            return
        try:
            self.session.state_path.unlink(missing_ok=True)
        except Exception:
            logging.warning("m 模式清理快照失败: %s", self.session.state_path)
        session_temp = self.session.temp_root
        try:
            if session_temp.resolve().is_relative_to(self.temp_root.resolve()):
                shutil.rmtree(session_temp, ignore_errors=True)
        except Exception:
            logging.warning("m 模式清理临时目录失败: %s", session_temp)
        self.session = None

    def reset(self) -> None:
        with self.lock:
            source_package = self.source_package
            self._cleanup_session()
            self.exported = None
            self.source_package = None
            self.prepared_package = None
            if source_package is not None:
                try:
                    if source_package.resolve().is_relative_to(self.upload_root.resolve()):
                        source_package.unlink(missing_ok=True)
                        if source_package.parent != self.upload_root:
                            shutil.rmtree(source_package.parent, ignore_errors=True)
                except Exception:
                    logging.warning("m 模式清理上传临时文件失败: %s", source_package)

    def phase(self) -> str:
        if self.exported is not None:
            return "export"
        if self.session is not None:
            return "review"
        return "import"

    def render_root_page(self) -> str:
        with self.lock:
            if self.exported is not None:
                return render_export_page(self.exported)
            if self.session is not None:
                return HTML_PAGE
            return IMPORT_HTML_PAGE

    def state(self) -> dict[str, Any]:
        with self.lock:
            if self.session is None:
                return {"phase": "import", "exported": None}
            payload = self.session.state()
            payload["phase"] = self.phase()
            payload["sourcePackage"] = self.source_package.name if self.source_package else None
            if self.exported is not None:
                payload["exported"] = {
                    "package": str(self.exported.package),
                    "checklist": str(self.exported.checklist),
                }
            return payload

    def import_package_stream(self, filename: str, source: Any, length: int) -> dict[str, Any]:
        if length <= 0:
            raise ValueError("上传文件为空")
        safe_name = safe_upload_filename(filename)
        timestamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
        upload_dir = self.upload_root / timestamp
        upload_dir.mkdir(parents=True, exist_ok=True)
        uploaded_path = upload_dir / safe_name
        session_temp = self.temp_root / f"session_{timestamp}_{os.getpid()}"
        try:
            with uploaded_path.open("wb") as output:
                remaining = length
                while remaining > 0:
                    chunk = source.read(min(1024 * 1024, remaining))
                    if not chunk:
                        raise ValueError("上传文件读取不完整")
                    output.write(chunk)
                    remaining -= len(chunk)
            prepared_path = prepare_package_for_review(
                uploaded_path,
                self.rebuild_output_root,
                self.options.config_path,
                force_rebuild=True,
                database_path=self.options.database_path,
            )
            session_temp.mkdir(parents=True, exist_ok=True)
            session = ReviewSession(
                prepared_path,
                self.review_output_root,
                session_temp,
                **self._session_kwargs(),
            )
        except Exception:
            try:
                uploaded_path.unlink(missing_ok=True)
            except Exception:
                logging.warning("m 模式清理上传失败: %s", uploaded_path)
            shutil.rmtree(upload_dir, ignore_errors=True)
            shutil.rmtree(session_temp, ignore_errors=True)
            raise

        with self.lock:
            self.reset()
            self.source_package = uploaded_path
            self.prepared_package = prepared_path
            self.session = session
            logging.info("m 模式导入完成: %s -> %s", uploaded_path, prepared_path)
            return self.state()

    def load_package(self, package_path: Path) -> None:
        with self.lock:
            self.reset()
            timestamp = datetime.now().strftime("%Y%m%d%H%M%S%f")
            session_temp = self.temp_root / f"session_{timestamp}_{os.getpid()}"
            session_temp.mkdir(parents=True, exist_ok=True)
            self.source_package = package_path
            self.prepared_package = package_path
            self.session = ReviewSession(
                package_path,
                self.review_output_root,
                session_temp,
                **self._session_kwargs(),
            )
            logging.info("m 模式载入命令行指定包: %s", package_path)

    def export(self) -> ExportedFiles:
        with self.lock:
            if self.session is None:
                raise ValueError("尚未导入离线包")
            exported = self.session.export()
            self.exported = exported
            logging.info("m 模式导出完成: %s", exported.package)
            return exported

    def download_path(self, kind: str) -> Path:
        with self.lock:
            if self.exported is None:
                raise ValueError("尚未导出文件")
            if kind == "package":
                return self.exported.package
            if kind == "checklist":
                return self.exported.checklist
            raise ValueError(f"未知下载类型: {kind}")


class ReviewRequestHandler(BaseHTTPRequestHandler):
    server: "ReviewHTTPServer"

    def log_message(self, format: str, *args: Any) -> None:
        logging.info("Web %s - %s", self.address_string(), format % args)

    def send_json(self, value: Any, status: int = HTTPStatus.OK) -> None:
        payload = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(payload)

    def send_html(self, value: str, status: int = HTTPStatus.OK) -> None:
        payload = value.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(payload)

    def send_file(self, path: Path, filename: str) -> None:
        if not path.is_file():
            self.send_json({"error": "文件不存在"}, HTTPStatus.NOT_FOUND)
            return
        content_type = {
            ".zip": "application/zip",
            ".txt": "text/plain; charset=utf-8",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }.get(path.suffix.lower(), "application/octet-stream")
        payload_name = quote(filename)
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(path.stat().st_size))
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Disposition", f'attachment; filename="{payload_name}"; filename*=UTF-8\'\'{payload_name}')
        self.end_headers()
        with path.open("rb") as file:
            shutil.copyfileobj(file, self.wfile)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self.send_html(self.server.workspace.render_root_page())
            return
        if parsed.path == "/api/state":
            try:
                self.send_json(self.server.workspace.state())
            except Exception as exc:
                self.send_json({"error": str(exc)}, HTTPStatus.INTERNAL_SERVER_ERROR)
            return
        if parsed.path == "/download/package":
            try:
                path = self.server.workspace.download_path("package")
                self.send_file(path, path.name)
            except Exception as exc:
                self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
            return
        if parsed.path == "/download/checklist":
            try:
                path = self.server.workspace.download_path("checklist")
                self.send_file(path, path.name)
            except Exception as exc:
                self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
            return
        if parsed.path.startswith("/image/"):
            self.serve_image(parsed.path)
            return
        self.send_json({"error": "未找到资源"}, HTTPStatus.NOT_FOUND)

    def serve_image(self, path: str) -> None:
        parts = path.split("/")
        if len(parts) != 4:
            self.send_json({"error": "图片路径无效"}, HTTPStatus.BAD_REQUEST)
            return
        algo_type = unquote(parts[2])
        image_name = unquote(parts[3])
        if Path(algo_type).name != algo_type or Path(image_name).name != image_name:
            self.send_json({"error": "图片路径无效"}, HTTPStatus.BAD_REQUEST)
            return
        session = self.server.workspace.session
        if session is None:
            self.send_json({"error": "当前没有可浏览的包"}, HTTPStatus.NOT_FOUND)
            return
        image_path = session.image_tree / algo_type / image_name
        if not image_path.is_file():
            self.send_json({"error": "图片不存在"}, HTTPStatus.NOT_FOUND)
            return
        content_type = {
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".png": "image/png",
            ".webp": "image/webp",
        }.get(image_path.suffix.lower(), "application/octet-stream")
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(image_path.stat().st_size))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        with image_path.open("rb") as image:
            shutil.copyfileobj(image, self.wfile)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        length = int(self.headers.get("Content-Length", "0"))
        try:
            if parsed.path == "/api/import":
                filename = parse_qs(parsed.query).get("filename", ["offline_package.zip"])[0]
                result = self.server.workspace.import_package_stream(filename, self.rfile, length)
                self.send_json({"package": result.get("package"), "phase": result.get("phase", "review")})
                return
            if parsed.path == "/api/reset":
                self.server.workspace.reset()
                self.send_json({"ok": True, "phase": "import"})
                return
            body = self.rfile.read(length)
            payload = json.loads(body.decode("utf-8")) if body else {}
            session = self.server.workspace.session
            if parsed.path == "/api/mark":
                if session is None:
                    raise ValueError("尚未导入离线包")
                result = session.mark(
                    str(payload.get("algoType") or ""),
                    int(payload.get("index")),
                    int(payload.get("auditStatus")),
                    bool(payload.get("advance", True)),
                )
                self.send_json(result)
                return
            if parsed.path == "/api/navigate":
                if session is None:
                    raise ValueError("尚未导入离线包")
                result = session.navigate(
                    int(payload.get("categoryIndex")),
                    int(payload.get("imageIndex")),
                )
                self.send_json(result)
                return
            if parsed.path == "/api/mark-category-correct":
                if session is None:
                    raise ValueError("尚未导入离线包")
                result = session.mark_category_correct(
                    str(payload.get("algoType") or ""),
                )
                self.send_json(result)
                return
            if parsed.path == "/api/mark-recognition":
                if session is None:
                    raise ValueError("尚未导入离线包")
                result = session.mark_same_recognition(
                    str(payload.get("algoType") or ""),
                    str(payload.get("recognition") or ""),
                    int(payload.get("auditStatus")),
                )
                self.send_json(result)
                return
            if parsed.path == "/api/export":
                output = self.server.workspace.export()
                self.send_json({"package": str(output.package), "checklist": str(output.checklist)})
                if self.server.workspace.options.exit_after_export:
                    threading.Thread(target=self.server.shutdown, daemon=True).start()
                return
            self.send_json({"error": "未找到接口"}, HTTPStatus.NOT_FOUND)
        except (ValueError, TypeError, KeyError, IndexError, json.JSONDecodeError) as exc:
            self.send_json({"error": str(exc)}, HTTPStatus.BAD_REQUEST)
        except Exception as exc:
            logging.exception("Web 请求失败")
            self.send_json({"error": str(exc)}, HTTPStatus.INTERNAL_SERVER_ERROR)


class ReviewHTTPServer(ThreadingHTTPServer):
    # Each m-mode process must own its port independently on Windows.
    allow_reuse_address = False

    def __init__(self, address: tuple[str, int], workspace: ReviewWorkspace) -> None:
        super().__init__(address, ReviewRequestHandler)
        self.workspace = workspace


def create_review_server(host: str, preferred_port: int, workspace: ReviewWorkspace) -> ReviewHTTPServer:
    if preferred_port == 0:
        return ReviewHTTPServer((host, 0), workspace)
    if preferred_port < 0 or preferred_port > 65535:
        raise ValueError(f"端口号无效: {preferred_port}")

    last_error: OSError | None = None
    unavailable_errors = {errno.EADDRINUSE, errno.EACCES, 10013, 10048}
    for port in range(preferred_port, 65536):
        try:
            server = ReviewHTTPServer((host, port), workspace)
        except OSError as exc:
            if exc.errno not in unavailable_errors and getattr(exc, "winerror", None) not in unavailable_errors:
                raise
            last_error = exc
            logging.info("m 模式端口不可用，尝试下一个端口: %s", port)
            continue
        if port != preferred_port:
            logging.info("m 模式改用空闲端口: %s", port)
        return server

    raise OSError(f"从端口 {preferred_port} 开始没有找到空闲端口") from last_error


def local_lan_ip() -> str | None:
    candidates: list[str] = []
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as sock:
            sock.connect(("8.8.8.8", 80))
            append_lan_candidate(candidates, sock.getsockname()[0])
    except OSError:
        pass

    for hostname in {socket.gethostname(), socket.getfqdn()}:
        if not hostname:
            continue
        try:
            addresses = socket.getaddrinfo(hostname, None, socket.AF_INET, socket.SOCK_DGRAM)
        except OSError:
            continue
        for address in addresses:
            append_lan_candidate(candidates, address[4][0])

    for candidate in candidates:
        if ipaddress.ip_address(candidate).is_private:
            return candidate
    return candidates[0] if candidates else None


def usable_lan_ipv4(value: str | None) -> bool:
    if not value:
        return False
    try:
        ip = ipaddress.ip_address(value)
    except ValueError:
        return False
    return (
        ip.version == 4
        and not ip.is_loopback
        and not ip.is_unspecified
        and not ip.is_link_local
        and not ip.is_multicast
    )


def append_lan_candidate(candidates: list[str], value: str | None) -> None:
    if usable_lan_ipv4(value) and value not in candidates:
        candidates.append(value)


def raise_process_priority() -> None:
    if os.name != "nt":
        return
    try:
        import ctypes
        from ctypes import wintypes

        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.GetCurrentProcess.restype = wintypes.HANDLE
        kernel32.SetPriorityClass.argtypes = (wintypes.HANDLE, wintypes.DWORD)
        kernel32.SetPriorityClass.restype = wintypes.BOOL
        handle = kernel32.GetCurrentProcess()
        above_normal_priority_class = 0x00008000
        if not kernel32.SetPriorityClass(handle, above_normal_priority_class):
            logging.warning("m 模式提升进程优先级失败: Windows error %s", ctypes.get_last_error())
            return
        logging.info("m 模式进程优先级已提升为 Above normal")
    except Exception as exc:
        logging.warning("m 模式提升进程优先级失败: %s", exc)


def run_interactive_mode(
    package: Path | None,
    output_dir: Path,
    newwrap_dir: Path,
    config_path: Path,
    host: str,
    port: int,
    no_browser: bool,
    interaction_mode: str = "lan",
    database_path: Path | None = None,
) -> int:
    from tempfile import TemporaryDirectory

    save_problem_images = load_save_problem_images(config_path)
    recognition_red_keywords = load_yaml_list(config_path, RECOGNITION_RED_KEY)
    point_name_red_keywords = load_yaml_list(config_path, POINT_NAME_RED_KEY)
    point_name_end_open, point_name_end_characters, point_name_end_character_color = load_point_name_end_config(config_path)
    default_zoom = load_default_zoom(config_path)
    auto_click_times_ms = load_yaml_int_list(config_path, AUTO_CLICK_TIMES_KEY, [1000, 2000, 3000])
    show_auto_click = load_yaml_yes(config_path, SHOW_AUTO_CLICK_KEY)
    show_mark_all_correct = load_yaml_yes(config_path, SHOW_MARK_ALL_CORRECT_KEY)
    show_mark_same_recognition = load_yaml_yes(config_path, SHOW_MARK_SAME_RECOGNITION_KEY)
    web_hotkeys = load_web_hotkeys(config_path)
    newkeywords = load_yaml_list(config_path, NEW_KEYWORDS_KEY, keep_empty=True)
    default_wrong_keywords = load_yaml_list(config_path, DEFAULT_WRONG_KEY)
    default_problem_keywords = load_yaml_list(config_path, DEFAULT_PROBLEM_KEY)
    extra_xlsx_recognition_enabled = load_yaml_bool(config_path, EXTRA_XLSX_RECOGNITION_ENABLED_KEY, False)
    extra_xlsx_recognition_keywords = load_yaml_list(config_path, EXTRA_XLSX_RECOGNITION_KEY, keep_empty=True)
    show_algorithm_normal = load_yaml_bool(config_path, SHOW_ALGORITHM_NORMAL_KEY, False)
    checklist_image_width_enabled = load_yaml_bool(config_path, CHECKLIST_IMAGE_WIDTH_ENABLED_KEY, False)
    solution_rules = load_solution_rules(config_path)
    solution_options = load_solution_options(config_path)
    responsible_rules = load_responsible_rules(config_path)
    algo_type_labels = load_algo_type_labels(config_path)
    raise_process_priority()
    if interaction_mode == "local":
        package_path = choose_local_package(output_dir, package)
        server_host = "127.0.0.1"
        force_rebuild = False
    elif interaction_mode == "lan":
        if package is not None:
            package_candidate = package.expanduser().resolve()
            package_path = find_input_package(output_dir, package_candidate) if package_candidate.is_dir() else package_candidate
        else:
            package_path = None
        server_host = host
        force_rebuild = False
    else:
        raise ValueError(f"invalid m interaction mode: {interaction_mode}")
    if package_path is None:
        logging.info("m 模式启动导入页，等待上传离线包")
    else:
        logging.info("m 模式读取离线包: %s", package_path)
    options = ReviewSessionOptions(
        config_path=config_path,
        database_path=database_path,
        save_problem_images=save_problem_images,
        recognition_red_keywords=recognition_red_keywords,
        point_name_red_keywords=point_name_red_keywords,
        point_name_end_open=point_name_end_open,
        point_name_end_characters=point_name_end_characters,
        point_name_end_character_color=point_name_end_character_color,
        default_zoom=default_zoom,
        auto_click_times_ms=auto_click_times_ms,
        show_auto_click=show_auto_click,
        show_mark_all_correct=show_mark_all_correct,
        show_mark_same_recognition=show_mark_same_recognition,
        exit_after_export=True,
        web_hotkeys=web_hotkeys,
        newkeywords=newkeywords,
        default_wrong_keywords=default_wrong_keywords,
        default_problem_keywords=default_problem_keywords,
        extra_xlsx_recognition_enabled=extra_xlsx_recognition_enabled,
        extra_xlsx_recognition_keywords=extra_xlsx_recognition_keywords,
        show_algorithm_normal=show_algorithm_normal,
        checklist_image_width_enabled=checklist_image_width_enabled,
        algo_type_labels=algo_type_labels,
        solution_rules=solution_rules,
        solution_options=solution_options,
        responsible_rules=responsible_rules,
    )
    with TemporaryDirectory(prefix="review_package_") as temp_name:
        workspace = ReviewWorkspace(
            output_dir,
            newwrap_dir,
            Path(temp_name),
            options,
        )
        if package_path is not None:
            workspace.load_package(
                prepare_package_for_review(
                    package_path,
                    output_dir,
                    config_path,
                    force_rebuild=force_rebuild,
                    database_path=database_path,
                )
            )
        server = create_review_server(server_host, port, workspace)
        browser_host = "127.0.0.1" if server_host in ("", "0.0.0.0", "::") else server_host
        local_url = f"http://{browser_host}:{server.server_port}/"
        lan_ip = local_lan_ip() if interaction_mode == "lan" and server_host in ("", "0.0.0.0", "::") else None
        lan_url = f"http://{lan_ip}:{server.server_port}/" if lan_ip else None
        if interaction_mode == "lan" and not lan_url:
            workspace._cleanup_session()
            server.server_close()
            raise RuntimeError("无法检测服务主机局域网 IPv4，已拒绝启动不可供局域网访问的审核地址")
        url = lan_url or local_url
        logging.info("m 模式 Web 已启动: %s", url)
        print(f"m 模式 Web: {url}")
        if lan_url:
            logging.info("m 模式局域网访问: %s", lan_url)
            print(f"局域网访问: {lan_url}")
        if not no_browser:
            webbrowser.open(url)
        try:
            server.serve_forever()
        except KeyboardInterrupt:
            logging.info("m 模式 Web 手动停止")
            print("\nm 模式 Web 已停止")
        finally:
            exported = workspace.exported
            workspace._cleanup_session()
            server.server_close()
        if exported:
            logging.info("m 模式结束，输出: %s", exported.package)
    return 0


def write_reviewed_package_checklist(
    package: Path,
    output_dir: Path,
    config_path: Path,
) -> Path:
    from tempfile import TemporaryDirectory

    solution_rules = load_solution_rules(config_path)
    solution_options = load_solution_options(config_path)
    responsible_rules = load_responsible_rules(config_path)
    newkeywords = load_yaml_list(config_path, NEW_KEYWORDS_KEY, keep_empty=True)
    extra_xlsx_recognition_enabled = load_yaml_bool(config_path, EXTRA_XLSX_RECOGNITION_ENABLED_KEY, False)
    extra_xlsx_recognition_keywords = load_yaml_list(config_path, EXTRA_XLSX_RECOGNITION_KEY, keep_empty=True)
    show_algorithm_normal = load_yaml_bool(config_path, SHOW_ALGORITHM_NORMAL_KEY, False)
    checklist_image_width_enabled = load_yaml_bool(config_path, CHECKLIST_IMAGE_WIDTH_ENABLED_KEY, False)
    point_name_red_keywords = load_yaml_list(config_path, POINT_NAME_RED_KEY)
    algo_type_labels = load_algo_type_labels(config_path)
    with TemporaryDirectory(prefix="reviewed_checklist_") as temp_name:
        session = ReviewSession(
            package,
            output_dir,
            Path(temp_name),
            config_path=config_path,
            save_problem_images=True,
            recognition_red_keywords=[],
            point_name_red_keywords=point_name_red_keywords,
            point_name_end_open=False,
            point_name_end_characters=0,
            point_name_end_character_color="",
            default_zoom=1.0,
            auto_click_times_ms=[1000],
            show_auto_click=False,
            show_mark_all_correct=False,
            show_mark_same_recognition=False,
            exit_after_export=False,
            web_hotkeys=DEFAULT_WEB_HOTKEYS.copy(),
            newkeywords=newkeywords,
            default_wrong_keywords=[],
            default_problem_keywords=[],
            extra_xlsx_recognition_enabled=extra_xlsx_recognition_enabled,
            extra_xlsx_recognition_keywords=extra_xlsx_recognition_keywords,
            show_algorithm_normal=show_algorithm_normal,
            checklist_image_width_enabled=checklist_image_width_enabled,
            algo_type_labels=algo_type_labels,
            solution_rules=solution_rules,
            solution_options=solution_options,
            responsible_rules=responsible_rules,
        )
        workbook_path = session._write_problem_images_workbook()
        session._write_newkeywords_report(
            report_path=workbook_path.with_suffix(".txt"),
            ensure_exists=True,
        )
        try:
            session.state_path.unlink(missing_ok=True)
        except Exception:
            logging.warning("清单生成后清理快照失败: %s", session.state_path)
        return workbook_path
