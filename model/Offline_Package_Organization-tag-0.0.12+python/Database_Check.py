from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import datetime
import io
import json
import logging
import os
from pathlib import Path
import re
import shutil
import sqlite3
import sys
import uuid
import zipfile

from annotation_db import (
    INDEX_TABLE,
    STATION_COLUMNS,
    AnnotationDatabase,
    AnnotationRecord,
    find_metadata_value,
)


ROOT = Path(__file__).resolve().parent
DEFAULT_DATABASE = ROOT / "annotations.sqlite3"
DEFAULT_LOG = ROOT / "log.txt"
DEFAULT_SQL = ROOT / "annotations.sql"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="SQLite 标注数据库检查工具。")
    parser.add_argument(
        "mode",
        nargs="?",
        choices=("1", "2", "3", "4"),
        help="1=导入 ZIP；2=备份并导入 SQL；3=导出 Excel；4=拆分并校验数据库",
    )
    parser.add_argument("--input-dir", type=Path, help="模式 1 的 ZIP 文件夹")
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE, help="SQLite 数据库路径")
    parser.add_argument("--sql", type=Path, default=DEFAULT_SQL, help="模式 2 要导入的 SQL 文件")
    parser.add_argument("--backup-dir", type=Path, default=ROOT / "database_backups", help="模式 2 数据库备份目录")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "database_exports", help="模式 3 Excel 输出目录")
    parser.add_argument("--log", type=Path, default=DEFAULT_LOG, help="日志文件路径")
    return parser.parse_args()


def setup_logging(log_path: Path) -> None:
    log_path.parent.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(message)s",
        handlers=[
            logging.FileHandler(log_path, encoding="utf-8"),
            logging.StreamHandler(),
        ],
    )


def choose_mode(mode: str | None) -> str:
    if mode:
        return mode
    if not sys.stdin.isatty():
        return "1"

    print("Database_Check mode:")
    print("  1. 导入指定文件夹内 ZIP 的已标注 accuracy.json")
    print("  2. 备份数据库后导入 SQL")
    print("  3. 按 stationSN 分表导出 Excel")
    print("  4. 备份、拆分并校验现有数据库")
    selected = input("请选择 1/2/3/4，回车默认 1 [1] ").strip()
    if selected in ("", "1", "2", "3", "4"):
        return selected or "1"
    raise ValueError(f"模式无效: {selected}")


def require_path(path: Path | None, label: str) -> Path:
    if path is not None:
        return path.expanduser().resolve()
    if not sys.stdin.isatty():
        raise ValueError(f"{label} 必须通过命令行指定")
    value = input(f"请输入{label}: ").strip()
    if not value:
        raise ValueError(f"{label}不能为空")
    return Path(value).expanduser().resolve()


def parse_accuracy_time(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    normalized = text.replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(normalized).replace(tzinfo=None).isoformat(timespec="seconds")
    except ValueError:
        pass
    for pattern in ("%Y%m%d%H%M%S", "%Y%m%d"):
        try:
            return datetime.strptime(text, pattern).isoformat(timespec="seconds")
        except ValueError:
            continue
    if text.isdigit():
        numeric = int(text)
        if numeric > 10_000_000_000:
            numeric /= 1000
        try:
            return datetime.fromtimestamp(numeric).isoformat(timespec="seconds")
        except (OverflowError, OSError, ValueError):
            return None
    return None


def accuracy_updated_at(accuracy: dict[object, object], package_path: Path) -> str:
    for key in ("taskRunTime", "taskRuntime", "updatedAt", "updateTime", "runTime", "statisticDate"):
        timestamp = parse_accuracy_time(find_metadata_value(accuracy, key))
        if timestamp:
            return timestamp
    fallback = datetime.fromtimestamp(package_path.stat().st_mtime).isoformat(timespec="seconds")
    logging.warning("未找到 accuracy.json 审核时间，使用 ZIP 修改时间: %s -> %s", package_path, fallback)
    return fallback


def accuracy_documents(package_path: Path) -> list[dict[object, object]]:
    documents: list[dict[object, object]] = []
    with zipfile.ZipFile(package_path) as outer_archive:
        inner_names = [
            name
            for name in outer_archive.namelist()
            if Path(name).name.lower().startswith("accuracy_") and Path(name).suffix.lower() == ".zip"
        ]
        for inner_name in inner_names:
            with zipfile.ZipFile(io.BytesIO(outer_archive.read(inner_name))) as accuracy_archive:
                json_names = [
                    name
                    for name in accuracy_archive.namelist()
                    if Path(name).name.lower() == "accuracy.json"
                ]
                if len(json_names) != 1:
                    raise ValueError(f"{inner_name} 内 accuracy.json 数量异常: {len(json_names)}")
                document = json.loads(accuracy_archive.read(json_names[0]).decode("utf-8-sig"))
                if not isinstance(document, dict):
                    raise ValueError(f"{inner_name} 内 accuracy.json 不是对象")
                documents.append(document)
    return documents


def marked_count(accuracy: dict[object, object]) -> int:
    count = 0
    statistics = accuracy.get("data", {})
    if not isinstance(statistics, dict):
        return 0
    for statistic in statistics.get("statistic", []):
        if not isinstance(statistic, dict):
            continue
        for item in statistic.get("imagesData", []):
            if not isinstance(item, dict):
                continue
            try:
                audit_status = int(item.get("auditStatus", 0))
            except (TypeError, ValueError):
                continue
            if audit_status in (1, 2, 3):
                count += 1
    return count


def import_reviewed_packages(input_dir: Path, database_path: Path) -> tuple[int, int, int]:
    if not input_dir.is_dir():
        raise NotADirectoryError(f"ZIP 文件夹不存在: {input_dir}")

    database = AnnotationDatabase(database_path)
    package_count = 0
    document_count = 0
    marked_total = 0
    for package_path in sorted(input_dir.rglob("*.zip"), key=lambda path: str(path).casefold()):
        try:
            documents = accuracy_documents(package_path)
        except zipfile.BadZipFile:
            logging.warning("跳过损坏 ZIP: %s", package_path)
            continue
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            logging.warning("跳过非审核外层 ZIP: %s (%s)", package_path, exc)
            continue
        if not documents:
            continue

        package_count += 1
        for accuracy in documents:
            updated_at = accuracy_updated_at(accuracy, package_path)
            marked = marked_count(accuracy)
            database.upsert_marked_accuracy(
                accuracy,
                updated_at=updated_at,
                only_if_newer=True,
            )
            document_count += 1
            marked_total += marked
            logging.info(
                "导入审核标注: package=%s updated_at=%s marked=%s",
                package_path,
                updated_at,
                marked,
            )

    logging.info(
        "模式 1 完成: packages=%s accuracy_documents=%s marked=%s database=%s",
        package_count,
        document_count,
        marked_total,
        database_path,
    )
    return package_count, document_count, marked_total


def backup_and_import_sql(database_path: Path, sql_path: Path, backup_dir: Path) -> Path | None:
    if not sql_path.is_file():
        raise FileNotFoundError(f"SQL 文件不存在: {sql_path}")
    backup_path: Path | None = None
    if database_path.exists():
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / f"{database_path.stem}_{datetime.now():%Y%m%d%H%M%S}{database_path.suffix}"
        shutil.copy2(database_path, backup_path)
        logging.info("模式 2 已备份数据库: %s", backup_path)
    else:
        logging.info("模式 2 未找到现有数据库，直接导入 SQL: %s", database_path)

    temporary_database = database_path.with_name(
        f".{database_path.stem}.{uuid.uuid4().hex}{database_path.suffix}"
    )
    try:
        database_path.parent.mkdir(parents=True, exist_ok=True)
        sql_text = sql_path.read_text(encoding="utf-8-sig")
        connection = sqlite3.connect(temporary_database)
        try:
            if not re.search(
                r"\bCREATE\s+TABLE(?:\s+IF\s+NOT\s+EXISTS)?\s+[\"`\[]?annotations\b",
                sql_text,
                re.IGNORECASE,
            ):
                create_legacy_schema(connection)
            connection.executescript(sql_text)
            connection.commit()
        finally:
            connection.close()
        AnnotationDatabase(temporary_database)
        connection = sqlite3.connect(temporary_database)
        try:
            validate_annotation_schema(connection)
        finally:
            connection.close()
        database_path.parent.mkdir(parents=True, exist_ok=True)
        os.replace(temporary_database, database_path)
    except Exception:
        temporary_database.unlink(missing_ok=True)
        raise

    logging.info("模式 2 已导入 SQL: %s -> %s", sql_path, database_path)
    return backup_path


def create_legacy_schema(connection: sqlite3.Connection) -> None:
    connection.execute(
        """
        CREATE TABLE annotations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            station_sn TEXT NOT NULL,
            point_code TEXT NOT NULL,
            recognition TEXT NOT NULL,
            audit_status INTEGER NOT NULL CHECK (audit_status IN (1, 2, 3)),
            updated_at TEXT NOT NULL,
            UNIQUE (station_sn, point_code)
        )
        """
    )


def validate_annotation_schema(connection: sqlite3.Connection) -> None:
    table_names = {
        str(row[0])
        for row in connection.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type = 'table' AND name NOT LIKE 'sqlite_%'
            """
        ).fetchall()
    }
    if "annotations" in table_names:
        columns = {
            str(row[1])
            for row in connection.execute("PRAGMA table_info(annotations)").fetchall()
        }
        missing_columns = {
            "station_sn",
            "point_code",
            "recognition",
            "audit_status",
            "updated_at",
        } - columns
        if missing_columns:
            raise ValueError(
                f"SQL 导入后 annotations 表缺少字段: {sorted(missing_columns)}"
            )
        raise ValueError("SQL 导入后仍存在旧 annotations 大表")

    if INDEX_TABLE not in table_names:
        raise ValueError(f"SQL 导入后缺少 {INDEX_TABLE} 索引表")

    index_columns = {
        str(row[1])
        for row in connection.execute(
            f'PRAGMA table_info("{INDEX_TABLE}")'
        ).fetchall()
    }
    if index_columns != {"station_sn"}:
        raise ValueError(f"SQL 导入后 {INDEX_TABLE} 表字段不符合要求")
    index_info = connection.execute(
        f'PRAGMA table_info("{INDEX_TABLE}")'
    ).fetchone()
    if index_info is None or int(index_info[5]) != 1:
        raise ValueError(f"SQL 导入后 {INDEX_TABLE}.station_sn 必须是主键")

    indexed_stations = {
        str(row[0])
        for row in connection.execute(
            f'SELECT station_sn FROM "{INDEX_TABLE}"'
        ).fetchall()
    }
    station_tables: set[str] = set()
    for table_name in table_names - {INDEX_TABLE}:
        columns = {
            str(row[1])
            for row in connection.execute(
                f'PRAGMA table_info("{table_name.replace(chr(34), chr(34) * 2)}")'
            ).fetchall()
        }
        if columns == STATION_COLUMNS:
            station_tables.add(table_name)

    if any(station.casefold() == "testcode" for station in indexed_stations | station_tables):
        raise ValueError("SQL 导入后不允许存在 testCode stationSN")
    if indexed_stations != station_tables:
        missing_tables = sorted(indexed_stations - station_tables)
        missing_index = sorted(station_tables - indexed_stations)
        raise ValueError(
            "SQL 导入后索引表与 stationSN 子表不一致: "
            f"missing_tables={missing_tables} missing_index={missing_index}"
        )


def migrate_and_validate_database(database_path: Path, backup_dir: Path) -> Path | None:
    backup_path: Path | None = None
    if database_path.exists():
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / (
            f"{database_path.stem}_pre_split_{datetime.now():%Y%m%d%H%M%S}"
            f"{database_path.suffix}"
        )
        shutil.copy2(database_path, backup_path)
        logging.info("模式 4 已备份待拆分数据库: %s", backup_path)

    database = AnnotationDatabase(database_path)
    connection = sqlite3.connect(database_path)
    try:
        validate_annotation_schema(connection)
    finally:
        connection.close()
    records = len(database.get_all())
    stations = len(database.station_names())
    logging.info(
        "模式 4 数据库拆分并校验完成: stations=%s records=%s database=%s",
        stations,
        records,
        database_path,
    )
    return backup_path


def export_station_tables(database_path: Path, output_dir: Path) -> Path:
    if not database_path.is_file():
        raise FileNotFoundError(f"数据库不存在: {database_path}")
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError("模式 3 导出 Excel 需要 openpyxl") from exc

    records = AnnotationDatabase(database_path).get_all()
    grouped: dict[str, list[AnnotationRecord]] = defaultdict(list)
    for record in records:
        grouped[record.station_sn].append(record)

    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"annotations_{datetime.now():%Y%m%d%H%M%S}.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    if not grouped:
        sheet = workbook.create_sheet("empty")
        sheet.append(["stationSN", "pointCode", "recognition", "auditStatus", "updatedAt"])
    else:
        used_titles: set[str] = set()
        for index, (station_sn, station_records) in enumerate(sorted(grouped.items()), start=1):
            sheet = workbook.create_sheet(sheet_title(station_sn, used_titles))
            headers = ["stationSN", "pointCode", "recognition", "auditStatus", "updatedAt"]
            sheet.append(headers)
            for record in station_records:
                sheet.append(
                    [
                        record.station_sn,
                        record.point_code,
                        record.recognition,
                        record.audit_status,
                        record.updated_at,
                    ]
                )
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            table = Table(displayName=f"StationTable{index}", ref=f"A1:E{len(station_records) + 1}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
            sheet.freeze_panes = "A2"
            for column, width in zip(("A", "B", "C", "D", "E"), (28, 48, 24, 14, 22)):
                sheet.column_dimensions[column].width = width
    workbook.save(output_path)
    logging.info("模式 3 已导出数据库: records=%s stations=%s output=%s", len(records), len(grouped), output_path)
    return output_path


def sheet_title(station_sn: str, used_titles: set[str]) -> str:
    invalid = '[]:*?/\\'
    base = "".join("_" if character in invalid else character for character in station_sn).strip() or "station"
    base = base[:31]
    title = base
    counter = 2
    while title.casefold() in used_titles:
        suffix = f"_{counter}"
        title = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used_titles.add(title.casefold())
    return title


def main() -> int:
    args = parse_args()
    setup_logging(args.log)
    try:
        mode = choose_mode(args.mode)
        if mode == "1":
            input_dir = require_path(args.input_dir, "模式 1 ZIP 文件夹")
            import_reviewed_packages(input_dir, args.database)
        elif mode == "2":
            sql_path = require_path(args.sql, "模式 2 SQL 文件")
            backup_and_import_sql(args.database, sql_path, args.backup_dir)
        elif mode == "3":
            output_path = export_station_tables(args.database, args.output_dir)
            print(f"导出完成: {output_path}")
        else:
            backup_path = migrate_and_validate_database(args.database, args.backup_dir)
            if backup_path:
                print(f"拆分完成，备份: {backup_path}")
            else:
                print(f"拆分完成: {args.database}")
        return 0
    except Exception:
        logging.exception("Database_Check 处理失败")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
